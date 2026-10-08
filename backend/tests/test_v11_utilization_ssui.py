"""v1.1 验收门禁 G3(利用决策) / G4(课题三 SSUI 独立流程) / G6(追溯引导与权限)。"""
from __future__ import annotations

import io
import os

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
GEJIU = os.path.join(ROOT, "data", "raw", "3.20250731_重金属污染场地数据表(云南个旧)_最终版.xlsx")

import utilization as U  # noqa: E402  (conftest 已加入 ml/evaluation)
import ssui_v11 as SV  # noqa: E402

FEAS = {"value": 70.0, "label": "可行", "feasible": True, "source": "test"}
INFEAS = {"value": 30.0, "label": "不可行", "feasible": False, "source": "test"}
METALS_7 = ["As", "Cd", "Cr(VI)", "Cu", "Pb", "Hg", "Ni"]


def _clean_point(i=1, ph=6.8, **over):
    p = {"point": f"P{i}", "pH": ph, "Cd": 0.1, "Hg": 0.1, "As": 8, "Pb": 30, "Cr": 60, "Cu": 20, "Ni": 25,
         "Zn": 70, "Cr(VI)": 0.5}
    p.update(over)
    return p


# ─────────────── G3 利用决策引擎 ───────────────
def test_gejiu_real_data_fails_both_gates_and_score_cannot_offset():
    import pandas as pd
    df = pd.read_excel(GEJIU)
    cols = [c for c in df.columns if "_" in c and "mg/kg" in c]
    pts = [dict(point=r["采样点编号"], pH=r["pH"], **{c: r[c] for c in cols}) for _, r in df.iterrows()]
    d = U.decide("pre_remediation", pts, production_score=FEAS, ecology_score=FEAS, data_origin="client_real")
    assert d["decision_state"] == "neither_supported"
    for t in ("production", "ecology"):
        g = d[t]["gate"]
        assert g["state"] == "fail"
        assert {"As", "Pb"} <= set(g["exceed_control"])
        assert g["factors"]["Pb"]["n_exceed_control"] == 134
    assert "不可被评分抵消" in d["conclusion_text"]
    assert d["conclusion_text"].startswith("【修复前情景判断")
    assert d["is_post_remediation_conclusion"] is False
    assert set(d["remediation_targets"]["production"]) >= {"As", "Pb"}


def test_missing_required_factor_never_passes():
    p = _clean_point(); p.pop("Cd")
    d = U.decide("post_remediation", [p], production_score=FEAS, ecology_score=FEAS, eco_required=METALS_7)
    assert d["production"]["gate"]["state"] == "insufficient"
    assert d["ecology"]["gate"]["state"] == "insufficient"
    assert d["decision_state"] == "insufficient_evidence"
    assert any("Cd" in m for m in d["missing_evidence"])


def test_clean_site_production_supported_ecology_basic_items_missing():
    d = U.decide("post_remediation", [_clean_point(1), _clean_point(2)], production_score=FEAS, ecology_score=FEAS)
    assert d["production"]["gate"]["state"] == "pass"
    assert d["ecology"]["gate"]["state"] == "insufficient"  # GB 36600 VOC/SVOC 基本项目缺测
    assert d["decision_state"] == "production_supported"
    assert any("第一类用地" in a for a in d["assumptions"])


def test_both_supported_and_score_decides_only_after_gates():
    pts = [_clean_point(1), _clean_point(2)]
    d = U.decide("post_remediation", pts, production_score=FEAS, ecology_score=FEAS, eco_required=METALS_7)
    assert d["decision_state"] == "both_supported" and d["comparison"] is not None
    d2 = U.decide("post_remediation", pts, production_score=INFEAS, ecology_score=FEAS, eco_required=METALS_7)
    assert d2["decision_state"] == "ecology_supported"
    d3 = U.decide("post_remediation", pts, production_score=INFEAS, ecology_score=INFEAS, eco_required=METALS_7)
    assert d3["decision_state"] == "neither_supported"
    d4 = U.decide("post_remediation", pts, production_score=None, ecology_score=None, eco_required=METALS_7)
    assert d4["decision_state"] == "insufficient_evidence"


def test_ph_unknown_is_conservative_both_ways():
    # 其他农用地 Cd 筛选值 0.3(pH≤7.5)/0.6(pH>7.5); 0.45 在 pH 未知时无法判定
    d = U.decide("post_remediation", [_clean_point(ph=None, Cd=0.45)], production_score=FEAS)
    assert d["production"]["gate"]["state"] == "insufficient"
    assert "Cd" in d["production"]["gate"]["undetermined"]
    ok = U.decide("post_remediation", [_clean_point(ph=8.0, Cd=0.45)], production_score=FEAS, farmland_type="其他")
    assert ok["production"]["gate"]["state"] == "pass"
    # 超过所有 pH 档最宽管制值(4.0) → pH 未知也确定失败
    bad = U.decide("post_remediation", [_clean_point(ph=None, Cd=5.0)], production_score=FEAS)
    assert bad["production"]["gate"]["state"] == "fail"


def test_production_conditional_is_supported_with_safe_use_conditions():
    d = U.decide("post_remediation", [_clean_point(ph=6.8, Cd=0.8)], production_score=FEAS,
                 farmland_type="其他", eco_required=METALS_7, ecology_score=INFEAS)
    assert d["production"]["gate"]["state"] == "conditional"  # 0.3 < 0.8 ≤ 3.0
    assert d["decision_state"] == "production_supported"
    assert any("安全利用" in c for c in d["production"]["conditions"])


def test_ecology_conditional_requires_risk_assessment():
    d = U.decide("post_remediation", [_clean_point(Pb=500)], ecology_score=FEAS, eco_required=METALS_7,
                 production_score=INFEAS)
    assert d["ecology"]["gate"]["state"] == "conditional"  # 第一类 400 < 500 ≤ 800
    assert d["ecology"]["track_status"] == "insufficient"
    assert d["decision_state"] == "insufficient_evidence"


def test_demo_origin_is_labelled():
    d = U.decide("post_remediation", [_clean_point()], data_origin="monte_carlo_demo")
    assert any("仅供测试/演示" in a for a in d["assumptions"])


# ─────────────── G4 SSUI 计算 ───────────────
def test_ssui_weights_are_per_track_from_pptx():
    W = SV.load_weights()
    by = {i["code"]: i for i in W["indicators"]}
    assert (by["D1"]["w_production"], by["D1"]["w_ecology"]) == (0.1209, 0.1153)
    assert (by["D10"]["w_production"], by["D10"]["w_ecology"]) == (0.1348, 0.1406)
    assert (by["D22"]["w_production"], by["D22"]["w_ecology"]) == (0.3095, 0.3847)
    assert len(by) == 25


def test_ssui_not_clipped_and_formula_exact():
    W = SV.load_weights()
    scores = {i["code"]: 1.0 for i in W["indicators"]}
    r = SV.compute(scores, "production", t=0, M=1.1, W=W)
    v = W["criterion_weights"]["production"]
    gs = {c: sum(i["w_production"] for i in W["indicators"] if i["criterion"] == c) for c in v}
    expect = sum(v[c] * gs[c] for c in v) * 1.0 * 1.1
    assert r["status"] == "ok"
    assert r["ssui"] == pytest.approx(expect, abs=1e-6)
    assert r["ssui"] > 1.0 and r["exceeds_unit_range"] is True
    assert any("C4 组内权重和" in w for w in r["warnings"])
    half = SV.compute({k: 0.5 for k in scores}, "ecology", t=2, M=1.075, W=W)
    v2 = W["criterion_weights"]["ecology"]
    gs2 = {c: sum(i["w_ecology"] for i in W["indicators"] if i["criterion"] == c) for c in v2}
    assert half["ssui"] == pytest.approx(sum(v2[c] * 0.5 * gs2[c] for c in v2) * 1.06 * 1.075, abs=1e-6)


def test_ssui_missing_and_invalid_inputs():
    W = SV.load_weights()
    scores = {i["code"]: 0.7 for i in W["indicators"]}
    scores.pop("D25")
    r = SV.compute(scores, "production", 1, 1.15, W)
    assert r["status"] == "insufficient" and r["ssui"] is None and r["missing_indicators"] == ["D25"]
    scores["D25"] = 0.7
    assert SV.compute(scores, "production", 1, 1.3, W)["status"] == "invalid"  # M 超区间
    assert SV.compute(scores, "ecology", 1, 1.15, W)["status"] == "invalid"  # 生态 M ≤ 1.1


# ─────────────── G4 端到端: 模板 → 预览 → 确认 → 导出 → 修复后决策 ───────────────
def _login(c, user="admin"):
    tok = c.post("/api/v1/auth/login", json={"username": user, "password": "Demo@2026"}).json()["access_token"]
    return {"Authorization": f"Bearer {tok}"}


def _make_site(db, code="SSUIPOSTTEST"):
    from app.models import Site
    s = Site(site_code=code, name="课题三测试场地", pollution_type="heavy_metal", land_use_type="耕地")
    db.add(s); db.commit()
    return s.id, s.site_code


def _fill_template(content: bytes, site_code: str, score=0.72, bad_row=None, origin="真实数据(甲方/课题组提供)",
                   pollutants=True):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content))
    m = wb["批次信息"]
    m["B2"] = site_code; m["B4"] = 2026; m["B5"] = 3; m["B6"] = "中等强度"; m["B8"] = origin
    s = wb["指标得分"]
    for r in range(2, 27):
        s.cell(r, 7, score)
    if bad_row:
        s.cell(bad_row, 7, 1.5)
    if pollutants:
        p = wb["修复后污染物检测"]
        p.delete_rows(2, 1)
        r = 2
        for pc in ("A1", "A2"):
            for f, v in (("镉", 0.12), ("汞", 0.1), ("砷", 9), ("铅", 35), ("铬", 60), ("铜", 22), ("镍", 25), ("锌", 80),
                         ("六价铬", 0.4)):
                p.cell(r, 1, pc); p.cell(r, 2, 6.9); p.cell(r, 3, f); p.cell(r, 4, v); p.cell(r, 5, "mg/kg"); r += 1
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def test_ssui_post_full_workflow_preview_confirm_export_decide():
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from app.db.session import SessionLocal
    from app.main import app
    from app.models import (POST_REMEDIATION, EvaluationResult, Measurement, SSUIImportBatch, SSUIRecord,
                            UtilizationDecision)
    db = SessionLocal()
    try:
        sid, code = _make_site(db)
    finally:
        db.close()
    c = TestClient(app); h = _login(c)
    t = c.get("/api/v1/templates/ssui-post?track=production&site_code=" + code, headers=h)
    assert t.status_code == 200 and t.content[:2] == b"PK"
    filled = _fill_template(t.content, code)
    files = {"file": ("post.xlsx", filled, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    pv = c.post(f"/api/v1/sites/{sid}/ssui-post/preview", headers=h, files=files, data={"track": "production"})
    assert pv.status_code == 200, pv.text
    pj = pv.json()
    assert pj["can_confirm"] is True and pj["n_errors"] == 0, pj["errors"]
    assert pj["preview_calc"]["status"] == "ok" and pj["n_points"] == 2
    db = SessionLocal()
    try:  # 预览不写业务数据
        assert db.query(SSUIRecord).count() == 0
        assert db.query(Measurement).filter_by(site_id=sid).count() == 0
    finally:
        db.close()
    cf = c.post(f"/api/v1/ssui-post/batches/{pj['batch_id']}/confirm", headers=h)
    assert cf.status_code == 200, cf.text
    cj = cf.json()
    assert cj["calc"]["ssui"] == pytest.approx(pj["preview_calc"]["ssui"])
    db = SessionLocal()
    try:
        assert db.query(SSUIRecord).filter_by(batch_id=pj["batch_id"]).count() == 25
        ms = db.query(Measurement).filter_by(site_id=sid).all()
        assert len(ms) == 18 and all(m.stage == POST_REMEDIATION for m in ms)
        ev = db.get(EvaluationResult, cj["evaluation_id"])
        assert ev.stage == POST_REMEDIATION and ev.subproject == "S3" and ev.method_status == "provisional"
        assert db.get(SSUIImportBatch, pj["batch_id"]).status == "confirmed"
    finally:
        db.close()
    # 重复导入被拒
    pv2 = c.post(f"/api/v1/sites/{sid}/ssui-post/preview", headers=h, files=files).json()
    assert pv2["can_confirm"] is False and any("重复" in e["message"] for e in pv2["errors"])
    assert c.post(f"/api/v1/ssui-post/batches/{pv2['batch_id']}/confirm", headers=h).status_code == 409
    # 导出
    ex = c.get(f"/api/v1/ssui-post/batches/{pj['batch_id']}/export", headers=h)
    assert ex.status_code == 200
    wb = load_workbook(io.BytesIO(ex.content))
    kv = {r[0].value: r[1].value for r in wb["结果"].iter_rows() if r[0].value}
    assert kv["SSUI(未截断)"] == pytest.approx(cj["calc"]["ssui"]) and kv["数据阶段"] == "修复后(课题三)"
    assert wb["指标贡献"].max_row == 26
    # 修复后决策: 生产门禁通过 + SSUI → 结论; 生态缺 VOC/SVOC → 不支持生态结论
    dj = c.post(f"/api/v1/sites/{sid}/utilization?stage=post_remediation", headers=h).json()
    assert dj["is_post_remediation_conclusion"] is True
    assert dj["production"]["gate"]["state"] == "pass"
    assert dj["ecology"]["gate"]["state"] == "insufficient"
    exp = "production_supported" if cj["calc"]["feasible"] else "insufficient_evidence"
    if not cj["calc"]["feasible"]:
        exp = "neither_supported" if dj["ecology"]["track_status"] == "not_supported" else "insufficient_evidence"
    assert dj["decision_state"] == exp
    g = c.get(f"/api/v1/sites/{sid}/utilization?stage=post_remediation", headers=h).json()
    assert g["decision"]["decision_id"] == dj["decision_id"]
    db = SessionLocal()
    try:
        assert db.query(UtilizationDecision).filter_by(site_id=sid).count() == 1
    finally:
        db.close()


def test_ssui_post_validation_errors_carry_sheet_row_col():
    from fastapi.testclient import TestClient
    from app.db.session import SessionLocal
    from app.main import app
    db = SessionLocal()
    try:
        sid, code = _make_site(db, "SSUIPOSTBAD")
    finally:
        db.close()
    c = TestClient(app); h = _login(c)
    t = c.get("/api/v1/templates/ssui-post?track=ecology", headers=h).content
    filled = _fill_template(t, code, bad_row=5)
    r = c.post(f"/api/v1/sites/{sid}/ssui-post/preview", headers=h,
               files={"file": ("bad.xlsx", filled, "application/octet-stream")}).json()
    assert r["can_confirm"] is False
    assert {"sheet": "指标得分", "row": 5, "col": "G"}.items() <= next(
        e for e in r["errors"] if e["sheet"] == "指标得分").items()
    assert c.post(f"/api/v1/ssui-post/batches/{r['batch_id']}/confirm", headers=h).status_code == 409
    # 缺工作表 → 422
    from openpyxl import Workbook
    b = io.BytesIO(); Workbook().save(b)
    assert c.post(f"/api/v1/sites/{sid}/ssui-post/preview", headers=h,
                  files={"file": ("x.xlsx", b.getvalue(), "application/octet-stream")}).status_code == 422


def test_ssui_post_confirm_is_transactional(monkeypatch):
    from fastapi.testclient import TestClient
    from app.db.session import SessionLocal
    from app.main import app
    from app.models import Measurement, SSUIImportBatch, SSUIRecord
    from app.services import ssui_post_service as SP
    db = SessionLocal()
    try:
        sid, code = _make_site(db, "SSUIPOSTTX")
    finally:
        db.close()
    c = TestClient(app, raise_server_exceptions=False); h = _login(c)
    t = c.get("/api/v1/templates/ssui-post?track=production", headers=h).content
    pj = c.post(f"/api/v1/sites/{sid}/ssui-post/preview", headers=h,
                files={"file": ("tx.xlsx", _fill_template(t, code), "application/octet-stream")}).json()

    def boom(*a, **k):
        raise RuntimeError("计算故障注入")
    monkeypatch.setattr(SP.SV, "compute", boom)
    assert c.post(f"/api/v1/ssui-post/batches/{pj['batch_id']}/confirm", headers=h).status_code == 500
    db = SessionLocal()
    try:
        assert db.query(SSUIRecord).count() == 0
        assert db.query(Measurement).filter_by(site_id=sid).count() == 0
        assert db.get(SSUIImportBatch, pj["batch_id"]).status == "previewed"
    finally:
        db.close()


def test_demo_origin_marked_in_export():
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from app.db.session import SessionLocal
    from app.main import app
    db = SessionLocal()
    try:
        sid, code = _make_site(db, "SSUIPOSTDEMO")
    finally:
        db.close()
    c = TestClient(app); h = _login(c)
    t = c.get("/api/v1/templates/ssui-post?track=production", headers=h).content
    pj = c.post(f"/api/v1/sites/{sid}/ssui-post/preview", headers=h,
                files={"file": ("d.xlsx", _fill_template(t, code, origin="模拟数据——仅供测试/演示"),
                                "application/octet-stream")}).json()
    assert pj["data_origin"] == "monte_carlo_demo"
    c.post(f"/api/v1/ssui-post/batches/{pj['batch_id']}/confirm", headers=h)
    wb = load_workbook(io.BytesIO(c.get(f"/api/v1/ssui-post/batches/{pj['batch_id']}/export", headers=h).content))
    assert wb["结果"]["B1"].value == "模拟数据——仅供测试/演示, 不得用于正式报告"
    d = c.post(f"/api/v1/sites/{sid}/utilization?stage=post_remediation", headers=h).json()
    assert d["data_origin"] == "monte_carlo_demo"


# ─────────────── G6 追溯引导 / 进度 / 权限 ───────────────
def test_trace_guide_and_progress_create_no_records():
    from fastapi.testclient import TestClient
    from app.db.session import SessionLocal
    from app.main import app
    from app.models import WorkflowRecord
    from app.services.pipeline import run_import
    db = SessionLocal()
    try:
        before = db.query(WorkflowRecord).count()
        sid, _ = _make_site(db, "TRACEGUIDE")
    finally:
        db.close()
    c = TestClient(app); h = _login(c)
    g = c.get("/api/v1/trace/guide", headers=h).json()
    assert [s["name"] for s in g["stages"]] == ["调查评估", "方案审批", "施工监理", "效果评估", "后期管护"]
    assert all(c.get(t["url"], headers=h).status_code == 200 for t in g["templates"])
    p = c.get(f"/api/v1/sites/{sid}/trace/progress", headers=h).json()
    assert p["workflow_initialized"] is False and p["completed"] == 0
    assert p["next_step"] == "修复前数据导入"
    db = SessionLocal()
    try:
        assert db.query(WorkflowRecord).count() == before
        imp = run_import(db, GEJIU, "yunnan_gejiu")
        gid = imp["site_id"] if isinstance(imp, dict) else imp
    finally:
        db.close()
    p2 = c.get(f"/api/v1/sites/{gid}/trace/progress", headers=h).json()
    assert p2["milestones"][0]["done"] is True and p2["next_step"] == "障碍因子识别(课题一)"


def test_pre_template_is_importable(tmp_path):
    from app.api.v11 import build_pre_template
    from app.services.import_service import parse, smart_detect_and_map
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(build_pre_template()))
    ws = wb["检测数据"]
    for r in range(2, 5):
        vals = [f"S{r}", 103.1 + r / 100, 23.3, 0, 20, "红壤", 6.5, 0.2, 0.1, 15, 40, 70, 30, 30, 90, 0.5, 20, 1.2, 15, ""]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
    f = tmp_path / "pre.xlsx"; wb.save(f)
    _, mapping, _ = smart_detect_and_map(str(f))
    parsed = parse(str(f), mapping)
    assert parsed.n_points == 3 and parsed.n_measurements >= 3 * 9


def test_regulator_cannot_import_post_data_but_can_view():
    from fastapi.testclient import TestClient
    from app.db.session import SessionLocal
    from app.main import app
    db = SessionLocal()
    try:
        sid, code = _make_site(db, "PERMTEST")
    finally:
        db.close()
    c = TestClient(app); h = _login(c, "regulator")
    t = c.get("/api/v1/templates/ssui-post", headers=h)
    assert t.status_code == 200
    r = c.post(f"/api/v1/sites/{sid}/ssui-post/preview", headers=h,
               files={"file": ("x.xlsx", _fill_template(t.content, code), "application/octet-stream")})
    assert r.status_code == 403
    assert c.post(f"/api/v1/sites/{sid}/utilization", headers=h).status_code == 403
    assert c.get(f"/api/v1/sites/{sid}/utilization", headers=h).status_code == 200
    assert c.get(f"/api/v1/sites/{sid}/trace/progress", headers=h).status_code == 200


def test_gejiu_full_chain_import_evaluate_decide_via_api():
    """G3 参考示例: 真实个旧数据经导入→重构评价→修复前决策, 法规门禁否决(旧版判'可行 63.29')。"""
    from fastapi.testclient import TestClient
    from app.db.session import SessionLocal
    from app.main import app
    from app.services.evaluation_service import run_evaluation
    from app.services.pipeline import run_import
    db = SessionLocal()
    try:
        imp = run_import(db, GEJIU, "yunnan_gejiu")
        sid = imp["site_id"] if isinstance(imp, dict) else imp
        ev = run_evaluation(db, sid)
    finally:
        db.close()
    c = TestClient(app); h = _login(c)
    d = c.post(f"/api/v1/sites/{sid}/utilization?stage=pre_remediation", headers=h).json()
    assert d["decision_state"] == "neither_supported", d["conclusion_text"]
    assert d["production"]["gate"]["state"] == "fail" and d["ecology"]["gate"]["state"] == "fail"
    assert d["production"]["gate"]["factors"]["Pb"]["n_obs"] == 134
    assert d["data_origin"] in ("client_real", "field")
    assert not any("单位无法识别" in a for a in d["assumptions"]), d["assumptions"]
    # 删除准则层降维兜底(D-07)后, 个旧两轨重构均为"证据不足", 不再给出"可行 63.29"
    for k in ("reconstruction_prod", "reconstruction_eco"):
        assert ev[k]["grade"].startswith("证据不足"), (k, ev[k]["grade"])
    assert d["production"]["score"]["feasible"] is None
