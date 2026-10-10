"""v1.2.2 聚焦收尾回归(T01 SSUI 有效域 / T02 分析物身份 / T03 用途与适用性)。

独立期望值:
- T01: 25 项得分均 0.9, 生产 t=3, M=1.2 → 原值 = (1+0.03·3)·Σv_j·S_j·1.2, 由权重文件逐项计算(不调用被测函数)。
  等级区间仅定义到 [0, 1.0](ssui_weights_pptx_v1.json levels), 超出 → out_of_domain, 不给等级/支持。
- T02: GB 36600-2018 表2 官方记录(gb36600_2018_official.csv): DEHP(117-81-7) 一类筛选 42 / 二类 121 mg/kg;
  邻苯二甲酸二正辛酯(117-84-0) 一类 390 / 二类 2812 mg/kg; 无“邻苯二甲酸酯总量”限值。
- T03: 用途未选 → needs_manual_use_selection, 不得出现正式“支持”结论。
"""
from __future__ import annotations

import csv
import io
import os
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(ROOT, "ml", "evaluation"))
import ssui_v11 as S  # noqa: E402
import utilization as U  # noqa: E402

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
W = S.load_weights()
CODES = [i["code"] for i in W["indicators"]]


def _independent_ssui(score: float, track: str, t: float, M: float) -> float:
    inds = W["indicators"]
    v = W["criterion_weights"][track]
    tot = 0.0
    for c in ("C1", "C2", "C3", "C4"):
        tot += v[c] * sum(i["w_" + track] * score for i in inds if i["criterion"] == c)
    return (1 + 0.03 * t) * tot * M


# ───────────────────────── T01 ─────────────────────────
def test_t01_published_replay_is_out_of_domain_and_withheld():
    r = S.compute({c: 0.9 for c in CODES}, "production", 3, 1.2, W)
    exp = _independent_ssui(0.9, "production", 3, 1.2)
    assert exp == pytest.approx(1.191019, abs=1e-6)
    assert r["ssui"] == pytest.approx(exp, abs=1e-6)           # 原值保留, 不截断
    assert r["status"] == "out_of_domain"
    assert r["grade"] is None and r["feasible"] is None
    assert r["support_interpretation"] == "withheld"
    assert r["exceeds_unit_range"] is True and r["method_version"] == S.METHOD_VERSION


@pytest.mark.parametrize("raw,grade,feasible,status", [
    (1.0, "高度可持续", True, "ok"), (1.0 + 1e-6, None, None, "out_of_domain"),
    (0.8, "高度可持续", True, "ok"), (0.7999, "中度可持续", True, "ok"),
    (0.6, "中度可持续", True, "ok"), (0.5999, "低度可持续", False, "ok"),
    (0.4, "低度可持续", False, "ok"), (0.3999, "不可持续", False, "ok"), (0.0, "不可持续", False, "ok"),
    (-1e-6, None, None, "out_of_domain"),
])
def test_t01_classification_boundaries(raw, grade, feasible, status):
    c = S.classify(raw, W)
    assert (c["status"], c["grade"], c["feasible"]) == (status, grade, feasible)
    if status == "ok":
        assert c["classification_scope"] == "provisional_demonstration"


def test_t01_in_domain_result_is_labelled_provisional():
    r = S.compute({c: 0.7 for c in CODES}, "production", 0, 1.1, W)
    assert r["ssui"] == pytest.approx(_independent_ssui(0.7, "production", 0, 1.1), abs=1e-6)
    assert r["status"] == "ok" and r["classification_scope"] == "provisional_demonstration"
    assert r["support_interpretation"] == "provisional_demonstration"


def test_t01_missing_and_invalid_scores():
    sc = {c: 0.5 for c in CODES}; sc.pop(CODES[0])
    assert S.compute(sc, "ecology", 1, 1.05, W)["status"] == "insufficient"
    sc = {c: 0.5 for c in CODES}; sc[CODES[1]] = 1.2
    r = S.compute(sc, "ecology", 1, 1.05, W)
    assert r["status"] == "invalid" and r["grade"] is None and r.get("feasible") is None


def _pts(clean=True):
    base = {f: 0.1 for f in U.GB15618_REQUIRED}
    if not clean:
        base["Cd"] = 5.0  # > GB 15618 管制值(各 pH 档 ≤ 4.0)
    return [dict({"point": f"P{i}", "pH": 6.8}, **base, **{f: 0.0001 for f in U.GB36600_BASIC if f not in base})
            for i in range(5)]


def _score(r):
    return {"value": r.get("ssui"), "label": r.get("grade"), "feasible": r.get("feasible"), "source": "S3",
            "status": r.get("status")}


def test_t01_out_of_domain_score_never_creates_support():
    r = S.compute({c: 0.9 for c in CODES}, "production", 3, 1.2, W)
    d = U.decide("post_remediation", _pts(), farmland_type="其他", eco_land_class="第一类用地",
                 production_score=_score(r), ecology_score=_score(r))
    assert d["decision_state"] not in ("both_supported", "production_supported", "ecology_supported")
    assert d["production"]["track_status"] == "insufficient" and "有效域" in " ".join(d["production"]["conditions"])


def test_t01_gate_failure_still_unsupported_with_invalid_score():
    r = S.compute({c: 0.9 for c in CODES}, "production", 3, 1.2, W)
    d = U.decide("post_remediation", _pts(clean=False), farmland_type="其他", eco_land_class="第一类用地",
                 production_score=_score(r), ecology_score=_score(r))
    assert d["production"]["track_status"] == "not_supported"


def test_t01_score_from_eval_withholds_out_of_domain():
    from app.services.utilization_service import _score_from_eval

    class _E:
        id, eval_type, score, grade = 1, "ssui_post_production", 1.191019, None
        dimensions = {"status": "out_of_domain", "feasible": None, "support_interpretation": "withheld"}
    s = _score_from_eval(_E(), "课题三生产 SSUI")
    assert s["feasible"] is None and "有效域" in s["reason"]


# ───────────────────────── T02 ─────────────────────────
def _official():
    p = os.path.join(ROOT, "data", "standards", "gb36600_2018_official.csv")
    return {r["pollutant"]: r for r in csv.DictReader(open(p, encoding="utf-8-sig"))}


@pytest.mark.parametrize("name,canonical,identity", [
    ("邻苯二甲酸酯总量(PAEs)(mg/kg)", "SumPAE_ugkg", "family_total"),
    ("邻苯二甲酸酯_PAEs(mg/kg)", "SumPAE_ugkg", "family_total"),
    ("邻苯二甲酸二(2-乙基己基)酯(mg/kg)", "邻苯二甲酸二(2-乙基己基)酯", "exact_official"),
    ("DEHP(mg/kg)", "邻苯二甲酸二(2-乙基己基)酯", "exact_official"),
    ("邻苯二甲酸二正辛酯(mg/kg)", "邻苯二甲酸二正辛酯", "exact_official"),
    ("DnOP(μg/kg)", "邻苯二甲酸二正辛酯", "exact_official"),
    ("117-84-0(mg/kg)", "邻苯二甲酸二正辛酯", "exact_official"),
    ("邻苯二甲酸丁基苄酯(mg/kg)", "邻苯二甲酸丁基苄酯", "exact_official"),
    ("邻苯二甲酸二甲酯(mg/kg)", None, "mapping_review_required"),   # 未登记族成员
    ("邻苯二甲酸二异辛酯(mg/kg)", None, "mapping_review_required"),  # DIOP ≠ DEHP
    ("二甲苯(mg/kg)", "Xylenes_total", "family_total"),
    ("邻二甲苯(mg/kg)", "邻-二甲苯", "exact_official"),
    ("间二甲苯+对二甲苯(mg/kg)", "间-二甲苯+对-二甲苯", "exact_official"),
    ("1,2-二氯乙烷(mg/kg)", "1,2-二氯乙烷", "exact_official"),
    ("二氯乙烷(mg/kg)", None, "mapping_review_required"),
    ("二苯并呋喃(mg/kg)", None, "mapping_review_required"),
])
def test_t02_identity(name, canonical, identity):
    from app.services import factor_normalizer as FN
    c, m = FN.normalize_factor_name(name)
    assert (c, m["identity"]) == (canonical, identity)
    if identity == "exact_official":
        assert m["cas"] == _official()[m["official_name"]]["cas"]


def test_t02_substring_does_not_assign_other_analyte():
    from app.services import factor_normalizer as FN
    assert FN.normalize_factor_name("对硝基甲苯(mg/kg)")[0] is None   # 不是甲苯
    assert FN.normalize_factor_name("1-甲基萘(mg/kg)")[0] is None     # 不是萘
    assert FN.normalize_factor_name("甲苯(mg/kg)")[0] == "BTEX_Toluene"


def test_t02_header_parser_keeps_identity_parentheses():
    from app.services.import_service import split_header_unit as f
    assert f("邻苯二甲酸二(2-乙基己基)酯(mg/kg)") == ("邻苯二甲酸二(2-乙基己基)酯", "mg/kg")
    assert f("石油烃(C10-C40)(mg/kg)") == ("石油烃(C10-C40)", "mg/kg")
    assert f("多氯联苯(总量)") == ("多氯联苯(总量)", None)
    assert f("六价铬_Cr(VI)(mg/kg)") == ("六价铬", "mg/kg")
    assert f("阳离子交换量(cmol(+)/kg)") == ("阳离子交换量", "cmol(+)/kg")


def test_t02_resolver_binds_own_records_and_no_family_representative(client):
    from app.db.session import SessionLocal
    from app.services import threshold_resolver as TR
    off = _official()
    db = SessionLocal()
    try:
        for name in ("邻苯二甲酸二(2-乙基己基)酯", "邻苯二甲酸二正辛酯", "邻苯二甲酸丁基苄酯"):
            for land, col in (("第一类用地", "screening_cat1"), ("第二类用地", "screening_cat2")):
                r = TR.resolve_threshold_from_db(db, name, track="eco", land_use_type=land)
                assert r["threshold_resolution_status"] == "resolved", (name, land, r)
                assert r["threshold_value"] == pytest.approx(float(off[name][col])), (name, land, r)
        for fam in ("SumPAE_ugkg", "Xylenes_total", "BTEX_Xylene", "Phenol_Chlorophenol"):
            r = TR.resolve_threshold_from_db(db, fam, track="eco", land_use_type="第一类用地")
            assert r["threshold_resolution_status"] != "resolved" and r.get("threshold_value") is None, (fam, r)
            fb = TR.resolve_threshold_fallback(db, fam, track="eco")
            assert fb.get("threshold_value") is None, (fam, fb)
    finally:
        db.close()


def _login(c):
    tok = c.post("/api/v1/auth/login", json={"username": "admin", "password": "Demo@2026"}).json()["access_token"]
    return {"Authorization": f"Bearer {tok}"}


@pytest.fixture()
def client():
    from fastapi.testclient import TestClient
    from app.main import app
    with TestClient(app) as c:
        yield c, _login(c)


def _import(c, h, rows, hdr, name):
    from openpyxl import Workbook
    before = {s["id"] for s in c.get("/api/v1/sites", params={"page_size": 100}, headers=h).json()["items"]}
    wb = Workbook(); ws = wb.active; ws.append(hdr)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf)
    r = c.post("/api/v1/import", headers=h, data={"mapping_id": "auto", "on_conflict": "skip"},
               files={"file": (name, buf.getvalue(), XLSX)})
    assert r.status_code == 200, r.text[:300]
    return [s["id"] for s in c.get("/api/v1/sites", params={"page_size": 100}, headers=h).json()["items"]
            if s["id"] not in before][0]


PAE_HDR = ["采样点编号", "经度", "纬度", "深度_上限(cm)", "深度_下限(cm)", "土壤类型", "pH",
           "邻苯二甲酸酯总量_PAEs(mg/kg)", "邻苯二甲酸二(2-乙基己基)酯(mg/kg)", "邻苯二甲酸二正辛酯(μg/kg)",
           "邻苯二甲酸二甲酯(mg/kg)"]


@pytest.mark.parametrize("land,expect_official", [("第一类用地", ["邻苯二甲酸二(2-乙基己基)酯"]), ("第二类用地", [])])
def test_t02_api_real_db_phthalates(client, land, expect_official):
    c, h = client
    rows = [[f"Q{i:02d}", 120.2 + i * 0.001, 30.2, 0, 20, "潮土", 7.0, 50, 50, 50000, 5] for i in range(1, 6)]
    sid = _import(c, h, rows, PAE_HDR, "【模拟数据——仅供测试/演示】邻苯夹具_修复前检测.xlsx")
    k = c.post(f"/api/v1/sites/{sid}/kos-diagnosis?track=eco&subset=all&top_n=10&eco_land_class={land}",
               headers=h).json()
    official = [x["factor"] for x in k["key_obstacles"]]
    assert official == expect_official
    for x in k["key_obstacles"] + k["exploratory_obstacles"]:
        assert x["factor"] not in ("SumPAE_ugkg", "邻苯二甲酸二正辛酯")  # 总量无限值; DnOP 50 mg/kg < 390
    st = k["per_point_stats"]
    dnop = st["邻苯二甲酸二正辛酯"]
    assert dnop["n_exceed_points"] == 0 and dnop["max_value"] == pytest.approx(50.0)  # 50000 μg/kg = 50 mg/kg
    thr = {p.get("threshold") for p in dnop["point_details"]}
    assert thr == {390.0 if land == "第一类用地" else 2812.0}
    assert "SumPAE_ugkg" not in [x["factor"] for x in k["key_obstacles"]]
    md = {m["original_name"]: m for m in k["mapping_details"]}
    assert md["邻苯二甲酸酯总量"]["identity"] == "family_total"
    assert md["邻苯二甲酸二(2-乙基己基)酯"]["cas"] == "117-81-7" and md["邻苯二甲酸二正辛酯"]["cas"] == "117-84-0"
    # 未登记族成员: 不赋阈值, 列入未映射并在族群提示中要求人工复核
    assert "邻苯二甲酸二甲酯" in k["unmapped"]
    fam = {x["original_name"]: x for x in k["family_alerts"]}
    assert fam["邻苯二甲酸二甲酯"]["review_required"] is True
    assert not [x for x in k["key_obstacles"] if x["factor"] in ("邻苯二甲酸二甲酯",)]


# ───────────────────────── T03 ─────────────────────────
def _ok_scores():
    return _score(S.compute({c: 0.7 for c in CODES}, "production", 0, 1.1, W))


def test_t03_use_omitted_needs_manual_selection():
    d = U.decide("post_remediation", _pts(), production_score=_ok_scores(), ecology_score=_ok_scores())
    assert d["decision_state"] == "needs_manual_use_selection"
    assert d["use_state"]["production"] == "needs_manual_use_selection"
    assert d["use_state"]["ecology"] == "needs_manual_use_selection"
    assert d["hypothetical_screen"]["label"] == "hypothetical_conservative_screen"
    assert d["production"]["track_status"] == "withheld_use" and d["ecology"]["track_status"] == "withheld_use"


def test_t03_explicit_production_only():
    d = U.decide("post_remediation", _pts(), farmland_type="其他", production_score=_ok_scores(),
                 ecology_score=_ok_scores())
    assert d["production"]["track_status"] == "supported" and d["ecology"]["track_status"] == "withheld_use"
    assert d["decision_state"] == "production_supported"
    assert d["use_state"]["ecology"] == "needs_manual_use_selection"


def test_t03_non_construction_ecological_use_is_unresolved():
    d = U.decide("post_remediation", _pts(), farmland_type="其他", eco_land_class="非建设用地生态用途",
                 production_score=_ok_scores(), ecology_score=_ok_scores())
    assert d["use_state"]["ecology"] == "regulatory_applicability_unresolved"
    assert d["ecology"]["track_status"] == "withheld_use"


def test_t03_explicit_both_uses_formal():
    d = U.decide("post_remediation", _pts(), farmland_type="其他", eco_land_class="第一类用地",
                 production_score=_ok_scores(), ecology_score=_ok_scores())
    assert d["decision_state"] == "both_supported"
    assert d["use_state"] == {"production": "explicit", "ecology": "explicit"}


def test_t03_hard_fail_under_lenient_use_still_unsupported():
    d = U.decide("post_remediation", _pts(clean=False), production_score=_ok_scores(), ecology_score=_ok_scores())
    assert d["production"]["track_status"] == "not_supported"   # GB 15618 管制值与农用地类型无关
    assert d["decision_state"] != "both_supported"


def test_t03_api_invalid_and_missing_use(client):
    c, h = client
    rows = [[f"U{i:02d}", 120.3 + i * 0.001, 30.3, 0, 20, "潮土", 6.8, 0.1, 0.1, 5, 20, 50, 20, 20, 60]
            for i in range(1, 6)]
    hdr = ["采样点编号", "经度", "纬度", "深度_上限(cm)", "深度_下限(cm)", "土壤类型", "pH", "镉_Cd(mg/kg)",
           "汞_Hg(mg/kg)", "砷_As(mg/kg)", "铅_Pb(mg/kg)", "铬_Cr(mg/kg)", "铜_Cu(mg/kg)", "镍_Ni(mg/kg)", "锌_Zn(mg/kg)"]
    sid = _import(c, h, rows, hdr, "【模拟数据——仅供测试/演示】用途夹具_修复前检测.xlsx")
    bad = c.post(f"/api/v1/sites/{sid}/utilization", headers=h, params={"stage": "pre_remediation", "farmland_type": "旱地"})
    assert bad.status_code in (400, 422)
    bad2 = c.post(f"/api/v1/sites/{sid}/utilization", headers=h,
                  params={"stage": "pre_remediation", "eco_land_class": "生态用地"})
    assert bad2.status_code in (400, 422)
    d = c.post(f"/api/v1/sites/{sid}/utilization", headers=h, params={"stage": "pre_remediation"}).json()
    assert d["use_state"]["production"] == "needs_manual_use_selection"
    assert d["decision_state"] in ("needs_manual_use_selection", "insufficient_evidence")
    assert "supported" not in d["decision_state"]


# ═════════════════════ T04: 企业场地归属 / 操作人留痕 / 退回原因 / 阶段→证据追溯 ═════════════════════
def _ent_login(c, h, username, org):
    r = c.post("/api/v1/auth/register", json={"username": username, "password": "Test@2026!", "display_name": f"模拟{username}",
                                              "organization_name": org, "role_code": "enterprise"})
    if r.status_code == 200:
        assert c.post(f"/api/v1/auth/approve/{r.json()['user_id']}", headers=h).status_code == 200
    t = c.post("/api/v1/auth/login", json={"username": username, "password": "Test@2026!"}).json()["access_token"]
    return {"Authorization": f"Bearer {t}"}


def _ent_import(c, he, name):
    import io as _io
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "检测数据"
    ws.append(["采样点编号", "经度", "纬度", "深度_上限(cm)", "深度_下限(cm)", "pH", "镉_Cd(mg/kg)"])
    for i in range(3):
        ws.append([f"P{i}", 120.1 + i / 1000, 30.1, 0, 20, 6.2, 0.9 + i / 10])
    bio = _io.BytesIO(); wb.save(bio)
    r = c.post("/api/v1/import", data={"mapping_id": "auto", "on_conflict": "skip"}, headers=he,
               files={"file": (name, bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text[:300]
    return r.json()["site_id"]


def test_t04_enterprise_import_owns_site(client):
    c, h = client
    he = _ent_login(c, h, "t04_ent_a", "T04模拟企业A")
    sid = _ent_import(c, he, "MCT04-A_修复前检测数据_模拟数据_仅供测试演示.xlsx")
    ids = [x["id"] for x in c.get("/api/v1/sites", params={"page_size": 100}, headers=he).json()["items"]]
    assert sid in ids, "企业用户导入的场地应归属本企业并出现在其场地列表"
    assert c.get(f"/api/v1/sites/{sid}/workflow", headers=he).status_code == 200
    ho = _ent_login(c, h, "t04_ent_b", "T04模拟企业B")
    assert c.get(f"/api/v1/sites/{sid}/workflow", headers=ho).status_code == 403, "其他企业不得访问"


def test_t04_operator_from_token_and_return_needs_reason(client):
    c, h = client
    he = _ent_login(c, h, "t04_ent_c", "T04模拟企业C")
    sid = _ent_import(c, he, "MCT04-C_修复前检测数据_模拟数据_仅供测试演示.xlsx")
    assert c.post(f"/api/v1/sites/{sid}/workflow/init", headers=he).status_code == 200
    me = c.get("/api/v1/auth/me", headers=he).json()
    admin_id = c.get("/api/v1/auth/me", headers=h).json()["id"]
    r = c.post(f"/api/v1/sites/{sid}/workflow/survey", headers=he, json={"status": "in_progress", "operator_id": admin_id})
    s = next(x for x in r.json()["stages"] if x["stage"] == "survey")
    assert s["operator_id"] == me["id"] != admin_id, "operator_id 必须取自登录用户"
    r = c.post(f"/api/v1/sites/{sid}/workflow/survey", headers=h, json={"status": "returned", "is_returned": True})
    assert r.status_code == 400 and "原因" in r.text
    up = c.post(f"/api/v1/sites/{sid}/workflow/survey/attachment", headers=he, data={"file_role": "调查评估报告"},
                files={"file": ("H1_模拟材料.pdf", b"%PDF-1.4 synthetic test", "application/pdf")})
    assert up.status_code == 200
    r = c.post(f"/api/v1/sites/{sid}/workflow/survey", headers=h,
               json={"status": "completed", "is_completed": True, "review_comment": "模拟审核通过"})
    assert r.status_code == 200
    snap = c.get(f"/api/v1/sites/{sid}/evaluation-snapshot", headers=he).json()["snapshot"]
    sv = next(x for x in snap["workflow"]["business_stages"] if x["stage"] == "survey")
    import hashlib as _h
    assert sv["status"] == "completed" and sv["operator_name"] and sv["review_comment"] == "模拟审核通过"
    assert [e["sha256_12"] for e in sv["evidence"]] == [_h.sha256(b"%PDF-1.4 synthetic test").hexdigest()[:12]]


def test_t01_legacy_v121_record_out_of_domain_on_read(client):
    """v1.2.1 入库(dims 无 status, grade=高度可持续)的 SSUI=1.019189 记录: 升级后读取即判为域外, 不分级、不支持;
    库内原等级保留供审计。用真实的 v1.2.1 旧得分夹具经 API 入库, 再还原为 v1.2.1 的存储形态。"""
    import glob
    from openpyxl import load_workbook
    c, h = client
    rows = [[f"L{i:02d}", 120.3 + i * 0.001, 30.3, 0, 20, "潮土", 6.8, 0.2, 0.3, 30, 10] for i in range(1, 4)]
    sid = _import(c, h, rows, ["采样点编号", "经度", "纬度", "深度_上限(cm)", "深度_下限(cm)", "土壤类型", "pH",
                               "镉_Cd(mg/kg)", "汞_Hg(mg/kg)", "铅_Pb(mg/kg)", "砷_As(mg/kg)"],
                  "【模拟数据——仅供测试/演示】旧记录夹具_修复前检测.xlsx")
    code = c.get(f"/api/v1/sites/{sid}", headers=h).json()["site_code"]
    fx = glob.glob(os.path.join(ROOT, "demo", "mc_v12", "fixtures", "ssui_out_of_domain", "A_production_*.xlsx"))[0]
    wb = load_workbook(fx); wb["批次信息"]["B2"] = code
    buf = io.BytesIO(); wb.save(buf)
    pv = c.post(f"/api/v1/sites/{sid}/ssui-post/preview", headers=h, data={"track": "production"},
                files={"file": ("legacy.xlsx", buf.getvalue(), XLSX)}).json()
    assert pv["can_confirm"], pv
    calc = c.post(f"/api/v1/ssui-post/batches/{pv['batch_id']}/confirm", headers=h).json()["calc"]
    assert calc["status"] == "out_of_domain" and calc["grade"] is None
    from app.db.session import SessionLocal
    from app.models import EvaluationResult
    db = SessionLocal()
    try:
        ev = (db.query(EvaluationResult).filter_by(site_id=sid, eval_type="ssui_post_production")
              .order_by(EvaluationResult.id.desc()).first())
        ev.grade = "高度可持续"   # v1.2.1 的存储形态: ssui≥1 记为最高档, dims 无 status
        ev.dimensions = {k: v for k, v in (ev.dimensions or {}).items()
                         if k not in ("status", "support_interpretation", "classification_scope", "validity_domain", "domain_note")}
        ev.dimensions = {**ev.dimensions, "feasible": True}
        db.commit(); eid = ev.id
    finally:
        db.close()
    sj = c.get(f"/api/v1/sites/{sid}/evaluation-snapshot", headers=h).json()
    r = sj["snapshot"]["ssui"]["post"]["production"]["evaluation"]
    assert r["evaluation_id"] == eid
    assert r["status"] == "out_of_domain" and r["grade"] is None and r["feasible"] is None
    assert r["legacy_record"] is True and r["stored_grade"] == "高度可持续" and r["score"] == pytest.approx(1.019189)
    assert sj["snapshot"]["ssui"]["headline_grade"]["production"] == "超出有效域, 不分级"
    bl = c.get(f"/api/v1/sites/{sid}/ssui-post/batches", headers=h).json()["batches"]
    assert bl[0]["grade"] is None and bl[0]["ssui_status"] == "out_of_domain"
    dq = c.post(f"/api/v1/sites/{sid}/utilization", headers=h,
                params={"stage": "post_remediation", "farmland_type": "水田", "eco_land_class": "第一类用地"}).json()
    assert dq["production"]["track_status"] != "supported"
    assert dq["decision_state"] not in ("both_supported", "production_supported")
