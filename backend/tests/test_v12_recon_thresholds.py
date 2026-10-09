"""v1.2 年度验收门禁: 冻结方法 M-REC-2025、28 项指标导入、GB 36600/GB 15618 官方阈值。"""
from __future__ import annotations

import csv
import io
import os

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
STD = os.path.join(ROOT, "data", "standards")
SUBPROJECT_WB = "/Users/lensis/Desktop/SRS/05_其他子课题提供/Special for Python testing (Gejiu, Yunnan).xlsx"

import reconstruction_m2025 as M  # noqa: E402
import utilization as U  # noqa: E402


def _off():
    with open(os.path.join(STD, "gb36600_2018_official.csv"), encoding="utf-8") as fh:
        return {r["pollutant"]: r for r in csv.DictReader(fh)}


# ───────────── GB 36600 / GB 15618 ─────────────
VOC12 = {  # 标准原文(生态环境部 PDF 表1): (CAS, 筛1, 筛2, 管1, 管2)
    "氯甲烷": ("74-87-3", 12, 37, 21, 120), "1,1-二氯乙烷": ("75-34-3", 3, 9, 20, 100),
    "1,1-二氯乙烯": ("75-35-4", 12, 66, 40, 200), "顺-1,2-二氯乙烯": ("156-59-2", 66, 596, 200, 2000),
    "反-1,2-二氯乙烯": ("156-60-5", 10, 54, 31, 163), "1,2-二氯丙烷": ("78-87-5", 1, 5, 5, 47),
    "1,1,1,2-四氯乙烷": ("630-20-6", 2.6, 10, 26, 100), "1,1,2,2-四氯乙烷": ("79-34-5", 1.6, 6.8, 14, 50),
    "1,1,1-三氯乙烷": ("71-55-6", 701, 840, 840, 840), "1,1,2-三氯乙烷": ("79-00-5", 0.6, 2.8, 5, 15),
    "1,2,3-三氯丙烷": ("96-18-4", 0.05, 0.5, 0.5, 5), "间-二甲苯+对-二甲苯": ("108-38-3,106-42-3", 163, 570, 500, 570),
}


def test_gb36600_official_table_complete_and_12_voc_verified():
    off = _off()
    assert len(off) == 85
    assert sum(1 for r in off.values() if r["table"] == "表1") == 45
    for name, (cas, s1, s2, c1, c2) in VOC12.items():
        r = off[name]
        assert r["cas"] == cas
        assert (float(r["screening_cat1"]), float(r["screening_cat2"]), float(r["control_cat1"]), float(r["control_cat2"])) == (s1, s2, c1, c2)


def test_utilization_loads_all_45_basic_items_with_both_values():
    std = U.load_standards()
    for item in U.GB36600_BASIC:
        for lc in ("第一类用地", "第二类用地"):
            rec = std["gb36600"].get((U._key(item), lc))
            assert rec and rec["screening"] is not None and rec["control"] is not None, (item, lc)


@pytest.mark.parametrize("name,cls,s,c", [("氯甲烷", "第一类用地", 12, 21), ("氯甲烷", "第二类用地", 37, 120),
                                          ("1,2,3-三氯丙烷", "第一类用地", 0.05, 0.5)])
def test_voc_boundary_classification(name, cls, s, c):
    pts = [{"point": "P1", "pH": 7.0, "As": 1, "Cd": 0.1, "Cr(VI)": 0.1, "Cu": 1, "Pb": 1, "Hg": 0.01, "Ni": 1}]
    base = {k: 0.0 for k in U.GB36600_BASIC if k not in pts[0]}
    for v, expect in ((s, "pass"), (s * 1.01, "conditional"), (c * 1.01, "fail")):
        p = {**pts[0], **base, name if name != "间-二甲苯+对-二甲苯" else "间二甲苯+对二甲苯": v}
        g = U.ecology_gate([p], U.load_standards(), land_class=cls)
        assert g["state"] == expect, (name, cls, v, g["state"])


def test_seed_rows_are_official_only():
    from app.db.load_standard_thresholds import seed_rows
    rows = [r for r in seed_rows() if r["standard_code"] == "GB 36600-2018"]
    names = {r["factor_name"] for r in rows}
    for bad in ("荧蒽", "芘", "蒽", "菲", "苯并[j]荧蒽", "多环芳烃总量", "有机氯农药", "DDT类", "六六六"):
        assert bad not in names
    pcb = [r for r in rows if r["factor_name"] == "多氯联苯" and r["land_use_type"] == "第一类用地"][0]
    assert (pcb["screening_value"], pcb["control_value"]) == (0.14, 1.4)
    assert all(r["control_value"] is not None for r in rows)


def test_gb15618_copper_orchard_vs_other_and_other_items():
    std = U.load_standards()
    # 表1 铜: 果园 150/150/200/200, 其他 50/50/100/100; 水田属"其他"
    assert std["gb15618"][("Cu", "screening", "果园")]["pH<=5.5"]["value"] == 150
    assert std["gb15618"][("Cu", "screening", "水田")]["pH<=5.5"]["value"] == 50
    assert std["gb15618"][("BHC", "screening", "其他")]["pH>7.5"]["value"] == 0.10
    assert std["gb15618"][("BaP", "screening", "水田")]["pH<=5.5"]["value"] == 0.55
    assert std["gb15618"][("Cd", "control", "all")]["pH>7.5"]["value"] == 4.0


def test_kb_parser_corrects_chloromethane_scenario_and_exponents():
    import sys
    sys.path.insert(0, os.path.join(ROOT, "ml", "etl"))
    import load_knowledge_base as LK
    _, rules = LK.parse_knowledge_base(os.path.join(ROOT, "data", "knowledge_base", "统一障碍因子知识库_V1.0.csv"))
    fixes = {(c["factor"], c["field"]) for c in LK.KB_CORRECTIONS}
    assert ("氯甲烷", "application_scenario") in fixes and ("二噁英类(总毒性当量)", "threshold_max") in fixes
    assert len(LK.KB_CORRECTIONS) == 10
    from app.services.threshold_resolver import build_pollutant_limits
    L = build_pollutant_limits(os.path.join(ROOT, "data", "knowledge_base", "统一障碍因子知识库_V1.0.csv"))
    assert L["氯甲烷"]["ecology"]["特殊绿地"][0]["limit"] == 12.0
    assert L["氯甲烷"]["ecology"]["一般绿地"][0]["limit"] == 37.0


# ───────────── M-REC-2025 方法 ─────────────
def test_weight_tables_are_raw_and_not_normalised():
    assert round(sum(M.T218.values()), 4) == 0.9643 and len(M.T218) == 26
    assert round(sum(M.T219.values()), 4) == 1.0001
    assert round(sum(M.T221.values()), 4) == 1.0063
    assert "soc" not in M.T218 and "profile" not in M.T218


@pytest.mark.parametrize("fid,raw,F", [
    ("bulk_density", "Moderate", 100), ("bulk_density", "Slightly light", 50), ("bulk_density", "heavy", 50),
    ("biodiversity", "Not rich", 50), ("biodiversity", "Moderate", 80), ("biodiversity", "rich", 100),
    ("salinization", "Moderate to severe", 30), ("salinization", "Mild to moderate", 60), ("salinization", "None to mild", 100),
    ("irrigation_drainage", "Basic satisfaction or dissatisfaction", 20),
    ("irrigation_drainage", "Satisfaction or basic satisfaction", 50),
    ("irrigation_drainage", "Fully satisfaction or satisfaction", 100),
    ("texture", "Gravely soil", 40), ("texture", "sandy soil", 70), ("texture", "clay", 90), ("texture", "loam", 100),
    ("carbon_factor", "Temperate/Northern Temperate humid areas ", 69), ("carbon_factor", 0.8, 80),
    ("profile", "clay/sand/clay、Whole body clay、clay/sand/sand", 50), ("profile", "sand/clay/clay", 60),
    ("profile", "sand/clay/sand、loam/clay/clay、loam/sand/sand", 70), ("profile", "Whole body loam、loam/sand/loam", 100),
    ("profile", "Whole body sandy", 40)])
def test_categorical_mapping_production(fid, raw, F):
    it = M.score_indicator(fid, raw, "production")
    assert it["status"] == "scored" and it["F"] == F, it


def test_ecology_profile_group_without_rule_is_not_scored():
    it = M.score_indicator("profile", "clay/sand/clay、Whole body clay、clay/sand/sand", "ecology")
    assert it["status"] == "no_rule" and it["F"] is None


def test_invalid_category_is_invalid_not_default():
    it = M.score_indicator("texture", "very sticky", "production")
    assert it["status"] == "invalid" and it["F"] is None


@pytest.mark.parametrize("v,F", [(4.9, 20), (5.2, 40), (5.5, 60), (6.0, 80), (6.5, 100), (7.5, 100), (7.9, 80), (8.4, 60), (8.7, 40), (9.0, 20)])
def test_ph_production_grades(v, F):
    assert M._ph_prod(v) == F


def test_pollutant_three_grades_and_no_control_two_grades():
    assert M.score_pollutant("cd", 0.2, "production", 6.8, "其他")[0] == 100
    assert M.score_pollutant("cd", 0.5, "production", 6.8, "其他")[0] == 50
    assert M.score_pollutant("cd", 3.5, "production", 6.8, "其他")[0] == 10  # 管制值 3.0
    assert M.score_pollutant("cu", 9999, "production", 6.8, "其他")[0] == 50  # 无管制值 → 仅 50/100
    assert M.score_pollutant("bhc", 5.0, "production", 6.8, None)[0] == 50
    assert M.score_pollutant("as", 130, "ecology", None, None)[0] == 10  # GB 36600 一类管制值 120
    assert M.score_pollutant("zn", 10, "ecology", None, None)[0] is None


def _full_point(**over):
    p = {"cd": 0.2, "hg": 0.1, "as": 10, "pb": 30, "cr": 60, "cu": 20, "ni": 25, "zn": 80, "bhc": 0.01, "ddt": 0.01, "bap": 0.1,
         "soil_depth": 120, "ph": 6.8, "bulk_density": "适中", "biodiversity": "丰富", "salinization": "无、轻度",
         "total_n": 1.5, "avail_p": 20, "avail_k": 150, "cec": 20, "light_temp": 3500, "slope": 1, "irrigation_drainage": "充分满足、满足",
         "groundwater_depth": 4, "texture": "壤土", "carbon_factor": "温带/北温带干燥", "soc": 25, "profile": "通体壤"}
    p.update(over)
    return p


def test_full_path_is_sum_f_times_t_without_normalisation():
    r = M.evaluate(_full_point(), "production", land_subtype="旱地")
    assert r["path"] == "full"
    nonmax = {"light_temp": 80, "carbon_factor": 80, "soil_depth": 90}  # 3500 指数 / 温带干燥 0.8 / 120 cm
    exp = sum(M.T218[f] * nonmax.get(f, 100) for f in M.T218)
    assert abs(r["score"] - round(exp, 2)) < 1e-9
    assert r["score"] < 96.44  # 理论最高 96.43(表2.18 Σ=0.9643)


def test_missing_data_path_uses_table_2_19_and_nemerow():
    p = _full_point(); p.pop("bhc"); p.pop("light_temp")
    r = M.evaluate(p, "production", land_subtype="旱地")
    assert r["is_insufficient"] and "光温生产潜力/气候生产潜力" in r["missing_indicators"]
    p = _full_point(); p.pop("bhc")
    r = M.evaluate(p, "production", land_subtype="旱地")
    assert r["path"] == "missing_data" and r["score"] is not None
    assert any("内梅罗" in (d.get("rule") or "") for d in r["dimensions"])


def test_safety_gate_still_overrides_feasible_score():
    p = _full_point(cd=5.0)
    r = M.evaluate(p, "production", land_subtype="旱地")
    assert r["grade"] == "可行"  # 评分本身可能>50
    d = U.decide("pre_remediation", [{"point": "P1", "pH": 6.8, "Cd": 5.0, "Hg": 0.1, "As": 10, "Pb": 30, "Cr": 60,
                                       "Cu": 20, "Ni": 25, "Zn": 80}],
                 production_score={"value": r["score"], "label": r["grade"], "feasible": True, "source": "S2"})
    assert d["production"]["gate"]["state"] == "fail"
    assert d["decision_state"] != "production_supported"


# ───────────── 导入流程(API) ─────────────
def _login(c, user="admin"):
    tok = c.post("/api/v1/auth/login", json={"username": user, "password": "Demo@2026"}).json()["access_token"]
    return {"Authorization": f"Bearer {tok}"}


def _site(code):
    from app.db.session import SessionLocal
    from app.models import Site
    db = SessionLocal()
    try:
        s = Site(site_code=code, name="课题二导入测试", pollution_type="heavy_metal", land_use_type="耕地")
        db.add(s); db.commit()
        return s.id
    finally:
        db.close()


def _wide_xlsx(rows, header=None):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    header = header or (["location"] + [M.FEATURES[f]["en"][0] for f in M.RECON_28])
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def _row(label, **over):
    p = _full_point(**over)
    return [label] + [p.get(f) for f in M.RECON_28]


def test_recon_import_full_workflow_with_numeric_text_and_excel_error():
    from fastapi.testclient import TestClient
    from app.main import app
    sid = _site("RECONIMPA")
    c = TestClient(app); h = _login(c)
    t = c.get("/api/v1/templates/recon-pre?site_code=RECONIMPA", headers=h)
    assert t.status_code == 200 and t.content[:2] == b"PK"
    rows = [_row("A1", total_n="1.2\u00a0"), _row("A2", avail_p="15.5\t\t", cec="#VALUE!"), _row("A3")]
    files = {"file": ("wide.xlsx", _wide_xlsx(rows), "application/octet-stream")}
    pv = c.post(f"/api/v1/sites/{sid}/recon/preview", headers=h, files=files,
                data={"data_origin": "monte_carlo_demo", "provenance_status": "unverified", "land_subtype": "旱地"}).json()
    assert pv["can_confirm"], pv["errors"]
    assert {x["cell"] for x in pv["cleaning_log"]} >= {"R2", "S3"}
    assert any("#VALUE!" in w["message"] for w in pv["warnings"])
    cf = c.post(f"/api/v1/recon/batches/{pv['batch_id']}/confirm", headers=h)
    assert cf.status_code == 200, cf.text
    ev = cf.json()["evaluation"]
    assert ev["production"]["grade"] in ("可行", "不可行") and ev["production"]["path"] == "full"  # A2 缺 CEC 不影响中位数
    ex = c.get(f"/api/v1/recon/batches/{pv['batch_id']}/export", headers=h)
    assert ex.status_code == 200 and ex.content[:2] == b"PK"
    # 重复导入同一文件被拒绝
    pv2 = c.post(f"/api/v1/sites/{sid}/recon/preview", headers=h, files=files, data={"data_origin": "monte_carlo_demo"}).json()
    assert not pv2["can_confirm"] and any("已在批次" in e["message"] for e in pv2["errors"])
    # 利用决策读取课题二点位污染物
    d = c.post(f"/api/v1/sites/{sid}/utilization", headers=h, params={"stage": "pre_remediation"}).json()
    assert any("课题二指标批次" in a for a in d["assumptions"])


def test_recon_import_rejects_invalid_category_duplicate_point_and_post_stage():
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from app.main import app
    sid = _site("RECONIMPB")
    c = TestClient(app); h = _login(c)
    rows = [_row("B1", texture="sticky"), _row("B1"), _row("B3", ph=15)]
    pv = c.post(f"/api/v1/sites/{sid}/recon/preview", headers=h,
                files={"file": ("bad.xlsx", _wide_xlsx(rows), "application/octet-stream")}).json()
    msgs = " ".join(e["message"] for e in pv["errors"])
    assert not pv["can_confirm"] and "无法识别的类别" in msgs and "重复" in msgs and "超出物理范围" in msgs
    t = c.get("/api/v1/templates/recon-pre?site_code=RECONIMPB", headers=h).content
    wb = load_workbook(io.BytesIO(t)); wb["批次信息"]["B3"] = "修复后"
    ws = wb["指标数据"]; ws.append(_row("C1"))
    buf = io.BytesIO(); wb.save(buf)
    pv = c.post(f"/api/v1/sites/{sid}/recon/preview", headers=h, files={"file": ("t.xlsx", buf.getvalue(), "x")}).json()
    assert any("修复后" in e["message"] for e in pv["errors"])
    # 确认含错误的批次 → 409
    assert c.post(f"/api/v1/recon/batches/{pv['batch_id']}/confirm", headers=h).status_code == 409


def test_constant_column_and_single_row_warnings():
    from app.services import recon_import_service as RI

    class _S:
        site_code = "X"; original_site_code = None
    res = RI.parse_and_validate(_wide_xlsx([_row("P1")]), "one.xlsx", _S())
    assert any("仅 1 个点位" in w["message"] for w in res["warnings"])
    res = RI.parse_and_validate(_wide_xlsx([_row("P1"), _row("P2")]), "two.xlsx", _S())
    assert any("常数列" in w["message"] for w in res["warnings"])


@pytest.mark.skipif(not os.path.exists(SUBPROJECT_WB), reason="子课题测试工作簿仅在项目主机上提供")
def test_supplied_subproject_workbook_all_28_production_features_scorable():
    from app.services import recon_import_service as RI

    class _S:
        site_code = "X"; original_site_code = None
    res = RI.parse_and_validate(open(SUBPROJECT_WB, "rb").read(), os.path.basename(SUBPROJECT_WB), _S())
    assert not res["errors"] and len(res["points"]) == 81 and len(res["mapping"]) == 28
    assert len(res["cleaning_log"]) == 18
    for p in res["points"]:
        for f in M.RECON_28:
            assert M.score_indicator(f, p.get(f), "production")["status"] == "scored", (p["point"], f)


@pytest.mark.parametrize("user", ["admin", "enterprise", "regulator"])
def test_file_management_list_is_accessible(user):
    """D-v12-04: file:read 未登记导致所有角色访问文件管理 403(Windows 截图 11_files 暴露)。"""
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    tok = c.post("/api/v1/auth/login", json={"username": user, "password": "Demo@2026"})
    if tok.status_code != 200:
        pytest.skip(f"演示账号 {user} 不存在")
    h = {"Authorization": f"Bearer {tok.json()['access_token']}"}
    assert c.get("/api/v1/files", headers=h).status_code == 200
