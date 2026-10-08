"""导入流水线: 解析 -> 校验 -> 入库, 一步到位。需 DB。"""
from __future__ import annotations

import os

from sqlalchemy.orm import Session

from app.services.import_service import load_mapping, parse
from app.services.ingest_service import ingest
from app.services.threshold_resolver import build_pollutant_limits
from app.services.validation_service import validate

from app.core.config import resource_root

ROOT = resource_root()
KB_CSV = os.path.join(ROOT, "data", "knowledge_base", "统一障碍因子知识库_V1.0.csv")
ORG_SUPP_CSV = os.path.join(ROOT, "data", "knowledge_base", "有机物阈值补充_GB36600.csv")

_LIMITS_CACHE: dict | None = None


def get_pollutant_limits() -> dict:
    global _LIMITS_CACHE
    if _LIMITS_CACHE is None:
        limits = build_pollutant_limits(KB_CSV)
        # 合并有机物阈值补充(GB36600 PAH/OCP/苯并芘, brief: 三类场地全覆盖, OP 系统支持)
        if os.path.exists(ORG_SUPP_CSV):
            for fac, scopes in build_pollutant_limits(ORG_SUPP_CSV).items():
                limits.setdefault(fac, {}).update(scopes)
        _LIMITS_CACHE = limits
    return _LIMITS_CACHE


SIMULATION_LABEL = "模拟数据——仅供测试/演示"


def _file_has_simulation_label(file_path: str) -> bool:
    """扫描工作簿所有单元格(上限 20 万格)是否含模拟数据标签。CSV 按文本扫描。"""
    try:
        if file_path.lower().endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(file_path, read_only=True, data_only=True)
            n = 0
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for v in row:
                        n += 1
                        if isinstance(v, str) and SIMULATION_LABEL in v:
                            return True
                        if n > 200_000:
                            return False
            return False
        with open(file_path, encoding="utf-8", errors="ignore") as fh:
            return SIMULATION_LABEL in fh.read(5_000_000)
    except Exception:  # noqa: BLE001
        return False


def mark_site_simulated(db: Session, site_id: int) -> None:
    from app.models import ImportBatch, Measurement, Site
    site = db.get(Site, site_id)
    if site and SIMULATION_LABEL not in (site.name or ""):
        site.name = f"【{SIMULATION_LABEL}】{site.name}"
    db.query(ImportBatch).filter_by(site_id=site_id).update({"data_origin": "monte_carlo_demo"})
    db.query(Measurement).filter_by(site_id=site_id).update({"data_origin": "monte_carlo_demo"})
    db.commit()


def run_import_with_mapping(db: Session, file_path: str, mapping: dict,
                             imported_by: int | None = None,
                             scope: str = "production",
                             land_subtype: str = "其他用地",
                             on_conflict: str = "skip") -> dict:
    """直接接受 mapping 字典（无需磁盘 JSON 文件）。供 wizard 接口和 run_import 共用。

    on_conflict:  P1-3 导入幂等策略(skip/overwrite/new_version), 透传 ingest。
    """
    parsed = parse(file_path, mapping)
    report = validate(parsed, mapping, pollutant_limits=get_pollutant_limits(),
                      scope=scope, land_subtype=land_subtype)
    # 透传 mapping + source_path + on_conflict: 入库时保存 mapping_snapshot、计算
    # source_sha256/mapping_hash 做全局幂等判重(brief 4.2 +  P1-3)
    result = ingest(db, parsed, mapping=mapping, validation_report=report,
                    imported_by=imported_by, source_path=file_path,
                    on_conflict=on_conflict)
    # v1.1 (G5): 文件内含"模拟数据——仅供测试/演示"标签 → 批次/测值/场地名统一标记, 防止模拟数据混入正式结果
    if _file_has_simulation_label(file_path) and result.get("site_id"):
        mark_site_simulated(db, result["site_id"])
        result["data_origin"] = "monte_carlo_demo"
    result["validation"] = {
        "n_points": report["n_points"], "n_measurements": report["n_measurements"],
        "n_errors": report["n_errors"], "n_warnings": report["n_warnings"],
        "n_exceed": report["n_exceed"], "passed": report["passed"],
        "exceed_factors": report["summary"]["exceed_factors"],
    }
    return result


def run_import(db: Session, file_path: str, mapping_id: str,
               imported_by: int | None = None,
               scope: str = "production", land_subtype: str = "其他用地",
               on_conflict: str = "skip") -> dict:
    # v1.0.2: 预设模板已删除, mapping_id 找不到时自动走 smart_detect
    try:
        mapping = load_mapping(mapping_id)
    except FileNotFoundError:
        # 预设模板不存在 → 用 smart_detect_and_map 自动识别
        from app.services.import_service import resolve_mapping_for_file
        _, mapping, _ = resolve_mapping_for_file("auto", file_path)
    return run_import_with_mapping(db, file_path, mapping,
                                   imported_by=imported_by,
                                   scope=scope, land_subtype=land_subtype,
                                   on_conflict=on_conflict)
