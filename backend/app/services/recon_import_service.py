"""v1.2 课题二: 28 项重构可行性指标导入 → 校验预览 → 事务确认 → 计算(M-REC-2025) → 持久化 → 导出。

接受两种输入:
  1) SRS 模板 RECON-PRE-v1.2(说明 / 批次信息 / 指标数据 / 类别取值);
  2) 子课题原始宽表(首个工作表, 首行为表头, 一行一个点位, 例如
     "Special for Python testing (Gejiu, Yunnan).xlsx" 的英文表头)。
所有单元格保留原文; 数值文本中的不间断空格/制表符按确定性规则去除并记录; Excel 错误值(#VALUE! 等)
视为缺测并告警, 不当作 0; 类别值必须能映射到表2.22 的分级, 否则逐格报错。
"""
from __future__ import annotations

import csv
import hashlib
import io
import math
import os
import re
import sys
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core.config import resource_root
from app.models import (PRE_REMEDIATION, EvaluationResult, ReconImportBatch, ReconObservation, Site)
from app.services import audit_service

ROOT = resource_root()
_ML = os.path.join(ROOT, "ml", "evaluation")
if _ML not in sys.path:
    sys.path.insert(0, _ML)
import reconstruction_m2025 as M  # noqa: E402

TEMPLATE_VERSION = "RECON-PRE-v1.2"
DATA_ORIGINS = {"client_real": "真实数据(甲方/课题组提供)", "field": "现场实测",
                "monte_carlo_demo": "模拟数据——仅供测试/演示", "test_fixture": "测试夹具"}
PROVENANCE = {"verified": "来源已核实", "unverified": "来源未核实(仅作软件演示, 不作正式结论)"}
POINT_HEADERS = {"location", "点位编号", "采样点编号", "点位", "point", "point_id", "sample_id"}
_NUMTXT = re.compile(r"^[\s\u00a0\t]*[-+]?\d+(\.\d+)?([eE][-+]?\d+)?[\s\u00a0\t]*$")
_DL = re.compile(r"^[\s\u00a0]*[<＜]\s*([0-9.]+)\s*$")
# 合理范围(超出给出告警, 不阻断; 物理不可能的给错误)
_HARD = {"ph": (0, 14), "slope": (0, 90)}
_SOFT = {"total_n": (0, 10, "全氮 > 10 g/kg 超出常见范围, 请核对单位"),
         "soc": (0, 120, "有机碳 > 120 g/kg 超出常见范围, 请核对是否为有机质或单位"),
         "cec": (0, 100, "CEC > 100 cmol(+)/kg 超出常见范围"),
         "soil_depth": (0, 500, "有效土层厚度 > 500 cm, 请核对单位(cm)"),
         "groundwater_depth": (0, 100, "地下水埋深 > 100 m, 请核对单位(m)")}


# 模板列: 28 项重构指标(子课题二指标集) + 六价铬 + 生态轨道表2.22 补充指标(选填)
TEMPLATE_FEATURES = M.RECON_28 + ["cr6"] + [f for f, sp in M.FEATURES.items() if sp.get("eco_only")]


class ReconImportError(ValueError):
    pass


# ───────────────────────── 模板 ─────────────────────────
def build_template(site_code: str = "") -> bytes:
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="8B4513")
    ws = wb.active
    ws.title = "说明"
    lines = [
        f"SRS 课题二 功能重构可行性指标导入模板 {TEMPLATE_VERSION} (修复前数据)",
        "1. 本模板用于【修复前】调查评估阶段; 修复后数据请使用'课题三 修复后 SSUI 导入'。",
        "2. '指标数据'表一行一个点位; 前 28 列为课题二重构指标, 其后为六价铬与生态轨道补充指标(选填); 数值列按表头单位填写;",
        "   类别列只能填写'类别取值'表中的取值(中文或英文均可)。",
        "3. 低于检出限填写 '<检出限' (如 <0.01); 缺测留空。缺测不会被当作 0, 也不会被当作达标。",
        f"4. 评价方法: {M.METHOD_VERSION}; 来源: {M.METHOD_SOURCE}。",
        "5. 总铬(Cr)不等同六价铬; 全氮不等同水解性氮; 有机碳不等同有机质。请按实际测定项目填写。",
        "6. 数据来源与来源核实状态必须如实选择; 模拟数据结果在所有页面与报告中显示'模拟数据——仅供测试/演示'。",
    ]
    for i, t in enumerate(lines, 1):
        ws.cell(i, 1, t)
    ws.column_dimensions["A"].width = 130
    m = wb.create_sheet("批次信息")
    meta = [("场地编号", site_code, "必填, 与系统中场地编号一致"),
            ("数据阶段", "修复前", "固定为 修复前"),
            ("数据来源", DATA_ORIGINS["client_real"], " / ".join(DATA_ORIGINS.values())),
            ("来源核实", PROVENANCE["verified"], " / ".join(PROVENANCE.values())),
            ("农用地类型", "未注明", "旱地 / 水田 / 未注明 (影响 全氮/有效磷/速效钾/CEC 分档与 GB 15618 筛选值)"),
            ("生态用地类别", "第一类用地", "第一类用地 / 第二类用地 (GB 36600; 未明确时按第一类)"),
            ("备注", "", "")]
    for c, h in enumerate(["字段", "值", "说明"], 1):
        x = m.cell(1, c, h); x.font = head; x.fill = fill
    for r, (k, v, n) in enumerate(meta, 2):
        m.cell(r, 1, k); m.cell(r, 2, v); m.cell(r, 3, n)
    for col, w in zip("ABC", (14, 34, 80)):
        m.column_dimensions[col].width = w
    d = wb.create_sheet("指标数据")
    hdr = ["点位编号"] + [f"{M.FEATURES[f]['cn']}({M.FEATURES[f]['unit']})" for f in TEMPLATE_FEATURES]
    for c, h in enumerate(hdr, 1):
        x = d.cell(1, c, h); x.font = head; x.fill = fill
        d.column_dimensions[get_column_letter(c)].width = max(10, min(28, len(h) * 1.6))
    c = wb.create_sheet("类别取值")
    r = 1
    for fid in ("bulk_density", "biodiversity", "salinization", "irrigation_drainage", "texture", "carbon_factor"):
        c.cell(r, 1, M.FEATURES[fid]["cn"]).font = Font(bold=True)
        for k, al in M.CATEGORY_MAP[fid].items():
            r += 1
            c.cell(r, 2, k); c.cell(r, 3, " / ".join(al[1:]))
            c.cell(r, 4, f"生产 F={M.CAT_SCORE['production'].get(fid, {}).get(k, '—')}; "
                         f"生态 F={M.CAT_SCORE['ecology'].get(fid, {}).get(k, '无规则')}")
        r += 2
    c.cell(r, 1, "剖面构型").font = Font(bold=True)
    for k, al in M.PROFILE_TYPES.items():
        r += 1
        c.cell(r, 2, k); c.cell(r, 3, " / ".join(al[1:]))
        c.cell(r, 4, f"生产 F={M.PROFILE_SCORE['production'].get(k, '—')}; 生态 F={M.PROFILE_SCORE['ecology'].get(k, '无规则')}")
    c.cell(r + 2, 1, "多个构型可用'、'连接(须属于同一分档)。")
    for col, w in zip("ABCD", (22, 22, 50, 34)):
        c.column_dimensions[col].width = w
    for s in wb.worksheets:
        for row in s.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=s.title == "说明")
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


# ───────────────────────── 解析 ─────────────────────────
def clean_cell(v):
    """→ (value, transform, qualifier)。value: float / str / None。"""
    if v is None:
        return None, "empty", None
    if isinstance(v, bool):
        return None, "invalid_bool", None
    if isinstance(v, (int, float)):
        return (float(v), "numeric", None) if math.isfinite(float(v)) else (None, "nonfinite", None)
    t = str(v)
    if t.strip().startswith("#"):
        return None, f"excel_error:{t.strip()}", None
    if not t.replace("\u00a0", " ").strip():
        return None, "empty", None
    if _NUMTXT.match(t):
        chars = sorted({{"\u00a0": "NBSP", "\t": "TAB", " ": "SPACE"}[ch] for ch in t if ch in "\u00a0\t "})
        return float(t.replace("\u00a0", " ").strip()), ("numeric_text_stripped(" + ",".join(chars) + ")" if chars else "numeric_text"), None
    m = _DL.match(t)
    if m:
        return float(m.group(1)) / 2.0, f"below_detection_limit(DL={m.group(1)}→DL/2)", "<"
    return t.replace("\u00a0", " ").strip(), "text", None


def _read_rows(content: bytes, filename: str):
    """→ (sheet_name, header, rows(list of list), meta dict)。"""
    meta = {}
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text)))
        rows = [r for r in rows if any(str(x).strip() for x in r)]
        return "csv", rows[0], rows[1:], meta
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001
        raise ReconImportError(f"无法读取 Excel 文件: {e}")
    if "批次信息" in wb.sheetnames:
        for r in wb["批次信息"].iter_rows(min_row=2, values_only=True):
            if r and r[0]:
                meta[str(r[0]).strip()] = r[1]
    name = "指标数据" if "指标数据" in wb.sheetnames else wb.sheetnames[0]
    rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    rows = [r for r in rows if any(v not in (None, "") for v in r)]
    if not rows:
        raise ReconImportError(f"工作表 {name} 为空")
    return name, rows[0], rows[1:], meta


def _norm_meta(meta: dict, form: dict) -> dict:
    def pick(k, fk):
        v = form.get(fk)
        if v not in (None, ""):
            return v
        return meta.get(k)
    inv_o = {v: k for k, v in DATA_ORIGINS.items()}
    inv_p = {v: k for k, v in PROVENANCE.items()}
    origin = pick("数据来源", "data_origin") or "client_real"
    origin = inv_o.get(origin, origin)
    prov = pick("来源核实", "provenance_status") or "unverified"
    prov = inv_p.get(prov, prov)
    sub = pick("农用地类型", "land_subtype")
    sub = None if sub in (None, "", "未注明") else sub
    eco = pick("生态用地类别", "eco_land_class") or "第一类用地"
    return {"site_code": meta.get("场地编号"), "stage": meta.get("数据阶段") or "修复前", "data_origin": origin,
            "provenance_status": prov, "land_subtype": sub, "eco_land_class": eco}


def parse_and_validate(content: bytes, filename: str, site: Site, form: dict | None = None) -> dict:
    form = form or {}
    sheet, header, rows, meta = _read_rows(content, filename)
    m = _norm_meta(meta, form)
    errors, warnings, clean_log = [], [], []
    if m["stage"] not in ("修复前", "pre_remediation"):
        errors.append({"sheet": "批次信息", "cell": "B3", "message": f"数据阶段为 '{m['stage']}': 本入口只接受修复前数据; 修复后数据请使用课题三导入"})
    if m["site_code"] and str(m["site_code"]).strip() not in {site.site_code, getattr(site, "original_site_code", None)}:
        errors.append({"sheet": "批次信息", "cell": "B2", "message": f"模板场地编号 {m['site_code']} 与所选场地 {site.site_code} 不一致"})
    if m["data_origin"] not in DATA_ORIGINS:
        errors.append({"sheet": "批次信息", "cell": "B4", "message": f"无法识别的数据来源: {m['data_origin']}"})
    if m["provenance_status"] not in PROVENANCE:
        errors.append({"sheet": "批次信息", "cell": "B5", "message": f"无法识别的来源核实状态: {m['provenance_status']}"})
    if m["land_subtype"] not in (None, "旱地", "水田"):
        errors.append({"sheet": "批次信息", "cell": "B6", "message": f"农用地类型只能为 旱地/水田/未注明: {m['land_subtype']}"})
    if m["eco_land_class"] not in ("第一类用地", "第二类用地"):
        errors.append({"sheet": "批次信息", "cell": "B7", "message": f"生态用地类别只能为 第一类用地/第二类用地"})
    # 表头映射
    point_col, mapping, unmapped, special = None, {}, [], []
    for j, h in enumerate(header):
        if h is None or str(h).strip() == "":
            continue
        hs = str(h).strip()
        if M._nc(hs) in {M._nc(x) for x in POINT_HEADERS} and point_col is None:
            point_col = j
            continue
        fid = M.feature_id(hs)
        if fid in ("__crvi__", "__om__"):
            special.append(hs)
            warnings.append({"sheet": sheet, "cell": f"{get_column_letter(j + 1)}1",
                             "message": f"列 '{hs}' 不等同于 28 项指标中的总铬/有机碳, 未映射(不做换算)"})
            continue
        if fid is None:
            unmapped.append(hs)
            warnings.append({"sheet": sheet, "cell": f"{get_column_letter(j + 1)}1", "message": f"未识别的列 '{hs}', 已忽略"})
            continue
        if fid in mapping.values():
            errors.append({"sheet": sheet, "cell": f"{get_column_letter(j + 1)}1", "message": f"列 '{hs}' 与其他列重复映射到 {M.FEATURES[fid]['cn']}"})
            continue
        mapping[j] = fid
    if point_col is None:
        errors.append({"sheet": sheet, "cell": "A1", "message": "未找到点位编号列(location/点位编号/采样点编号)"})
    if not mapping:
        errors.append({"sheet": sheet, "cell": "1", "message": "未识别到任何重构指标列"})
    points, seen = [], {}
    for i, r in enumerate(rows, start=2):
        label = None if point_col is None or point_col >= len(r) else r[point_col]
        label = None if label is None else str(label).strip()
        if not label:
            errors.append({"sheet": sheet, "cell": f"{get_column_letter((point_col or 0) + 1)}{i}", "message": "点位编号为空"})
            continue
        if label in seen:
            errors.append({"sheet": sheet, "cell": f"{get_column_letter(point_col + 1)}{i}",
                           "message": f"点位编号 {label} 重复(首次出现在第 {seen[label]} 行)"})
            continue
        seen[label] = i
        pt = {"point": label, "_row": i, "_cells": {}}
        for j, fid in mapping.items():
            raw = r[j] if j < len(r) else None
            v, tr, q = clean_cell(raw)
            col = get_column_letter(j + 1)
            ref = f"{col}{i}"
            spec = M.FEATURES[fid]
            if tr.startswith("excel_error"):
                warnings.append({"sheet": sheet, "cell": ref, "message": f"{spec['cn']} 为 Excel 错误值 {raw}, 按缺测处理(不当作 0)"})
            if tr.startswith(("numeric_text", "below_detection")):
                clean_log.append({"sheet": sheet, "cell": ref, "point": label, "feature": fid, "original": repr(raw),
                                  "cleaned": v, "transform": tr})
            if tr == "nonfinite":
                errors.append({"sheet": sheet, "cell": ref, "message": f"{spec['cn']} 非有限数值"})
            if v is None:
                pt["_cells"][fid] = {"raw": raw, "transform": tr, "col": col}
                continue
            if spec["kind"] in ("num", "pol"):
                if isinstance(v, str):
                    errors.append({"sheet": sheet, "cell": ref, "message": f"{spec['cn']} 应为数值, 实际为 '{v}'"})
                    continue
                if spec["kind"] == "pol" and v < 0:
                    errors.append({"sheet": sheet, "cell": ref, "message": f"{spec['cn']} 浓度为负值 {v}"})
                    continue
                if fid in _HARD and not (_HARD[fid][0] <= v <= _HARD[fid][1]):
                    errors.append({"sheet": sheet, "cell": ref, "message": f"{spec['cn']} = {v} 超出物理范围 {_HARD[fid]}"})
                    continue
                if fid in _SOFT and not (_SOFT[fid][0] <= v <= _SOFT[fid][1]):
                    warnings.append({"sheet": sheet, "cell": ref, "message": f"{_SOFT[fid][2]} (值 {v})"})
            elif spec["kind"] == "cat" or (spec["kind"] == "cat_or_num" and isinstance(v, str)):
                if fid == "carbon_factor" and isinstance(v, float):
                    cat, why = M.normalize_category(fid, v)
                else:
                    cat, why = M.normalize_category(fid, v) if isinstance(v, str) else (None, f"应为类别文本, 实际为数值 {v}")
                if cat is None:
                    errors.append({"sheet": sheet, "cell": ref, "message": f"{spec['cn']}: {why}"})
                    continue
            pt[fid] = v
            pt["_cells"][fid] = {"raw": raw, "transform": tr, "col": col, "qualifier": q}
        points.append(pt)
    if len(points) == 1:
        warnings.append({"sheet": sheet, "cell": "", "message": "仅 1 个点位: 场地代表值即该点位值, 无法反映空间变异"})
    for fid in set(mapping.values()):
        vals = [p[fid] for p in points if fid in p]
        if len(vals) >= 2 and len({str(v) for v in vals}) == 1:
            warnings.append({"sheet": sheet, "cell": "", "message": f"{M.FEATURES[fid]['cn']} 全部点位取值相同({vals[0]}), 为常数列"})
        missing = len(points) - len(vals)
        if missing:
            warnings.append({"sheet": sheet, "cell": "", "message": f"{M.FEATURES[fid]['cn']} 缺测 {missing}/{len(points)} 个点位"})
    absent = [M.FEATURES[f]["cn"] for f in M.RECON_28 if f not in mapping.values()]
    return {"sheet": sheet, "meta": m, "mapping": {get_column_letter(j + 1): {"header": str(header[j]), "feature": f,
                                                                             "cn": M.FEATURES[f]["cn"]} for j, f in mapping.items()},
            "unmapped_columns": unmapped, "special_columns": special, "absent_features": absent,
            "points": points, "errors": errors, "warnings": warnings, "cleaning_log": clean_log}


# ───────────────────────── 预览 / 确认 ─────────────────────────
def _public_points(points):
    return [{k: v for k, v in p.items() if not k.startswith("_")} for p in points]


def preview(db: Session, site: Site, content: bytes, filename: str, user_id: int | None, form: dict | None = None) -> dict:
    sha = hashlib.sha256(content).hexdigest()
    res = parse_and_validate(content, filename, site, form)
    # 内容指纹: 规范化后的点位×指标值(与文件字节无关; 另存为/改表头格式不影响判重)
    import json as _json
    content_fp = hashlib.sha256(_json.dumps(_public_points(res["points"]), sort_keys=True, ensure_ascii=False,
                                            default=str).encode()).hexdigest()
    for old in db.query(ReconImportBatch).filter_by(site_id=site.id, status="confirmed").all():
        if old.source_sha256 == sha or (old.validation_report or {}).get("content_sha256") == content_fp:
            res["errors"].append({"sheet": "", "cell": "",
                                  "message": f"相同内容已在批次 #{old.id} 确认导入(文件 sha256 {sha[:12]}…), 不重复导入"})
            break
    pts = res["points"]
    preview_eval = {}
    if not res["errors"] and pts:
        vals, meta = M.site_representative(_public_points(pts))
        for scope in ("production", "ecology"):
            r = M.evaluate(vals, scope, land_subtype=res["meta"]["land_subtype"], eco_land_class=res["meta"]["eco_land_class"])
            preview_eval[scope] = {"score": r["score"], "grade": r["grade"], "path": r["path"]}
    b = ReconImportBatch(site_id=site.id, source_file=os.path.basename(filename), source_sha256=sha,
                         template_version=TEMPLATE_VERSION if res["sheet"] == "指标数据" else "supplied-wide-table",
                         stage=PRE_REMEDIATION, subproject="S2", data_origin=res["meta"]["data_origin"],
                         provenance_status=res["meta"]["provenance_status"], land_subtype=res["meta"]["land_subtype"],
                         eco_land_class=res["meta"]["eco_land_class"], status="previewed",
                         point_count=len(pts), cell_count=sum(len([k for k in p if not k.startswith("_") and k != "point"]) for p in pts),
                         error_count=len(res["errors"]), warning_count=len(res["warnings"]),
                         validation_report={"errors": res["errors"], "warnings": res["warnings"],
                                            "absent_features": res["absent_features"], "unmapped_columns": res["unmapped_columns"],
                                            "points": _public_points(pts), "cells": {p["point"]: p["_cells"] for p in pts},
                                            "sheet": res["sheet"], "content_sha256": content_fp},
                         mapping_snapshot=res["mapping"], cleaning_log=res["cleaning_log"],
                         method_version=M.METHOD_VERSION, imported_by=user_id)
    db.add(b); db.flush()
    audit_service.log(db, "recon_preview", user_id, "recon_import_batch", b.id,
                      detail={"site_id": site.id, "sha256": sha, "errors": len(res["errors"])}, commit=False)
    db.commit()
    return {"batch_id": b.id, "status": b.status, "sha256": sha, "sheet": res["sheet"], "meta": res["meta"],
            "point_count": len(pts), "mapping": res["mapping"], "absent_features": res["absent_features"],
            "unmapped_columns": res["unmapped_columns"], "errors": res["errors"], "warnings": res["warnings"][:200],
            "warning_count": len(res["warnings"]), "cleaning_log": res["cleaning_log"],
            "data_origin_label": DATA_ORIGINS.get(res["meta"]["data_origin"]),
            "provenance_label": PROVENANCE.get(res["meta"]["provenance_status"]),
            "preview_evaluation": preview_eval, "can_confirm": not res["errors"] and bool(pts)}


def confirm(db: Session, batch_id: int, user_id: int | None) -> dict:
    b = db.get(ReconImportBatch, batch_id)
    if b is None:
        raise ReconImportError("批次不存在")
    if b.status != "previewed":
        raise ReconImportError(f"批次状态为 {b.status}, 只能确认 previewed 批次")
    if b.error_count:
        raise ReconImportError(f"批次存在 {b.error_count} 个错误, 不能确认")
    rep = b.validation_report or {}
    try:
        for old in db.query(ReconImportBatch).filter_by(site_id=b.site_id, status="confirmed").all():
            old.status = "superseded"
        n = 0
        for p in rep.get("points", []):
            cells = (rep.get("cells") or {}).get(p["point"], {})
            for fid, c in cells.items():
                v = p.get(fid)
                db.add(ReconObservation(
                    batch_id=b.id, site_id=b.site_id, point_label=p["point"], feature_id=fid,
                    feature_cn=M.FEATURES[fid]["cn"],
                    value_num=float(v) if isinstance(v, (int, float)) else None,
                    value_cat=v if isinstance(v, str) else None,
                    original_text=None if c.get("raw") is None else str(c.get("raw")),
                    transform=c.get("transform"), qualifier=c.get("qualifier"), unit=M.FEATURES[fid]["unit"],
                    source_sheet=rep.get("sheet"), source_row=None, source_col=c.get("col"),
                    stage=PRE_REMEDIATION, data_origin=b.data_origin))
                n += 1
        b.status = "confirmed"
        b.confirmed_at = datetime.utcnow()
        db.flush()
        ev = evaluate_batch(db, b, user_id, commit=False)
        audit_service.log(db, "recon_confirm", user_id, "recon_import_batch", b.id,
                          detail={"observations": n, "site_id": b.site_id}, commit=False)
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"batch_id": b.id, "status": b.status, "observations": n, "evaluation": ev}


def latest_confirmed(db: Session, site_id: int) -> ReconImportBatch | None:
    return (db.query(ReconImportBatch).filter_by(site_id=site_id, status="confirmed")
            .order_by(ReconImportBatch.id.desc()).first())


def batch_points(db: Session, b: ReconImportBatch) -> list[dict]:
    pts: dict[str, dict] = {}
    for o in db.query(ReconObservation).filter_by(batch_id=b.id).all():
        p = pts.setdefault(o.point_label, {"point": o.point_label})
        p[o.feature_id] = o.value_num if o.value_num is not None else o.value_cat
    return list(pts.values())


def evaluate_batch(db: Session, b: ReconImportBatch, user_id: int | None = None, commit: bool = True) -> dict:
    pts = batch_points(db, b)
    vals, meta = M.site_representative(pts)
    out = {}
    for scope, et in (("production", "reconstruction_prod"), ("ecology", "reconstruction_eco")):
        r = M.evaluate(vals, scope, land_subtype=b.land_subtype, eco_land_class=b.eco_land_class or "第一类用地")
        per_point = [M.evaluate(p, scope, land_subtype=b.land_subtype, eco_land_class=b.eco_land_class or "第一类用地")
                     for p in pts]
        dist = {}
        for pr in per_point:
            dist[pr["grade"]] = dist.get(pr["grade"], 0) + 1
        dims = {"dimensions": r["dimensions"], "missing_indicators": r["missing_indicators"],
                "calculation_trace": r["calculation_trace"], "not_scored": r["not_scored"], "path": r["path"],
                "assumptions": r["assumptions"], "site_values": vals, "site_value_method": meta,
                "point_grade_distribution": dist, "data_origin": b.data_origin,
                "provenance_status": b.provenance_status, "batch_id": b.id}
        ev = EvaluationResult(site_id=b.site_id, eval_type=et, data_version=f"recon-batch-{b.id}", score=r["score"],
                              grade=r["grade"], dimensions=dims, weights=r["weights"],
                              limiting_factors=r["limiting_factors"], explanation=r["explanation"],
                              stage=PRE_REMEDIATION, subproject="S2", method_version=M.METHOD_VERSION[:40],
                              method_status="frozen_baseline", source_batch_id=b.id)
        db.add(ev); db.flush()
        out[scope] = {"evaluation_id": ev.id, "score": r["score"], "grade": r["grade"], "path": r["path"],
                      "limiting_factors": r["limiting_factors"], "point_grade_distribution": dist,
                      "missing_indicators": r["missing_indicators"]}
    if commit:
        db.commit()
    return out


def export_batch(db: Session, b: ReconImportBatch) -> bytes:
    pts = batch_points(db, b)
    wb = Workbook()
    ws = wb.active; ws.title = "结果摘要"
    label = DATA_ORIGINS.get(b.data_origin, b.data_origin)
    ws.append([f"SRS 课题二 功能重构可行性评价导出  批次 #{b.id}  {label}  {PROVENANCE.get(b.provenance_status)}"])
    ws.append(["方法", M.METHOD_VERSION]); ws.append(["方法来源", M.METHOD_SOURCE])
    ws.append(["源文件", b.source_file, "sha256", b.source_sha256])
    ws.append([])
    ws.append(["轨道", "综合得分", "等级", "计算路径", "限制性指标(F≤60)", "点位等级分布"])
    for scope, et in (("production", "reconstruction_prod"), ("ecology", "reconstruction_eco")):
        ev = (db.query(EvaluationResult).filter_by(source_batch_id=b.id, eval_type=et, subproject="S2")
              .order_by(EvaluationResult.id.desc()).first())
        if ev:
            d = ev.dimensions or {}
            ws.append(["生产" if scope == "production" else "生态", ev.score, ev.grade, d.get("path"),
                       "、".join(ev.limiting_factors or []), str(d.get("point_grade_distribution"))])
    for scope in ("production", "ecology"):
        s = wb.create_sheet("点位F值_" + ("生产" if scope == "production" else "生态"))
        s.append(["点位编号"] + [M.FEATURES[f]["cn"] for f in M.RECON_28] + ["综合得分", "等级"])
        for p in pts:
            r = M.evaluate(p, scope, land_subtype=b.land_subtype, eco_land_class=b.eco_land_class or "第一类用地")
            s.append([p["point"]] + [r["items"][f]["F"] if r["items"][f]["status"] == "scored" else r["items"][f]["status"]
                                     for f in M.RECON_28] + [r["score"], r["grade"]])
    s = wb.create_sheet("原始值")
    s.append(["点位编号"] + [M.FEATURES[f]["cn"] for f in M.RECON_28])
    for p in pts:
        s.append([p["point"]] + [p.get(f) for f in M.RECON_28])
    s = wb.create_sheet("清洗记录")
    s.append(["工作表", "单元格", "点位", "指标", "原文", "清洗后", "变换"])
    for c in b.cleaning_log or []:
        s.append([c["sheet"], c["cell"], c["point"], c["feature"], c["original"], c["cleaned"], c["transform"]])
    s = wb.create_sheet("方法待确认项")
    s.append(["编号", "事项", "原文", "本版采用", "依据", "状态"])
    for dfr in M.DEFERRED:
        s.append([dfr["id"], dfr["item"], dfr["source_value"], dfr["applied"], dfr["basis"], dfr["status"]])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
