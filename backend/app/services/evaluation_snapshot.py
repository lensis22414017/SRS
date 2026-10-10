"""v1.2.1(R03): 场地评价快照 — 界面摘要、Excel、PDF、DOCX、比较 JSON 的唯一数据来源。

设计原则
- 按 阶段(修复前/修复后) × 子课题(S1/S2/S3) × 批次 分开统计, 不把修复前与修复后、两条修复后轨道的
  重复导入混成一个"采样点数"。同一物理样品在两个轨道工作簿中重复出现时按样品编号去重, 并列出重复关系。
- 场地级超标结论来自法规门禁(GB 15618 / GB 36600, 与利用方向结论同一算法), 按阶段给出;
  "最新导入批次校验"只作为该批次的导入质量信息单独标注, 不再代替场地结论。
- 课题三 SSUI 以修复后批次(每轨最新已确认)为准; 修复前旧口径 SSUI 只作为"参考评价"保留并标注。
- 五阶段业务记录(调查评估…后期管护)与七项软件操作里程碑分开; 无记录/无材料一律显示"未开展"。
- 快照内容做规范化 JSON 后取 SHA-256 作为 snapshot_id; 报告记录保存快照全文, 生成后不随新数据改变。
"""
from __future__ import annotations

import hashlib
import io
import json
import re
from collections import defaultdict
from datetime import datetime, timezone

from sqlalchemy.orm import Session

from app.models import (POST_REMEDIATION, PRE_REMEDIATION, DiagnosisResult, EvaluationResult, FactorDictionary,
                        ImportBatch, Measurement, Recommendation, SamplingPoint, Site, SSUIImportBatch,
                        TechnologyLibrary, UtilizationDecision)

SNAPSHOT_SCHEMA = "srs-eval-snapshot/1.0"
SYNTHETIC_ORIGINS = {"monte_carlo_demo", "test_fixture"}
STAGE_CN = {PRE_REMEDIATION: "修复前", POST_REMEDIATION: "修复后"}
POLLUTION_TYPE_CN = {"heavy_metal": "重金属污染", "organic": "有机污染", "composite": "复合污染",
                     "HM": "重金属污染", "OP": "有机污染", "HM+OP": "复合污染"}
GATE_STATE_CN = {"pass": "未超筛选值", "conditional": "超筛选值、未超管制值(需风险评估/管控)",
                 "fail": "超管制值", "insufficient": "证据不足(缺测/条件不明)"}
DECISION_CN = {"both_supported": "生产与生态均支持", "production_supported": "支持生产利用",
               "ecology_supported": "支持生态利用", "neither_supported": "均不支持",
               "insufficient_evidence": "证据不足"}
OFFICIAL_STATUS_CN = {"available": "正式结果可用", "partial": "正式因子不足 3 个(部分结果)",
                      "insufficient_evidence": "证据不足, 无正式排名"}
_SPECIATION_UNITS = {"VI", "III", "Ⅵ", "Ⅲ", "6+", "3+", "六价", "三价"}
_POST_PREFIX = re.compile(r"^POST\d*-")


def sample_identity(point_code: str | None) -> str:
    """修复后点位编号去掉导入批次前缀(POST{批次}-), 得到物理样品编号。"""
    return _POST_PREFIX.sub("", point_code or "")


def display_unit(unit: str | None, default_unit: str | None) -> tuple[str, str | None]:
    """返回 (显示单位, 说明)。形态标记被误存为单位时改用因子默认单位并说明。"""
    u = (unit or "").strip()
    du = default_unit if default_unit and default_unit.strip() not in _SPECIATION_UNITS else "mg/kg"
    if u in _SPECIATION_UNITS:
        return du, f"原记录单位为形态标记“{u}”, 按质量浓度单位显示"
    if not u and default_unit in ("无量纲", "-", "pH"):
        return "无量纲", None
    return (u or default_unit or "—"), (None if u else "原记录未登记单位")


def _gstate(g):
    """UtilizationDecision.production_gate/ecology_gate 存的是完整门禁字典; 取其 state。"""
    return g.get("state") if isinstance(g, dict) else g


def _r(v, n=6):
    try:
        return None if v is None else round(float(v), n)
    except (TypeError, ValueError):
        return None


# ───────────────────────── 数据清单(阶段 × 批次) ─────────────────────────
def _inventory(db: Session, site_id: int) -> dict:
    rows = (db.query(Measurement.id, Measurement.stage, Measurement.import_batch_id, Measurement.factor_id,
                     Measurement.value, Measurement.unit, Measurement.qa_status, Measurement.data_origin,
                     SamplingPoint.point_code)
            .outerjoin(SamplingPoint, SamplingPoint.id == Measurement.sampling_point_id)
            .filter(Measurement.site_id == site_id).all())
    batches = {b.id: b for b in db.query(ImportBatch).filter_by(site_id=site_id).all()}
    ssui_by_sha = {}
    for sb in db.query(SSUIImportBatch).filter_by(site_id=site_id).order_by(SSUIImportBatch.id).all():
        ssui_by_sha.setdefault(sb.source_sha256, sb)
    out = {}
    for stage in (PRE_REMEDIATION, POST_REMEDIATION):
        st_rows = [r for r in rows if (r.stage or PRE_REMEDIATION) == stage]
        per_batch = defaultdict(list)
        for r in st_rows:
            per_batch[r.import_batch_id].append(r)
        blist = []
        for bid in sorted(per_batch, key=lambda x: (x is None, x or 0)):
            b = batches.get(bid)
            rs = per_batch[bid]
            sb = ssui_by_sha.get(b.source_sha256) if b is not None else None
            blist.append({
                "batch_id": bid, "subproject": (b.subproject if b else None) or ("S3" if stage == POST_REMEDIATION else "S1"),
                "kind": ("课题三修复后批次(含污染物检测)" if sb else "检测数据导入"),
                "ssui_batch_id": sb.id if sb else None, "ssui_track": sb.track if sb else None,
                "ssui_batch_status": sb.status if sb else None,
                "source_file": b.source_file if b else None,
                "source_sha256_12": (b.source_sha256 or "")[:12] if b else None,
                "data_origin": (b.data_origin if b else None) or (rs[0].data_origin if rs else None),
                "import_status": b.status if b else None,
                "created_at": b.created_at.strftime("%Y-%m-%d %H:%M") if b and b.created_at else None,
                "n_point_records": len({r.point_code for r in rs if r.point_code}),
                "n_measurement_records": len(rs),
                "n_rejected": sum(1 for r in rs if r.qa_status == "rejected"),
            })
        sample_ids = {sample_identity(r.point_code) for r in st_rows if r.point_code}
        point_records = {r.point_code for r in st_rows if r.point_code}
        uniq, conflicts = {}, []
        for r in st_rows:
            if r.value is None or not r.point_code:
                continue
            k = (sample_identity(r.point_code), r.factor_id)
            if k in uniq and abs(uniq[k] - float(r.value)) > 1e-9:
                conflicts.append({"sample": k[0], "factor_id": k[1], "values": [uniq[k], float(r.value)]})
            uniq.setdefault(k, float(r.value))
        dup_groups = []
        if stage == POST_REMEDIATION and len(blist) > 1:
            by_batch = {b["batch_id"]: {sample_identity(r.point_code) for r in per_batch[b["batch_id"]] if r.point_code}
                        for b in blist}
            ids = list(by_batch)
            for i in range(len(ids)):
                for j in range(i + 1, len(ids)):
                    shared = by_batch[ids[i]] & by_batch[ids[j]]
                    if shared:
                        dup_groups.append({"batches": [ids[i], ids[j]], "n_shared_samples": len(shared)})
        n_unique = len(sample_ids)
        out[stage] = {
            "stage": stage, "stage_cn": STAGE_CN[stage], "batches": blist,
            "n_batches": len(blist),
            "n_point_records": len(point_records), "n_unique_samples": n_unique,
            "n_measurement_records": len(st_rows), "n_unique_measurements": len(uniq),
            "duplicate_groups": dup_groups, "value_conflicts": conflicts[:20], "n_value_conflicts": len(conflicts),
            "note": ((f"{len(blist)} 个修复后批次共含 {len(point_records)} 条点位记录, 按样品编号去重为 {n_unique} 个独立样品; "
                      "重复导入不代表额外采样。") if dup_groups else None),
        }
    # 课题二指标批次(修复前, 与 KOS 检测数据批次分离)
    try:
        from app.services import recon_import_service as RI
        rb = RI.latest_confirmed(db, site_id)
    except Exception:  # pragma: no cover
        rb = None
    out["s2_indicator_batch"] = ({"batch_id": rb.id, "source_file": rb.source_file, "n_points": rb.point_count,
                                  "data_origin": rb.data_origin, "provenance_status": rb.provenance_status,
                                  "stage": rb.stage} if rb else None)
    return out


# ───────────────────────── 法规门禁(按阶段) ─────────────────────────
def _gate_summary(g: dict, kind: str) -> dict:
    rows = []
    for f, d in sorted((g.get("factors") or {}).items()):
        if kind == "production":
            w = d.get("worst") or {}
            rows.append({"factor": f, "n_obs": d.get("n_obs"), "n_exceed_screening": d.get("n_exceed_screening"),
                         "n_exceed_control": d.get("n_exceed_control"), "n_undetermined": d.get("n_undetermined"),
                         "worst_point": w.get("point"), "worst_value": _r(w.get("value")),
                         "worst_screening": _r(w.get("screening")), "worst_ratio": _r(w.get("ratio_to_screening"), 3)})
        else:
            m = d.get("max") or {}
            rows.append({"factor": f, "n_obs": d.get("n_obs"), "n_exceed_screening": d.get("n_exceed_screening"),
                         "n_exceed_control": d.get("n_exceed_control"), "screening": _r(d.get("screening")),
                         "control": _r(d.get("control")), "worst_point": m.get("point"),
                         "worst_value": _r(m.get("value")), "worst_ratio": _r(m.get("ratio_to_screening"), 3),
                         "confidence": d.get("confidence")})
    exceed_any = sorted(set(g.get("exceed_control", [])) | set(g.get("exceed_screening_only", [])))
    return {"standard": g.get("standard"), "state": g.get("state"), "state_cn": GATE_STATE_CN.get(g.get("state"), g.get("state")),
            "farmland_type": g.get("farmland_type"), "land_class": g.get("land_class"),
            "exceed_control": g.get("exceed_control", []), "exceed_screening_only": g.get("exceed_screening_only", []),
            "exceed_factors": exceed_any, "n_exceed_factors": len(exceed_any),
            "missing_required": g.get("missing_required", []), "undetermined": g.get("undetermined", []),
            "threshold_missing": g.get("threshold_missing", []), "factors": rows,
            "notes": list(g.get("notes", []) or []) + list(g.get("assumptions", []) or [])}


def _gates(db: Session, site_id: int) -> dict:
    from app.services import utilization_service as US
    U = US.U
    std = U.load_standards()
    out = {}
    for stage in (PRE_REMEDIATION, POST_REMEDIATION):
        dec = (db.query(UtilizationDecision).filter_by(site_id=site_id, stage=stage)
               .order_by(UtilizationDecision.id.desc()).first())
        ev = (dec.evidence or {}) if dec else {}
        farmland = ((ev.get("production") or {}).get("gate") or {}).get("farmland_type")
        farmland = None if farmland in (None, "未指定") else farmland
        land_class = ((ev.get("ecology") or {}).get("gate") or {}).get("land_class") if dec else None
        if stage == PRE_REMEDIATION:
            points, notes, origin = US._pre_points(db, site_id)
            src = {"kind": "修复前检测数据" + ("+课题二指标批次污染物" if any(str(p.get("point", "")).startswith("S2:") for p in points) else ""),
                   "batches": None}
        else:
            bs = [b for b in (db.query(SSUIImportBatch).filter_by(site_id=site_id, track=t, status="confirmed")
                              .order_by(SSUIImportBatch.id.desc()).first() for t in ("production", "ecology")) if b]
            latest = max(bs, key=lambda b: b.id, default=None)
            points = (latest.validation_report or {}).get("points", []) if latest else []
            notes = [] if latest else ["尚无已确认的修复后(课题三)批次"]
            fp = {b.id: hashlib.sha256(json.dumps((b.validation_report or {}).get("points", []), sort_keys=True,
                                                  ensure_ascii=False, default=str).encode()).hexdigest() for b in bs}
            same = len(set(fp.values())) == 1 and len(fp) > 1
            src = {"kind": "课题三修复后批次污染物检测", "points_batch": latest.id if latest else None,
                   "batches": [b.id for b in bs], "tracks_share_same_samples": same}
            if same:
                notes.append(f"生产/生态两个修复后批次 #{bs[0].id}/#{bs[1].id} 为同一组样品(点位与浓度完全一致), 按一组样品计")
        if not points:
            out[stage] = {"available": False, "n_points": 0, "source": src, "notes": notes,
                          "decision_id": dec.id if dec else None}
            continue
        pg = U.production_gate(points, std, farmland)
        eg = U.ecology_gate(points, std, land_class)
        out[stage] = {"available": True, "n_points": len(points), "source": src, "notes": notes,
                      "production": _gate_summary(pg, "production"), "ecology": _gate_summary(eg, "ecology"),
                      "standards_sha256": std["sha256"][:16], "decision_id": dec.id if dec else None,
                      "consistent_with_decision": (None if not dec else
                                                   (_gstate(dec.production_gate) == pg["state"]
                                                    and _gstate(dec.ecology_gate) == eg["state"]))}
    return out


# ───────────────────────── 课题一 KOS ─────────────────────────
def _kos(db: Session, site_id: int) -> dict | None:
    d = (db.query(DiagnosisResult).filter_by(site_id=site_id, diagnosis_method="kos")
         .order_by(DiagnosisResult.id.desc()).first())
    if d is None or not d.result_payload:
        return None
    rp = d.result_payload
    pts = {p.id: p.point_code for p in db.query(SamplingPoint).filter_by(site_id=site_id).all()}

    def row(k):
        dp = k.get("decision_point_id")
        return {"rank": k.get("rank"), "factor": k.get("factor"), "KOS": _r(k.get("KOS"), 4),
                "value": _r(k.get("value")), "threshold_value": _r(k.get("threshold_value")),
                "threshold_unit": k.get("threshold_unit"), "threshold_type": k.get("threshold_type"),
                "threshold_standard": k.get("threshold_standard"),
                "threshold_resolution_status": k.get("threshold_resolution_status"),
                "evidence": k.get("evidence") or (k.get("components") or {}).get("E"),
                "exceedance_ratio": _r(k.get("exceedance_ratio"), 4),
                "decision_point_id": dp, "decision_point_code": pts.get(dp) if dp is not None else None,
                "threshold_condition": " ".join(x for x in (
                    ((k.get("evidence_chain") or {}).get("land_use_type") or ""),
                    ((k.get("evidence_chain") or {}).get("pH_condition") or "")) if x) or None}
    off = [row(k) for k in rp.get("key_obstacles", []) or []]
    exp = [row(k) for k in rp.get("exploratory_obstacles", []) or []]
    status = rp.get("official_ranking_status") or ("available" if off else "insufficient_evidence")
    return {"diagnosis_id": d.id, "stage": PRE_REMEDIATION, "track": d.track, "subset": d.subset,
            "created_at": d.created_at.strftime("%Y-%m-%d %H:%M") if d.created_at else None,
            "model_version": d.model_version, "threshold_version": rp.get("threshold_version"),
            "official_ranking_status": status, "official_ranking_status_cn": OFFICIAL_STATUS_CN.get(status, status),
            "official": off, "exploratory": exp, "n_official": len(off), "n_exploratory": len(exp),
            "unit_unresolved": rp.get("unit_unresolved", []),
            "legacy_payload": "official_ranking_status" not in rp,
            "farmland_type": (rp.get("threshold_condition") or {}).get("farmland_type"),
            "rule": rp.get("official_ranking_rule")}


# ───────────────────────── 课题二 / 课题三 ─────────────────────────
def _eval_row(e: EvaluationResult) -> dict:
    dims = e.dimensions or {}
    return {"evaluation_id": e.id, "eval_type": e.eval_type, "stage": e.stage, "subproject": e.subproject,
            "score": _r(e.score), "grade": e.grade, "method_version": e.method_version,
            "method_status": e.method_status, "source_batch_id": e.source_batch_id, "data_version": e.data_version,
            "status": dims.get("status"), "feasible": dims.get("feasible"),
            "exceeds_unit_range": dims.get("exceeds_unit_range"),
            "created_at": e.created_at.strftime("%Y-%m-%d %H:%M") if e.created_at else None}


def _reconstruction(db: Session, site_id: int) -> dict:
    out = {}
    for et, track in (("reconstruction_prod", "production"), ("reconstruction_eco", "ecology")):
        evs = (db.query(EvaluationResult).filter_by(site_id=site_id, eval_type=et)
               .order_by(EvaluationResult.id.desc()).all())
        frozen = [e for e in evs if e.source_batch_id is not None]
        head = frozen[0] if frozen else (evs[0] if evs else None)
        out[track] = {"headline": _eval_row(head) if head else None,
                      "headline_basis": ("课题二指标批次(冻结方法)" if frozen else ("修复前检测数据(旧通用评价)" if head else None)),
                      "all_records": [_eval_row(e) for e in evs]}
    return out


def _ssui(db: Session, site_id: int) -> dict:
    pre = (db.query(EvaluationResult).filter_by(site_id=site_id, eval_type="ssui")
           .order_by(EvaluationResult.id.desc()).first())
    post = {}
    for track in ("production", "ecology"):
        b = (db.query(SSUIImportBatch).filter_by(site_id=site_id, track=track, status="confirmed")
             .order_by(SSUIImportBatch.id.desc()).first())
        ev = None
        if b is not None:
            ev = (db.query(EvaluationResult).filter_by(site_id=site_id, eval_type=f"ssui_post_{track}",
                                                       source_batch_id=b.id)
                  .order_by(EvaluationResult.id.desc()).first())
        superseded = [x.id for x in db.query(SSUIImportBatch).filter_by(site_id=site_id, track=track, status="superseded").all()]
        post[track] = {"batch_id": b.id if b else None, "source_file": b.source_file if b else None,
                       "data_origin": b.data_origin if b else None,
                       "years_since_remediation": b.years_since_remediation if b else None,
                       "multiplier_m": b.multiplier_m if b else None,
                       "evaluation": _eval_row(ev) if ev else None, "superseded_batches": superseded}
    return {"post": post,
            "headline": {t: (post[t]["evaluation"] or {}).get("score") for t in post},
            "headline_grade": {t: (post[t]["evaluation"] or {}).get("grade") for t in post},
            "pre_reference": ({**_eval_row(pre), "label": "修复前旧口径综合评价(参考, 非课题三修复后 SSUI)"}
                              if pre else None),
            "method_status": "provisional",
            "note": "SSUI 为课题三修复后评价; 得分录入口径, 方法与 0.6 支持阈值待课题组签认(provisional)。"}


def _utilization(db: Session, site_id: int) -> dict:
    out = {}
    for stage in (PRE_REMEDIATION, POST_REMEDIATION):
        d = (db.query(UtilizationDecision).filter_by(site_id=site_id, stage=stage)
             .order_by(UtilizationDecision.id.desc()).first())
        out[stage] = None if d is None else {
            "decision_id": d.id, "state": d.decision_state, "state_cn": DECISION_CN.get(d.decision_state, d.decision_state),
            "production_gate": _gstate(d.production_gate), "ecology_gate": _gstate(d.ecology_gate),
            "production_score": _r(d.production_score), "ecology_score": _r(d.ecology_score),
            "conclusion": d.conclusion_text, "method_version": d.method_version, "method_status": d.method_status,
            "missing_evidence": d.missing_evidence or [], "created_at": d.created_at.strftime("%Y-%m-%d %H:%M") if d.created_at else None}
    return out


def _recommendations(db: Session, site_id: int) -> dict:
    recs = db.query(Recommendation).filter_by(site_id=site_id).order_by(Recommendation.rank).all()
    tech = {t.id: t.tech_name for t in db.query(TechnologyLibrary).all()}
    return {"n": len(recs), "status": "generated" if recs else "not_generated",
            "items": [{"rank": r.rank, "technology": tech.get(r.technology_id, "—"), "match_score": _r(r.match_score, 3)}
                      for r in recs[:10]]}


def _latest_batch_validation(db: Session, site_id: int) -> dict | None:
    b = db.query(ImportBatch).filter_by(site_id=site_id).order_by(ImportBatch.id.desc()).first()
    if b is None:
        return None
    vr = b.validation_report or {}
    return {"batch_id": b.id, "stage": b.stage, "stage_cn": STAGE_CN.get(b.stage, b.stage), "source_file": b.source_file,
            "passed": vr.get("passed", True), "n_errors": vr.get("n_errors", 0), "n_warnings": vr.get("n_warnings", 0),
            "n_exceed": vr.get("n_exceed", 0),
            "note": "仅为最新导入批次的导入校验结果, 不是全场地超标结论(全场地结论见法规门禁按阶段统计)"}


def _factor_summary(db: Session, site_id: int) -> dict:
    rows = (db.query(Measurement.stage, FactorDictionary.factor_code, FactorDictionary.level1_category,
                     FactorDictionary.default_unit, Measurement.value, Measurement.unit, SamplingPoint.point_code)
            .join(FactorDictionary, Measurement.factor_id == FactorDictionary.id)
            .outerjoin(SamplingPoint, SamplingPoint.id == Measurement.sampling_point_id)
            .filter(Measurement.site_id == site_id, Measurement.value.isnot(None)).all())
    agg: dict = {}
    for stage, code, cat, dunit, val, unit, pc in rows:
        st = stage or PRE_REMEDIATION
        du, note = ("无量纲", None) if str(code).lower() in ("ph", "ph值") else display_unit(unit, dunit)
        d = agg.setdefault((st, code, du), {"stage": st, "factor": code, "category": cat, "unit": du,
                                            "unit_note": note, "vals": {}})
        d["vals"].setdefault(sample_identity(pc) or f"m{len(d['vals'])}", float(val))
    out = {PRE_REMEDIATION: [], POST_REMEDIATION: []}
    for d in agg.values():
        vs = list(d["vals"].values())
        out[d["stage"]].append({"factor": d["factor"], "category": d["category"], "unit": d["unit"],
                                "unit_note": d["unit_note"], "n_samples": len(vs), "min": _r(min(vs), 4),
                                "mean": _r(sum(vs) / len(vs), 4), "max": _r(max(vs), 4)})
    for k in out:
        out[k].sort(key=lambda x: x["factor"])
    return out


# ───────────────────────── 主入口 ─────────────────────────
def _canonical(obj) -> str:
    return json.dumps(obj, ensure_ascii=False, sort_keys=True, default=str, separators=(",", ":"))


def build(db: Session, site_id: int) -> dict:
    from app.services import workflow_service as WS
    site = db.get(Site, site_id)
    if site is None:
        raise ValueError(f"场地不存在: {site_id}")
    origins = {o for (o,) in db.query(Measurement.data_origin).filter_by(site_id=site_id).distinct().all()}
    origins |= {o for (o,) in db.query(SSUIImportBatch.data_origin).filter_by(site_id=site_id).distinct().all()}
    synthetic = bool(origins & SYNTHETIC_ORIGINS)
    biz = WS.business_stage_status(db, site_id)
    # “全流程报告”里程碑: 快照是报告的数据依据, 以快照出具报告即完成该项; 不记次数,
    # 否则先出 PDF 再出 DOCX 会得到不同快照(报告数变化)。
    ms = [({**m, "count": None, "done": True, "note": "以本快照出具报告即完成"} if m["key"] == "report" else m)
          for m in WS.software_milestones(db, site_id)]
    body = {
        "schema": SNAPSHOT_SCHEMA,
        "site": {"site_id": site.id, "site_code": site.site_code, "name": site.name,
                 "pollution_type": site.pollution_type,
                 "pollution_type_cn": POLLUTION_TYPE_CN.get(site.pollution_type or "", site.pollution_type or "未填写"),
                 "land_use_type": site.land_use_type, "province": site.province, "city": site.city},
        "data_origin": {"origins": sorted(o or "field" for o in origins), "is_synthetic": synthetic,
                        "label": ("模拟数据——仅供测试/演示，不得用于正式报告" if synthetic else None),
                        "point_label": "模拟采样点" if synthetic else "采样点",
                        "factor_label": "模拟检测因子" if synthetic else "检测因子"},
        "inventory": _inventory(db, site_id),
        "gates": _gates(db, site_id),
        "latest_batch_validation": _latest_batch_validation(db, site_id),
        "factor_summary": _factor_summary(db, site_id),
        "kos": _kos(db, site_id),
        "reconstruction": _reconstruction(db, site_id),
        "ssui": _ssui(db, site_id),
        "utilization": _utilization(db, site_id),
        "recommendations": _recommendations(db, site_id),
        "workflow": {"business_stages": biz["stages"], "business_completed": biz["n_completed"],
                     "business_with_documents": biz["n_with_documents"], "business_total": len(biz["stages"]),
                     "software_milestones": ms, "software_completed": sum(1 for m in ms if m["done"]),
                     "software_total": len(ms),
                     "note": "七项软件操作里程碑记录系统内操作; 五阶段业务记录需上传调查/审批/施工/效果/管护材料, 二者不能互相替代。"},
    }
    sid = hashlib.sha256(_canonical(body).encode("utf-8")).hexdigest()
    body["snapshot_id"] = sid[:16]
    body["snapshot_sha256"] = sid
    body["generated_at"] = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    return body


def verify(snap: dict) -> bool:
    """重算 snapshot_sha256, 确认快照未被改动。"""
    body = {k: v for k, v in snap.items() if k not in ("snapshot_id", "snapshot_sha256", "generated_at")}
    return hashlib.sha256(_canonical(body).encode("utf-8")).hexdigest() == snap.get("snapshot_sha256")


def headline(snap: dict) -> dict:
    """各渠道(界面摘要/Excel/PDF/DOCX/比较 JSON)共用的首页关键数, 跨渠道一致性测试以此为准。"""
    inv, g, k, s, u = snap["inventory"], snap["gates"], snap.get("kos") or {}, snap["ssui"], snap["utilization"]
    pre, post = inv[PRE_REMEDIATION], inv[POST_REMEDIATION]

    def gate(stage, track):
        x = g.get(stage) or {}
        return (x.get(track) or {}).get("state") if x.get("available") else None

    def gexc(stage, track):
        x = g.get(stage) or {}
        return (x.get(track) or {}).get("exceed_factors", []) if x.get("available") else []
    rc = snap["reconstruction"]
    return {
        "snapshot_id": snap["snapshot_id"], "is_synthetic": snap["data_origin"]["is_synthetic"],
        "pre_n_samples": pre["n_unique_samples"], "pre_n_measurements": pre["n_measurement_records"],
        "post_n_samples": post["n_unique_samples"], "post_n_point_records": post["n_point_records"],
        "post_n_measurements": post["n_measurement_records"], "post_n_batches": post["n_batches"],
        "pre_gate_production": gate(PRE_REMEDIATION, "production"), "pre_gate_ecology": gate(PRE_REMEDIATION, "ecology"),
        "pre_exceed_production": gexc(PRE_REMEDIATION, "production"), "pre_exceed_ecology": gexc(PRE_REMEDIATION, "ecology"),
        "post_gate_production": gate(POST_REMEDIATION, "production"), "post_gate_ecology": gate(POST_REMEDIATION, "ecology"),
        "post_exceed_production": gexc(POST_REMEDIATION, "production"), "post_exceed_ecology": gexc(POST_REMEDIATION, "ecology"),
        "kos_official_status": k.get("official_ranking_status"),
        "kos_official": [x["factor"] for x in k.get("official", [])],
        "kos_exploratory": [x["factor"] for x in k.get("exploratory", [])],
        "recon_production": ((rc["production"]["headline"] or {}).get("score")),
        "recon_ecology": ((rc["ecology"]["headline"] or {}).get("score")),
        "ssui_post_production": s["headline"].get("production"), "ssui_post_ecology": s["headline"].get("ecology"),
        "decision_pre": (u.get(PRE_REMEDIATION) or {}).get("state"),
        "decision_post": (u.get(POST_REMEDIATION) or {}).get("state"),
        "recommendations": snap["recommendations"]["n"],
        "business_stages_completed": snap["workflow"]["business_completed"],
        "business_stages_total": snap["workflow"]["business_total"],
        "software_milestones_completed": snap["workflow"]["software_completed"],
        "software_milestones_total": snap["workflow"]["software_total"],
    }


def to_xlsx(snap: dict) -> bytes:
    """快照 Excel: 首页关键数 + 分阶段批次 + 门禁 + KOS 正式/探索性 + SSUI + 五阶段。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    hdr_fill = PatternFill("solid", fgColor="1F4E79")

    def sheet(title, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill
        for r in rows:
            ws.append([("、".join(map(str, v)) if isinstance(v, list) else v) for v in r])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(48, max(10, max(len(str(c.value or "")) for c in col) + 2))
        return ws
    h = headline(snap)
    rows = [["snapshot_id", snap["snapshot_id"]], ["snapshot_sha256", snap["snapshot_sha256"]],
            ["生成时间", snap["generated_at"]], ["场地", f"{snap['site']['name']} ({snap['site']['site_code']})"],
            ["数据来源标签", snap["data_origin"]["label"] or "非模拟数据"]] + [[k, v] for k, v in h.items()]
    sheet("首页关键数", ["字段", "值"], rows, first=True)
    inv = snap["inventory"]
    sheet("阶段与批次", ["阶段", "批次", "子课题", "类型", "课题三批次", "轨道", "来源文件", "点位记录", "检测记录", "数据来源"],
          [[inv[s]["stage_cn"], b["batch_id"], b["subproject"], b["kind"], b["ssui_batch_id"], b["ssui_track"],
            b["source_file"], b["n_point_records"], b["n_measurement_records"], b["data_origin"]]
           for s in (PRE_REMEDIATION, POST_REMEDIATION) for b in inv[s]["batches"]]
          + [[inv[s]["stage_cn"], "合计", "", "去重后独立样品", "", "", inv[s]["note"] or "", inv[s]["n_unique_samples"],
              inv[s]["n_measurement_records"], ""] for s in (PRE_REMEDIATION, POST_REMEDIATION)])
    grows = []
    for s in (PRE_REMEDIATION, POST_REMEDIATION):
        g = snap["gates"].get(s) or {}
        if not g.get("available"):
            grows.append([STAGE_CN[s], "—", "无数据", "", "", "", "", "", ""]); continue
        for t in ("production", "ecology"):
            gg = g[t]
            for f in gg["factors"]:
                grows.append([STAGE_CN[s], gg["standard"], gg["state_cn"], f["factor"], f["n_obs"],
                              f["n_exceed_screening"], f["n_exceed_control"], f["worst_value"], f["worst_ratio"]])
    sheet("法规门禁按阶段", ["阶段", "标准", "门禁结论", "因子", "样品数", "超筛选值点数", "超管制值点数", "最大值", "最大超筛选倍数"], grows)
    k = snap.get("kos") or {}
    sheet("KOS正式与探索性", ["层级", "排名", "因子", "KOS", "判定值", "阈值", "单位", "方向", "来源标准", "解析状态", "证据", "判定点"],
          [[lay, x["rank"], x["factor"], x["KOS"], x["value"], x["threshold_value"], x["threshold_unit"],
            x["threshold_type"], x["threshold_standard"], x["threshold_resolution_status"], x["evidence"],
            x["decision_point_code"]] for lay, key in (("正式", "official"), ("探索性(待复核)", "exploratory"))
           for x in k.get(key, [])])
    srows = []
    for t, cn in (("production", "生产"), ("ecology", "生态")):
        p = snap["ssui"]["post"][t]; e = p["evaluation"] or {}
        srows.append(["修复后", cn, p["batch_id"], p["source_file"], e.get("score"), e.get("grade"), e.get("method_status")])
    pr = snap["ssui"]["pre_reference"]
    if pr:
        srows.append(["修复前(参考, 旧口径)", "—", None, None, pr.get("score"), pr.get("grade"), pr.get("method_status")])
    sheet("SSUI", ["阶段", "轨道", "批次", "来源文件", "SSUI", "等级", "方法状态"], srows)
    wf = snap["workflow"]
    sheet("五阶段与软件里程碑", ["类别", "名称", "状态", "附件数"],
          [["五阶段业务记录", b["name"], b["status_cn"], b["n_attachments"]] for b in wf["business_stages"]]
          + [["软件操作里程碑", m["name"], "已执行" if m["done"] else "未执行", m.get("count")] for m in wf["software_milestones"]])
    bio = io.BytesIO(); wb.save(bio)
    return bio.getvalue()
