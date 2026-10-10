"""v1.2.1(R03/R04): 报告文档模型 — 同一份内容块列表渲染为 HTML / PDF / DOCX。

此前 HTML 模板、xhtml2pdf 与 python-docx 各自拼装内容, 首页与正文、PDF 与 DOCX 可能不一致,
PDF 在缺少 HTML 渲染库时退化为纯文本。现改为:
  collect() → 评价快照(evaluation_snapshot) → build_blocks() → render_html / render_pdf / render_docx
PDF 使用 ReportLab platypus 直接排版(表格、图件、页眉页脚、目录式章节编号), 优先嵌入 TrueType 中文字体。
"""
from __future__ import annotations

import base64
import html as _html
import io
import os
from typing import Any

PRE, POST = "pre_remediation", "post_remediation"
_CN_NUM = "零一二三四五六七八九十"


def cn_num(n: int) -> str:
    if n <= 10:
        return _CN_NUM[n]
    if n < 20:
        return "十" + (_CN_NUM[n - 10] if n > 10 else "")
    t, u = divmod(n, 10)
    return _CN_NUM[t] + "十" + (_CN_NUM[u] if u else "")


def _fmt(v, nd=4) -> str:
    if v is None or v == "":
        return "—"
    if isinstance(v, bool):
        return "是" if v else "否"
    if isinstance(v, float):
        s = f"{v:.{nd}f}".rstrip("0").rstrip(".")
        return s if s not in ("", "-0") else "0"
    if isinstance(v, (list, tuple)):
        return "、".join(_fmt(x) for x in v) if v else "无"
    return str(v)


def _png_from_data_url(u: str | None) -> bytes | None:
    if not u or not isinstance(u, str) or "base64," not in u:
        return None
    try:
        return base64.b64decode(u.split("base64,", 1)[1])
    except Exception:
        return None


# ───────────────────────── 图件 ─────────────────────────
def _mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager as fm
    cands = ["Microsoft YaHei", "SimHei", "PingFang SC", "Heiti SC", "STHeiti", "Hiragino Sans GB",
             "Noto Sans CJK SC", "WenQuanYi Zen Hei", "Arial Unicode MS"]
    avail = {f.name for f in fm.fontManager.ttflist}
    plt.rcParams["font.sans-serif"] = [c for c in cands if c in avail] + ["DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    return plt


def _fig_png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight", facecolor="white")
    import matplotlib.pyplot as plt
    plt.close(fig)
    return buf.getvalue()


def chart_gate_ratios(gate: dict, title: str) -> bytes | None:
    rows = [r for r in (gate or {}).get("factors", []) if r.get("worst_ratio") is not None]
    if not rows:
        return None
    try:
        plt = _mpl()
    except Exception:
        return None
    rows = sorted(rows, key=lambda r: -(r["worst_ratio"] or 0))[:16]
    names = [r["factor"] for r in rows][::-1]
    vals = [r["worst_ratio"] for r in rows][::-1]
    colors = ["#C0392B" if (r.get("n_exceed_control") or 0) > 0 else "#E67E22" if v > 1 else "#5B8DB8"
              for r, v in zip(rows[::-1], vals)]
    fig, ax = plt.subplots(figsize=(6.6, 0.32 * len(names) + 1.1))
    ax.barh(names, vals, color=colors)
    ax.axvline(1.0, color="#333", lw=0.8, ls="--")
    ax.text(1.0, len(names) - 0.4, " 筛选值", fontsize=7, va="bottom")
    for i, v in enumerate(vals):
        ax.text(v, i, f" {v:.2f}", va="center", fontsize=7)
    ax.set_xlabel("最不利点位浓度 / 筛选值", fontsize=8)
    ax.set_title(title, fontsize=9)
    ax.tick_params(labelsize=7.5)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    return _fig_png(fig)


def chart_kos(kos: dict) -> bytes | None:
    off, exp = (kos or {}).get("official", []), (kos or {}).get("exploratory", [])
    if not off and not exp:
        return None
    try:
        plt = _mpl()
    except Exception:
        return None
    items = [(f"{x['factor']}", x.get("KOS") or 0, "official") for x in off] + \
            [(f"{x['factor']}(探索)", x.get("KOS") or 0, "exploratory") for x in exp]
    items = items[:14][::-1]
    fig, ax = plt.subplots(figsize=(6.6, 0.32 * len(items) + 1.1))
    ax.barh([i[0] for i in items], [i[1] for i in items],
            color=["#1F4E79" if i[2] == "official" else "#F0AD4E" for i in items],
            hatch=None)
    for k, i in enumerate(items):
        ax.text(i[1], k, f" {i[1]:.3f}", va="center", fontsize=7)
    ax.set_xlim(0, 1)
    ax.set_xlabel("KOS 评分(深蓝=正式, 橙=探索性/待复核)", fontsize=8)
    ax.set_title("课题一 关键障碍因子 KOS 评分", fontsize=9)
    ax.tick_params(labelsize=7.5)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    return _fig_png(fig)


def chart_ssui(ssui: dict) -> bytes | None:
    post = (ssui or {}).get("post", {})
    vals = [(cn, ((post.get(t) or {}).get("evaluation") or {}).get("score")) for t, cn in (("production", "生产"), ("ecology", "生态"))]
    vals = [(a, b) for a, b in vals if b is not None]
    if not vals:
        return None
    try:
        plt = _mpl()
    except Exception:
        return None
    fig, ax = plt.subplots(figsize=(4.2, 2.4))
    ax.bar([v[0] for v in vals], [v[1] for v in vals], color=["#2E7D32", "#00897B"][:len(vals)], width=0.5)
    ax.axhline(0.6, color="#C0392B", lw=0.8, ls="--")
    ax.text(-0.45, 0.62, "0.6 暂定支持阈值", fontsize=7, ha="left", color="#C0392B")
    ax.axhline(1.0, color="#777", lw=0.6, ls=":")
    for i, v in enumerate(vals):
        ax.text(i, v[1], f"{v[1]:.3f}", ha="center", va="bottom", fontsize=8)
    ax.set_ylim(0, max(1.15, max(v[1] for v in vals) + 0.1))
    ax.set_title("课题三 修复后 SSUI(每轨最新已确认批次)", fontsize=9)
    ax.tick_params(labelsize=8)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    return _fig_png(fig)


# ───────────────────────── 内容块 ─────────────────────────
class _B:
    def __init__(self):
        self.blocks: list[dict] = []
        self.ch = 0
        self.sub = 0

    def h1(self, text):
        self.ch += 1
        self.sub = 0
        self.blocks.append({"t": "h1", "text": f"{cn_num(self.ch)}、{text}"})

    def h2(self, text):
        self.sub += 1
        self.blocks.append({"t": "h2", "text": f"{self.ch}.{self.sub} {text}"})

    def p(self, text, style="body"):
        if text:
            self.blocks.append({"t": "p", "text": str(text), "style": style})

    def kv(self, rows):
        self.blocks.append({"t": "kv", "rows": [(str(k), _fmt(v)) for k, v in rows]})

    def table(self, header, rows, widths=None, caption=None, empty="无"):
        if not rows:
            self.p(f"{caption + ': ' if caption else ''}{empty}", "muted")
            return
        self.blocks.append({"t": "table", "header": header, "rows": [[_fmt(c) for c in r] for r in rows],
                            "widths": widths, "caption": caption})

    def img(self, png: bytes | None, caption: str, width=0.92):
        if png:
            self.blocks.append({"t": "img", "png": png, "caption": caption, "width": width})


def build_blocks(ctx: dict) -> list[dict]:
    snap = ctx["snapshot"]
    hl = ctx.get("headline") or {}
    site, rep = ctx["site"], ctx["report"]
    org = snap["data_origin"]
    inv, gates, kos, ssui = snap["inventory"], snap["gates"], snap.get("kos"), snap["ssui"]
    b = _B()
    b.blocks.append({"t": "cover", "title": "污染场地全流程监管追溯报告",
                     "subtitle": "污染场地土壤生态-生产功能重构监管系统",
                     "rows": [("场地名称", site.get("name")), ("场地编号", site.get("site_code")),
                              ("报告版本", rep.get("version")), ("评价快照", f"{snap['snapshot_id']}（SHA-256 前 16 位）"),
                              ("生成时间", rep.get("generated_at")), ("数据来源", org.get("label") or "非模拟数据(以导入批次登记为准)"),
                              ("密级", "内部")],
                     "banner": org.get("label")})

    # 首页摘要 — 全部取自快照
    def gate_txt(stage, track):
        g = gates.get(stage) or {}
        if not g.get("available"):
            return "无该阶段数据"
        x = g[track]
        exc = "、".join(x["exceed_factors"]) or "无"
        return f"{x['state_cn']}；超筛选值因子: {exc}"
    pre_i, post_i = inv[PRE], inv[POST]
    summ = [
        ("修复前数据", f"{pre_i['n_unique_samples']} 个{org['point_label']} / {pre_i['n_measurement_records']} 条检测记录（{pre_i['n_batches']} 个批次）"),
        ("修复后数据", (f"{post_i['n_unique_samples']} 个独立样品 / {post_i['n_unique_measurements']} 条去重检测记录"
                    f"（{post_i['n_batches']} 个批次, {post_i['n_point_records']} 条点位记录）") if post_i["n_batches"] else "尚无修复后数据"),
        ("修复前 生产门禁 GB 15618", gate_txt(PRE, "production")),
        ("修复前 生态门禁 GB 36600", gate_txt(PRE, "ecology")),
        ("修复后 生产门禁 GB 15618", gate_txt(POST, "production")),
        ("修复后 生态门禁 GB 36600", gate_txt(POST, "ecology")),
        ("课题一 正式障碍因子", (f"{kos['official_ranking_status_cn']}: " + ("、".join(x["factor"] for x in kos["official"]) or "无")
                          + f"；探索性待复核 {kos['n_exploratory']} 项") if kos else "未运行"),
        ("课题二 重构可行性", " / ".join(
            f"{cn} {_fmt((snap['reconstruction'][t]['headline'] or {}).get('score'))}"
            f"（{(snap['reconstruction'][t]['headline'] or {}).get('grade') or '未计算'}）"
            for t, cn in (("production", "生产"), ("ecology", "生态")))),
        ("课题三 修复后 SSUI", " / ".join(
            f"{cn} {_fmt(ssui['headline'].get(t))}（{ssui['headline_grade'].get(t) or '无修复后批次'}）"
            for t, cn in (("production", "生产"), ("ecology", "生态"))) + "；方法状态 provisional"),
        ("修复前情景判断", (snap["utilization"].get(PRE) or {}).get("state_cn") or "未运行"),
        ("修复后利用结论", (snap["utilization"].get(POST) or {}).get("state_cn") or "未运行"),
        ("五阶段业务记录", f"{snap['workflow']['business_completed']}/{snap['workflow']['business_total']} 阶段已完成并有材料；"
                     f"{snap['workflow']['business_with_documents']} 个阶段上传了材料"),
        ("软件操作里程碑", f"{snap['workflow']['software_completed']}/{snap['workflow']['software_total']}（系统内操作, 不等于五阶段业务完成）"),
    ]
    b.blocks.append({"t": "summary", "title": "报告摘要（与正文同一评价快照）", "rows": summ})

    b.h1("场地基本信息")
    b.kv([("场地名称", site.get("name")), ("场地编号", site.get("site_code")),
          ("污染类型", snap["site"]["pollution_type_cn"]), ("用地类型", site.get("land_use_type") or "未填写"),
          ("省 / 市", f"{site.get('province') or '—'} / {site.get('city') or '—'}"),
          ("中心坐标", f"{_fmt(site.get('longitude'), 6)}, {_fmt(site.get('latitude'), 6)}")])

    b.h1("数据来源说明（按阶段与批次）")
    b.p("修复前与修复后数据分开统计; 同一样品在多个轨道工作簿中重复导入时按样品编号去重, 重复导入不代表额外采样。", "muted")
    rows = []
    for st in (PRE, POST):
        for x in inv[st]["batches"]:
            rows.append([inv[st]["stage_cn"], f"#{x['batch_id']}", x["subproject"], x["kind"],
                         (f"#{x['ssui_batch_id']} {x['ssui_track']}" if x["ssui_batch_id"] else "—"),
                         x["source_file"], x["n_point_records"], x["n_measurement_records"], x["data_origin"]])
    b.table(["阶段", "批次", "子课题", "类型", "课题三批次/轨道", "来源文件", "点位记录", "检测记录", "数据来源"], rows,
            widths=[6, 6, 6, 13, 11, 28, 8, 8, 12], caption="表 导入批次清单")
    b.table(["阶段", "批次数", "点位记录", "独立样品", "检测记录", "去重检测记录", "说明"],
            [[inv[st]["stage_cn"], inv[st]["n_batches"], inv[st]["n_point_records"], inv[st]["n_unique_samples"],
              inv[st]["n_measurement_records"], inv[st]["n_unique_measurements"], inv[st]["note"] or "—"] for st in (PRE, POST)],
            widths=[8, 8, 9, 9, 9, 10, 45], caption="表 分阶段样品与记录统计")
    s2 = inv.get("s2_indicator_batch")
    if s2:
        b.p(f"课题二 28 项指标批次 #{s2['batch_id']}（{s2['source_file']}，{s2['n_points']} 个点位，来源 {s2['data_origin']}，"
            f"溯源状态 {s2['provenance_status']}）与上述检测数据批次独立登记。", "muted")
    b.kv([("导入脚本版本", (ctx.get("data_source") or {}).get("script_version"))])

    b.h1("数据覆盖率与缺失率摘要")
    cbs = ctx.get("coverage_by_stage") or {}
    b.table(["阶段", "检测因子数", "独立样品", "有值单元", "应有单元", "覆盖率 %", "缺失率 %"],
            [[("修复前" if st == PRE else "修复后"), c["factor_count"], c["n_samples"], c["observed_cells"],
              c["expected_cells"], c["coverage_pct"], c["missing_pct"]] for st, c in cbs.items()],
            caption="表 分阶段覆盖率（未检测项不填 0）")

    b.h1("采样点信息（前 10 条）")
    sp = [p for p in ctx.get("sampling_points", []) if (p.get("stage") or PRE) == PRE][:10]
    b.table([org["point_label"], "区域", "经度", "纬度", "深度(cm)", "土壤类型"],
            [[p["point_code"], p.get("region"), _fmt(p.get("longitude"), 6), _fmt(p.get("latitude"), 6),
              f"{_fmt(p.get('depth_top_cm'))}–{_fmt(p.get('depth_bottom_cm'))}", p.get("soil_type")] for p in sp],
            caption="表 修复前采样点（前 10 条）")
    if post_i["n_unique_samples"]:
        b.p(f"修复后 {post_i['n_unique_samples']} 个独立样品来自课题三批次（无坐标登记时不在地图中显示）。", "muted")

    b.h1("地图图件与采样点空间分布")
    ms = ctx.get("map_summary") or {}
    _npre = pre_i["n_unique_samples"]
    b.kv([("修复前采样点 / 含坐标点", f"{_npre} / {ms.get('n_coord_points')}（坐标覆盖 {round(100 * (ms.get('n_coord_points') or 0) / max(_npre, 1), 2)}%）"),
          ("底图", "无(离线坐标散点, 无瓦片底图)")])
    b.img(_png_from_data_url(ms.get("map_image")),
          "图 采样点空间分布（离线渲染; 颜色为修复前点位八项重金属的最大 GB 15618 筛选值倍数, 按点位 pH 与修复前门禁登记的农用地类型取值, 与界面地图图层同一函数; 有机污染物与 GB 36600 生态口径不着色, 法规结论以法规门禁章节为准）")

    b.h1("检测数据摘要（按因子、按阶段）")
    fs = snap["factor_summary"]
    for st in (PRE, POST):
        rows = [[x["factor"], x["category"], x["unit"] + (" *" if x["unit_note"] else ""), x["n_samples"], x["min"], x["mean"], x["max"]]
                for x in fs.get(st, [])]
        b.table(["因子", "类别", "单位", "样品数", "最小", "均值", "最大"], rows,
                caption=f"表 {'修复前' if st == PRE else '修复后'}检测数据摘要（按独立样品）", empty="该阶段无数据")
    if any(x["unit_note"] for st in fs for x in fs[st]):
        b.p("* 单位已按因子默认单位显示: 旧版导入把形态标记(如 Cr(VI) 的“VI”)误存为单位, v1.2.1 已修复并在数据库中更正。", "muted")

    b.h1("数据质量校验结果")
    lv = snap.get("latest_batch_validation")
    if lv:
        b.kv([("校验对象", f"最新导入批次 #{lv['batch_id']}（{lv['stage_cn']}，{lv['source_file']}）"),
              ("校验结论", "通过" if lv["passed"] else "存在阻断性错误"),
              ("错误 / 警告", f"{lv['n_errors']} / {lv['n_warnings']}"), ("批次导入超标提示", lv["n_exceed"]),
              ("口径说明", lv["note"])])
    else:
        b.p("尚无导入批次。", "muted")
    b.table(["阶段", "批次", "导入状态", "被拒记录"],
            [[inv[st]["stage_cn"], f"#{x['batch_id']}", x["import_status"], x["n_rejected"]] for st in (PRE, POST) for x in inv[st]["batches"]],
            caption="表 各批次导入状态")

    b.h1("法规门禁与超标统计（按阶段）")
    b.p("按 GB 15618-2018(生产/农用地, 按点位 pH)与 GB 36600-2018(生态/建设用地)筛选值、管制值逐点判定; "
        "与利用方向结论使用同一算法。门禁不通过时功能评分不能抵消。", "muted")
    for st in (PRE, POST):
        g = gates.get(st) or {}
        b.h2(f"{'修复前' if st == PRE else '修复后'}（{g.get('n_points', 0)} 个点位）")
        if not g.get("available"):
            b.p("该阶段无可用于门禁的点位数据。", "muted")
            continue
        for n in g.get("notes", []):
            b.p(n, "muted")
        for track, cn in (("production", "生产"), ("ecology", "生态")):
            x = g[track]
            b.kv([(f"{cn}门禁 {x['standard']}", x["state_cn"]),
                  ("超管制值因子", x["exceed_control"]), ("仅超筛选值因子", x["exceed_screening_only"]),
                  ("缺测必测项目", x["missing_required"] if x["missing_required"] else "无"),
                  ("判定条件", f"农用地类型 {x.get('farmland_type') or '—'}" if track == "production" else f"用地类别 {x.get('land_class') or '—'}")])
            rows = [[f["factor"], f["n_obs"], f["n_exceed_screening"], f["n_exceed_control"], f.get("worst_point"),
                     f.get("worst_value"), f.get("worst_ratio")] for f in x["factors"]
                    if (f.get("n_exceed_screening") or f.get("n_exceed_control") or track == "production")]
            b.table(["因子", "点位数", "超筛选值点数", "超管制值点数", "最不利点", "最大值(mg/kg)", "最大超筛选倍数"], rows[:30],
                    caption=f"表 {('修复前' if st == PRE else '修复后')}{cn}门禁逐因子统计" + ("（仅列超标因子）" if track == "ecology" else ""),
                    empty="无超筛选值因子")
        b.img(chart_gate_ratios(g["production"], f"{'修复前' if st == PRE else '修复后'} GB 15618 最大超筛选倍数"),
              f"图 {'修复前' if st == PRE else '修复后'}生产门禁: 各因子最不利点位浓度/筛选值(红=超管制值, 橙=超筛选值)", 0.8)

    b.h1("关键障碍因子识别（课题一 KOS · 规则主导 + 模型辅助）")
    diag = ctx.get("diagnosis") or {}
    if kos:
        b.kv([("诊断记录", f"#{kos['diagnosis_id']}（{kos['created_at']}，修复前数据）"),
              ("模型", f"{diag.get('model_name')} {diag.get('model_version')}"),
              ("轨道 / 子集", f"{kos['track']} / {kos['subset']}"),
              ("阈值条件", f"GB 15618 农用地类型 {kos.get('farmland_type') or '未登记'}; 按点位 pH 选档" if kos["track"] == "prod" else "GB 36600 用地类别"),
              ("诊断方法", "KOS = B×(0.30R+0.25W+0.15M+0.20S+0.10E)"),
              ("正式排名状态", kos["official_ranking_status_cn"]),
              ("正式排名准入规则", kos.get("rule") or "—")])
        if kos.get("legacy_payload"):
            b.p("注意: 该诊断记录由 v1.2.1 之前版本生成, 未区分正式/探索性结果, 请重新运行诊断后再出具报告。", "warn")
        comp = {x.get("factor"): x for x in (diag.get("kos") or {}).get("key_obstacles", [])}
        gpf = ((gates.get(PRE) or {}).get("production") or {}).get("farmland_type")
        kf = kos.get("farmland_type")
        if kos["track"] == "prod":
            if not kf:
                b.p("注意: 该诊断记录未登记 GB 15618 农用地类型(旧版本记录), 无法确认与法规门禁使用同一阈值档。", "warn")
            elif gpf and gpf not in ("未指定",) and kf.replace("(默认)", "") != gpf:
                b.p(f"注意: KOS 使用农用地类型“{kf}”, 法规门禁使用“{gpf}”, 两处超标点数可能不同; "
                    "请以相同农用地类型重新运行障碍因子诊断。", "warn")
        b.h2("正式关键障碍因子 Top-N（仅本轨官方标准证据）")
        b.table(["排名", "因子", "KOS", "R/W/M/S/E", "判定值", "阈值", "方向", "标准/条件", "判定点"],
                [[x["rank"], x["factor"], x["KOS"], _comp_txt(comp.get(x["factor"], {})), x["value"],
                  f"{_fmt(x['threshold_value'])} {x['threshold_unit'] or ''}", _dir(x["threshold_type"]),
                  f"{x['threshold_standard']} {x.get('threshold_condition') or ''}", x["decision_point_code"] or x["decision_point_id"]]
                 for x in kos["official"]],
                widths=[6, 12, 8, 20, 9, 11, 8, 17, 9], empty="证据不足: 无因子满足正式排名准入条件")
        b.h2("探索性提示（文献参考/跨路径/启发式阈值 · 待复核, 不作正式结论）")
        b.table(["序", "因子", "参考分", "判定值", "参考值", "方向", "参考来源", "状态", "判定点"],
                [[x["rank"], x["factor"], x["KOS"], x["value"], f"{_fmt(x['threshold_value'])} {x['threshold_unit'] or ''}",
                  _dir(x["threshold_type"]), x["threshold_standard"], x["threshold_resolution_status"],
                  x["decision_point_code"] or x["decision_point_id"]] for x in kos["exploratory"]],
                widths=[5, 13, 8, 9, 11, 8, 24, 11, 11])
        if kos.get("unit_unresolved"):
            b.p("单位无法证明换算而被排除的因子: " + "、".join(f"{u.get('original_name')}({u.get('unit_raw')})" for u in kos["unit_unresolved"]), "warn")
        b.img(chart_kos(kos), "图 KOS 评分(深蓝=正式, 橙=探索性/待复核)", 0.8)
        k2 = diag.get("kos") or {}
        if k2.get("model_contribution_scope") == "local_point":
            b.p(f"模型贡献度说明: 当前为{org['point_label']} {k2.get('decision_point_code') or k2.get('decision_point_id') or '—'} 的局部 SHAP；"
                "所有模型输入来自同一点位。非因果强度, 法规判定以 R(阈值超标)为准。", "muted")
        else:
            b.p("模型贡献度说明: 局部解释不可用，当前仅展示训练集全局背景贡献，不代表本场地局部贡献；非因果强度。", "muted")
        b.img(_png_from_data_url((ctx.get("map_summary") or {}).get("shap_image")), "图 模型解释贡献份额(辅助识别)", 0.8)
        mc = k2.get("model_contribution") or []
        b.table(["因子", "贡献份额", "方向", "局部 SHAP", "范围"],
                [[m.get("factor"), m.get("contribution"), m.get("direction"), m.get("local_shap_value"), m.get("contribution_scope")] for m in mc[:12]],
                caption="表 模型贡献度")
        pps = k2.get("per_point_stats") or {}
        off_set = {x["factor"] for x in kos["official"]} | set(k2.get("formal_factors") or [])
        def _layer(f, s):
            if s.get("layer"):
                return s["layer"]
            sts = {d.get("threshold_resolution_status") for d in s.get("point_details") or []} or {s.get("threshold_resolution_status")}
            sts.discard(None)
            if not sts:  # 旧记录无状态信息: 以是否进入正式 Top-N 判定
                return "formal" if f in off_set else "exploratory"
            return "formal" if sts <= {"resolved"} else "exploratory"
        pps = {f: {**s, "layer": _layer(f, s)} for f, s in pps.items()}
        if pps and k2.get("n_sampling_points"):
            formal_rows = [[f, f"{s.get('n_exceed_points')}/{s.get('n_total_points')}", s.get("exceed_rate"), s.get("max_value"),
                            s.get("p95"), s.get("median"), s.get("max_exceedance_ratio")] for f, s in pps.items()
                           if s.get("layer", "formal") == "formal" and f != "pH"]
            b.table(["因子", "超标点/有效点", "超标率", "最大值", "P95", "中位数", "最大超标倍数"], formal_rows,
                    caption=f"表 修复前逐点超标统计（共 {k2.get('n_sampling_points')} 个点位, 仅本轨官方阈值因子）")
            exp_rows = [[f, _dir(s.get("threshold_type")), f"{s.get('n_exceed_points')}/{s.get('n_total_points')}",
                         s.get("max_value"), s.get("median"), s.get("threshold_resolution_status")] for f, s in pps.items()
                        if s.get("layer") == "exploratory" and f != "pH"]
            b.table(["因子", "方向", "不足/超参考值点数", "最大值", "中位数", "参考值状态"], exp_rows,
                    caption="表 探索性指标逐点统计（下限指标计数为低于参考值的点, 非法规超标）", empty="无")
        _ = off_set
        b.h2("开放集识别结果（辅助识别 · 非法规超标结论）")
        b.table(["原始名称", "规范名称", "候选原因", "需复核"],
                [[c.get("original_name"), c.get("canonical"), c.get("reason"), c.get("review_required")] for c in ctx.get("model_candidates") or []],
                caption="① 模型候选障碍")
        b.table(["原始名称", "匹配族群", "置信度", "原始单位"],
                [[f.get("original_name"), f.get("matched_family"), f.get("family_match_confidence"), f.get("original_unit")]
                 for f in ctx.get("family_alerts") or []], caption="② 族群级预警")
        b.table(["原始名称", "单位", "最大值", "点位数", "未识别原因"],
                [[u.get("original_name"), u.get("original_unit"), u.get("max"), u.get("n_points"), u.get("unknown_reason")]
                 for u in ctx.get("unknown_measured") or []], caption="③ 未识别检测因子")
    else:
        b.p("本场地的障碍因子诊断尚未运行。", "muted")

    b.h1("功能重构可行性评价（课题二）")
    rc = snap["reconstruction"]
    for t, cn in (("production", "生产"), ("ecology", "生态")):
        h = rc[t]["headline"]
        if not h:
            b.p(f"{cn}功能重构: 未计算。", "muted")
            continue
        b.kv([(f"{cn}功能重构", f"{_fmt(h['score'])}（{h['grade']}）"), ("依据", rc[t]["headline_basis"]),
              ("方法 / 状态", f"{h['method_version'] or '—'} / {h['method_status'] or '—'}"),
              ("来源批次", f"#{h['source_batch_id']}" if h["source_batch_id"] else "—"), ("评价记录", f"#{h['evaluation_id']}")])
    others = [[r["evaluation_id"], r["eval_type"], r["score"], r["grade"], r["method_status"], r["source_batch_id"]]
              for t in rc for r in rc[t]["all_records"][1:]]
    b.table(["记录", "类型", "得分", "等级", "方法状态", "来源批次"], others, caption="表 其它历史评价记录(不作为本报告结论)", empty="无")
    for ev in ctx.get("reconstruction") or []:
        if ev.get("limiting_factors"):
            b.p(f"{ev['title']}关键限制因子: {_fmt(ev['limiting_factors'] if isinstance(ev['limiting_factors'], list) else list(ev['limiting_factors']))}", "muted")

    b.h1("可持续利用评价（SSUI · 课题三修复后）")
    b.p(ssui["note"], "muted")
    rows = []
    for t, cn in (("production", "生产"), ("ecology", "生态")):
        p = ssui["post"][t]; e = p["evaluation"] or {}
        rows.append([cn, f"#{p['batch_id']}" if p["batch_id"] else "无", p["source_file"], e.get("score"), e.get("grade"),
                     p["years_since_remediation"], p["multiplier_m"], e.get("exceeds_unit_range"), e.get("method_status")])
    b.table(["轨道", "批次", "来源文件", "SSUI", "等级", "t(年)", "M", "超出单位区间", "方法状态"], rows,
            widths=[7, 7, 30, 9, 11, 7, 7, 10, 12], caption="表 修复后 SSUI（每轨最新已确认批次）")
    b.img(chart_ssui(ssui), "图 修复后 SSUI 与 0.6 暂定支持阈值", 0.55)
    pr = ssui.get("pre_reference")
    if pr:
        b.p(f"{pr['label']}: 记录 #{pr['evaluation_id']}，结果 {_fmt(pr['score'])}（{pr['grade'] or '—'}）。该项不是课题三修复后 SSUI, 不进入首页结论。", "muted")

    b.h1("利用方向结论（法规门禁 + 功能评分）")
    for st, lbl in ((PRE, "修复前情景判断"), (POST, "修复后利用结论")):
        u = snap["utilization"].get(st)
        if not u:
            b.p(f"{lbl}: 尚未运行。", "muted")
            continue
        b.kv([(lbl, u["state_cn"]), ("决策记录", f"#{u['decision_id']}（{u['created_at']}）"),
              ("门禁 生产/生态", f"{u['production_gate']} / {u['ecology_gate']}"),
              ("评分 生产/生态", f"{_fmt(u['production_score'])} / {_fmt(u['ecology_score'])}"),
              ("方法版本 / 状态", f"{u['method_version']} / {u['method_status']}")])
        b.p(f"{lbl}：{u['conclusion']}")
        for m in (u.get("missing_evidence") or [])[:6]:
            b.p(f"需补充：{m}", "muted")
    b.p("“支持利用”为系统辅助判断(方法状态 provisional), 不构成正式利用许可; 正式结论需课题组与主管部门签认。", "warn")

    b.h1("推荐重构方案与推荐修复方案矩阵")
    recs = ctx.get("recommendations") or []
    b.table(["排序", "技术", "匹配度", "成本/周期", "禁用条件", "推荐理由（摘要）"],
            [[r["rank"], r["technology"], r["match_score"], f"{r.get('cost_level') or '—'} / {r.get('duration_level') or '—'}",
              r.get("forbidden_conditions"), (r.get("reason") or "")[:120]] for r in recs],
            widths=[6, 16, 8, 12, 22, 36], empty="本场地尚未生成推荐方案（未在“方案推荐”页面执行匹配）。")

    b.h1("修复案例证据库")
    b.table(["案例", "区域/用地", "污染物", "技术", "效果", "证据来源"],
            [[c["case_id"], f"{c.get('region') or '—'} / {c.get('land_use') or '—'}", c.get("pollutants"),
              c.get("remediation_technology"), c.get("effectiveness"), c.get("evidence_source") or c.get("doi")]
             for c in ctx.get("remediation_cases") or []], widths=[10, 14, 16, 18, 22, 20])

    b.h1("五阶段全流程追溯记录")
    wf = snap["workflow"]
    b.p(wf["note"], "muted")
    b.table(["业务阶段", "状态", "附件数", "操作时间", "审批意见"],
            [[x["name"], x["status_cn"], x["n_attachments"], x.get("operated_at"), x.get("review_comment")] for x in wf["business_stages"]],
            caption="表 五阶段业务记录（调查评估→后期管护）")
    b.table(["软件操作里程碑", "是否执行", "次数"],
            [[m["name"], ("本报告" if m["key"] == "report" else "已执行" if m["done"] else "未执行"), m.get("count")]
             for m in wf["software_milestones"]],
            caption=f"表 七项软件操作里程碑（{wf['software_completed']}/{wf['software_total']}，不等于五阶段业务完成）")

    b.h1("附件清单")
    b.table(["阶段", "材料类型", "文件名"],
            [[a.get("stage_name"), a.get("file_role"), a.get("original_name")] for a in ctx.get("attachments") or []],
            empty="暂无附件（五阶段业务材料未上传）。")

    b.h1("模型版本、数据版本、标准版本、报告版本")
    b.kv([("模型版本", f"{diag.get('model_name')} {diag.get('model_version')}" if diag else "—"),
          ("数据版本", rep.get("data_version")), ("标准版本", rep.get("standard_version")),
          ("模板版本 / 报告版本", f"{rep.get('template_version')} / {rep.get('version')}"),
          ("评价快照", f"{snap['snapshot_id']}（{snap['schema']}）"), ("快照 SHA-256", snap["snapshot_sha256"]),
          ("门禁标准文件 SHA-256(前16位)", (gates.get(PRE) or gates.get(POST) or {}).get("standards_sha256"))])
    b.table(["标准号", "标准名称", "版本", "来源"],
            [[s["standard_code"], s["standard_name"], s["version"], s.get("source_reference")] for s in ctx.get("standard_versions") or []],
            widths=[16, 40, 10, 34])
    b.p("报告口径: 正式超标结论仅限身份明确、单位可证换算且阈值适用(本轨官方标准)的因子; 文献参考、跨路径与启发式阈值结果"
        "为探索性提示。本版本在 3 个原始场地与 5 类模拟演示场景上完成工程回归, 尚未开展跨区域独立验证。", "muted")

    b.h1("人工复核意见区")
    b.kv([("复核结论", "□ 采纳系统结论　□ 部分采纳　□ 驳回并要求补充数据"), ("复核意见", " "), ("复核人 / 复核时间", " ")])
    return b.blocks


def _dir(t):
    return {"upper": "上限(超标)", "lower": "下限(不足)", "interval": "区间", "unknown": "未知"}.get(t or "", t or "—")


def _comp_txt(k: dict) -> str:
    c = k.get("components") or {}
    if not c:
        return "—"
    return " / ".join(f"{n}={c.get(n) if c.get(n) is not None else '—'}" for n in ("R", "W", "M", "S", "E"))


def plain_text(blocks: list[dict]) -> str:
    out = []
    for bl in blocks:
        t = bl["t"]
        if t in ("h1", "h2", "p"):
            out.append(bl["text"])
        elif t in ("kv", "summary", "cover"):
            out += [bl.get("title", "")] + [f"{k} {v}" for k, v in bl["rows"]] + [bl.get("banner") or ""]
        elif t == "table":
            out += [bl.get("caption") or ""] + [" ".join(bl["header"])] + [" ".join(r) for r in bl["rows"]]
        elif t == "img":
            out.append(bl["caption"])
    return "\n".join(x for x in out if x)


# ───────────────────────── HTML ─────────────────────────
def render_html(ctx: dict, blocks: list[dict] | None = None) -> str:
    blocks = blocks or build_blocks(ctx)
    e = _html.escape
    css = ("body{font-family:'Microsoft YaHei','PingFang SC','SimHei',sans-serif;font-size:10.5pt;color:#222;margin:24px 36px}"
           "h1{font-size:15pt;color:#1F4E79;border-bottom:2px solid #1F4E79;padding-bottom:3px;margin-top:22px}"
           "h2{font-size:12pt;color:#1F4E79;margin-top:14px}table{border-collapse:collapse;width:100%;margin:6px 0}"
           "td,th{border:1px solid #bbb;padding:3px 5px;font-size:9pt;vertical-align:top}th{background:#1F4E79;color:#fff}"
           "tr:nth-child(even) td{background:#F4F7FB}.kv td:first-child{width:28%;background:#EEF3F8;font-weight:bold}"
           ".muted{color:#666;font-size:9pt}.warn{color:#A94400;font-weight:bold}.banner{background:#C0392B;color:#fff;"
           "padding:6px 10px;font-weight:bold}.cap{color:#444;font-size:9pt;margin-top:8px}.summary td:first-child{width:30%}"
           "img{max-width:100%;border:1px solid #eee}")
    h = [f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>{e(ctx['site'].get('name') or '')} 追溯报告</title>"
         f"<style>{css}</style></head><body>"]
    for bl in blocks:
        t = bl["t"]
        if t == "cover":
            if bl.get("banner"):
                h.append(f"<div class='banner'>{e(bl['banner'])}</div>")
            h.append(f"<h1 style='border:none;font-size:20pt'>{e(bl['title'])}</h1><div class='muted'>{e(bl['subtitle'])}</div>")
            h.append("<table class='kv'>" + "".join(f"<tr><td>{e(str(k))}</td><td>{e(_fmt(v))}</td></tr>" for k, v in bl["rows"]) + "</table>")
        elif t == "summary":
            h.append(f"<h2>{e(bl['title'])}</h2><table class='kv summary'>" +
                     "".join(f"<tr><td>{e(k)}</td><td>{e(_fmt(v))}</td></tr>" for k, v in bl["rows"]) + "</table>")
        elif t == "h1":
            h.append(f"<h1>{e(bl['text'])}</h1>")
        elif t == "h2":
            h.append(f"<h2>{e(bl['text'])}</h2>")
        elif t == "p":
            h.append(f"<p class='{bl['style']}'>{e(bl['text'])}</p>")
        elif t == "kv":
            h.append("<table class='kv'>" + "".join(f"<tr><td>{e(k)}</td><td>{e(v)}</td></tr>" for k, v in bl["rows"]) + "</table>")
        elif t == "table":
            if bl.get("caption"):
                h.append(f"<div class='cap'>{e(bl['caption'])}</div>")
            h.append("<table><tr>" + "".join(f"<th>{e(c)}</th>" for c in bl["header"]) + "</tr>" +
                     "".join("<tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>" for r in bl["rows"]) + "</table>")
        elif t == "img":
            h.append(f"<div><img src='data:image/png;base64,{base64.b64encode(bl['png']).decode()}' "
                     f"style='width:{int(bl['width'] * 100)}%'/></div><div class='cap'>{e(bl['caption'])}</div>")
    h.append(f"<p class='muted'>报告版本 {e(str(ctx['report'].get('version')))}｜评价快照 {e(ctx['snapshot']['snapshot_id'])}｜"
             f"本报告由“污染场地土壤生态-生产功能重构监管系统”生成。</p></body></html>")
    return "".join(h)


# ───────────────────────── PDF (ReportLab platypus) ─────────────────────────
_FONT_CANDIDATES = [
    (r"C:\Windows\Fonts\msyh.ttc", 0), (r"C:\Windows\Fonts\simhei.ttf", None), (r"C:\Windows\Fonts\simsun.ttc", 0),
    ("/System/Library/Fonts/STHeiti Medium.ttc", 0), ("/System/Library/Fonts/STHeiti Light.ttc", 0),
    ("/System/Library/Fonts/Hiragino Sans GB.ttc", 0), ("/Library/Fonts/Arial Unicode.ttf", None),
    ("/System/Library/Fonts/Supplemental/Arial Unicode.ttf", None),
    ("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc", 0), ("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc", 0),
    ("/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf", None),
]


def _register_font() -> dict:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    cands = list(_FONT_CANDIDATES)
    env = os.environ.get("SRS_REPORT_FONT")
    if env:
        cands.insert(0, (env, 0 if env.lower().endswith(".ttc") else None))
    windir = os.environ.get("WINDIR")
    if windir:
        cands = [(os.path.join(windir, "Fonts", "msyh.ttc"), 0), (os.path.join(windir, "Fonts", "simhei.ttf"), None)] + cands
    for path, idx in cands:
        if not os.path.isfile(path):
            continue
        try:
            f = TTFont("SRSCJK", path, subfontIndex=idx) if idx is not None else TTFont("SRSCJK", path)
            pdfmetrics.registerFont(f)
            return {"name": "SRSCJK", "embedded": True, "source": os.path.basename(path)}
        except Exception:
            continue
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return {"name": "STSong-Light", "embedded": False, "source": "Adobe CID STSong-Light(未嵌入, 依赖阅读器亚洲字体包)"}


def render_pdf(ctx: dict, blocks: list[dict] | None = None) -> tuple[bytes, dict]:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)
    blocks = blocks or build_blocks(ctx)
    font = _register_font()
    fn = font["name"]
    navy = colors.HexColor("#1F4E79")
    st = {
        "body": ParagraphStyle("b", fontName=fn, fontSize=9.5, leading=14, spaceAfter=3, wordWrap="CJK"),
        "muted": ParagraphStyle("m", fontName=fn, fontSize=8.3, leading=12, textColor=colors.HexColor("#555555"), spaceAfter=3, wordWrap="CJK"),
        "warn": ParagraphStyle("w", fontName=fn, fontSize=9, leading=13, textColor=colors.HexColor("#A94400"), spaceAfter=4, wordWrap="CJK"),
        "h1": ParagraphStyle("h1", fontName=fn, fontSize=14, leading=19, textColor=navy, spaceBefore=10, spaceAfter=5),
        "h2": ParagraphStyle("h2", fontName=fn, fontSize=11.5, leading=16, textColor=navy, spaceBefore=6, spaceAfter=3),
        "cell": ParagraphStyle("c", fontName=fn, fontSize=7.6, leading=10, wordWrap="CJK"),
        "head": ParagraphStyle("hd", fontName=fn, fontSize=7.8, leading=10, textColor=colors.white, wordWrap="CJK"),
        "cap": ParagraphStyle("cap", fontName=fn, fontSize=8, leading=11, textColor=colors.HexColor("#333333"), spaceBefore=4, spaceAfter=2),
        "title": ParagraphStyle("t", fontName=fn, fontSize=22, leading=30, alignment=TA_CENTER, textColor=navy, spaceAfter=6),
        "sub": ParagraphStyle("s", fontName=fn, fontSize=12, leading=18, alignment=TA_CENTER, textColor=colors.HexColor("#444444")),
        "banner": ParagraphStyle("bn", fontName=fn, fontSize=11, leading=16, alignment=TA_CENTER, textColor=colors.white),
    }
    W = A4[0] - 36 * mm - 14
    esc = _html.escape

    def P(text, s="cell"):
        return Paragraph(esc(str(text)).replace("\n", "<br/>"), st[s])

    def kv_table(rows, w0=0.30, shade="#EEF3F8"):
        t = Table([[P(k), P(v)] for k, v in rows], colWidths=[W * w0, W * (1 - w0)])
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B0B7C3")),
                               ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(shade)),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]))
        return t
    story = []
    snap = ctx["snapshot"]
    for bl in blocks:
        t = bl["t"]
        if t == "cover":
            story += [Spacer(1, 30 * mm)]
            if bl.get("banner"):
                bt = Table([[P(bl["banner"], "banner")]], colWidths=[W])
                bt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#C0392B")),
                                        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
                story += [bt, Spacer(1, 10 * mm)]
            story += [P(bl["title"], "title"), P(bl["subtitle"], "sub"), Spacer(1, 16 * mm),
                      kv_table([(k, _fmt(v)) for k, v in bl["rows"]], 0.28), PageBreak()]
        elif t == "summary":
            story += [P(bl["title"], "h1"), kv_table([(k, _fmt(v)) for k, v in bl["rows"]], 0.27, "#E3EDF7"), PageBreak()]
        elif t == "h1":
            story.append(P(bl["text"], "h1"))
        elif t == "h2":
            story.append(P(bl["text"], "h2"))
        elif t == "p":
            story.append(P(bl["text"], bl["style"] if bl["style"] in st else "body"))
        elif t == "kv":
            story += [kv_table(bl["rows"]), Spacer(1, 3)]
        elif t == "table":
            n = len(bl["header"])
            ws = bl.get("widths") or [1] * n
            tot = float(sum(ws))
            cw = [W * w / tot for w in ws]
            data = [[P(h, "head") for h in bl["header"]]] + [[P(c) for c in r] for r in bl["rows"]]
            tb = Table(data, colWidths=cw, repeatRows=1)
            sty = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B0B7C3")), ("BACKGROUND", (0, 0), (-1, 0), navy),
                   ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2),
                   ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
            for i in range(2, len(data), 2):
                sty.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F4F7FB")))
            tb.setStyle(TableStyle(sty))
            if bl.get("caption"):
                story.append(P(bl["caption"], "cap"))
            story += [tb, Spacer(1, 4)]
        elif t == "img":
            from reportlab.lib.utils import ImageReader
            ir = ImageReader(io.BytesIO(bl["png"]))
            iw, ih = ir.getSize()
            w = W * bl.get("width", 0.9)
            h = w * ih / iw
            if h > 175 * mm:
                h = 175 * mm; w = h * iw / ih
            story.append(KeepTogether([Image(io.BytesIO(bl["png"]), width=w, height=h), P(bl["caption"], "cap")]))
    title = f"{ctx['site'].get('site_code') or ''} 全流程监管追溯报告｜{(ctx['site'].get('name') or '')[:24]}"
    label = snap["data_origin"].get("label")

    def on_page(c, doc):
        c.saveState()
        c.setFont(fn, 7.5)
        c.setFillColor(colors.HexColor("#555555"))
        if doc.page > 1:
            c.drawString(18 * mm, A4[1] - 11 * mm, title[:60])
            c.drawRightString(A4[0] - 18 * mm, A4[1] - 11 * mm, f"评价快照 {snap['snapshot_id']}")
            c.setStrokeColor(navy); c.setLineWidth(0.5)
            c.line(18 * mm, A4[1] - 12.5 * mm, A4[0] - 18 * mm, A4[1] - 12.5 * mm)
        c.drawRightString(A4[0] - 18 * mm, 9 * mm, f"第 {doc.page} 页")
        c.drawString(18 * mm, 9 * mm, f"报告版本 {ctx['report'].get('version')}｜字体 {font['source']}")
        if label:
            c.setFillColor(colors.HexColor("#C0392B"))
            c.drawCentredString(A4[0] / 2, 9 * mm, label)
        c.restoreState()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=17 * mm,
                            bottomMargin=15 * mm, title=title, author="污染场地土壤生态-生产功能重构监管系统",
                            subject=f"snapshot {snap['snapshot_id']}")
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue(), font


# ───────────────────────── DOCX ─────────────────────────
def render_docx(ctx: dict, blocks: list[dict] | None = None) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    blocks = blocks or build_blocks(ctx)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.6)
    usable = sec.page_width - sec.left_margin - sec.right_margin
    navy = RGBColor(0x1F, 0x4E, 0x79)

    def set_font(run, size=10.5, bold=False, color=None, east="宋体", latin="Times New Roman"):
        run.font.size = Pt(size); run.font.bold = bold; run.font.name = latin
        rpr = run._element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        rf.set(qn("w:eastAsia"), east)
        if color is not None:
            run.font.color.rgb = color
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for lvl, size in (("Heading 1", 14), ("Heading 2", 12)):
        hs = doc.styles[lvl]
        hs.font.size = Pt(size); hs.font.color.rgb = navy; hs.font.name = "Times New Roman"
        hs.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    def shade(cell, hexcolor):
        tcPr = cell._element.get_or_add_tcPr()
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor)
        tcPr.append(sh)

    def cell_text(cell, text, size=8.5, bold=False, color=None, east="宋体"):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(str(text)), size, bold, color, east)

    def repeat_header(row):
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)

    def kv(rows, w0=0.3, fill="EEF3F8"):
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, v in rows:
            c = t.add_row().cells
            cell_text(c[0], k, 9, True, east="黑体"); shade(c[0], fill)
            cell_text(c[1], v, 9)
            c[0].width = int(usable * w0); c[1].width = int(usable * (1 - w0))
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def para(text, style="body"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        size, color = {"body": (10, None), "muted": (8.5, RGBColor(0x55, 0x55, 0x55)),
                       "warn": (9.5, RGBColor(0xA9, 0x44, 0x00))}.get(style, (10, None))
        set_font(p.add_run(text), size, style == "warn", color)
        return p
    for bl in blocks:
        t = bl["t"]
        if t == "cover":
            for _ in range(4):
                doc.add_paragraph()
            if bl.get("banner"):
                tb = doc.add_table(rows=1, cols=1); tb.style = "Table Grid"
                cell_text(tb.rows[0].cells[0], bl["banner"], 12, True, RGBColor(0xFF, 0xFF, 0xFF), "黑体")
                shade(tb.rows[0].cells[0], "C0392B")
                tb.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.add_paragraph()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(bl["title"]), 22, True, navy, "黑体")
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(bl["subtitle"]), 13, False, RGBColor(0x44, 0x44, 0x44), "黑体")
            if bl.get("banner"):
                para(bl["banner"], "warn").alignment = WD_ALIGN_PARAGRAPH.CENTER
            for _ in range(3):
                doc.add_paragraph()
            kv([(k, _fmt(v)) for k, v in bl["rows"]], 0.28)
            doc.add_page_break()
        elif t == "summary":
            doc.add_heading(bl["title"], level=1)
            kv([(k, _fmt(v)) for k, v in bl["rows"]], 0.27, "E3EDF7")
            doc.add_page_break()
        elif t == "h1":
            doc.add_heading(bl["text"], level=1)
        elif t == "h2":
            doc.add_heading(bl["text"], level=2)
        elif t == "p":
            para(bl["text"], bl["style"])
        elif t == "kv":
            kv(bl["rows"])
        elif t == "table":
            if bl.get("caption"):
                para(bl["caption"], "muted").paragraph_format.keep_with_next = True
            n = len(bl["header"])
            ws = bl.get("widths") or [1] * n
            tot = float(sum(ws))
            tb = doc.add_table(rows=1, cols=n)
            tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
            repeat_header(tb.rows[0])
            for i, h in enumerate(bl["header"]):
                c = tb.rows[0].cells[i]
                cell_text(c, h, 8.5, True, RGBColor(0xFF, 0xFF, 0xFF), "黑体"); shade(c, "1F4E79")
            for ri, r in enumerate(bl["rows"]):
                cells = tb.add_row().cells
                for i, v in enumerate(r):
                    cell_text(cells[i], v, 8)
                    if ri % 2 == 1:
                        shade(cells[i], "F4F7FB")
            for row in tb.rows:
                for i, c in enumerate(row.cells):
                    c.width = int(usable * ws[i] / tot)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif t == "img":
            doc.add_picture(io.BytesIO(bl["png"]), width=int(usable * bl.get("width", 0.9)))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            para(bl["caption"], "muted").alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 页眉页脚: 快照编号 + 页码(域代码) + 模拟数据标签
    snap = ctx["snapshot"]
    hp = sec.header.paragraphs[0]
    set_font(hp.add_run(f"{ctx['site'].get('site_code') or ''} 全流程监管追溯报告｜评价快照 {snap['snapshot_id']}"), 8,
             color=RGBColor(0x55, 0x55, 0x55))
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lbl = snap["data_origin"].get("label")
    if lbl:
        set_font(fp.add_run(lbl + "　"), 8, True, RGBColor(0xC0, 0x39, 0x2B))
    set_font(fp.add_run("第 "), 8)
    r = fp.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); r._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt; r._r.append(el)
    set_font(fp.add_run(" 页"), 8)
    cp = doc.core_properties
    cp.title = f"{ctx['site'].get('name') or ''} 全流程监管追溯报告"
    cp.subject = f"snapshot {snap['snapshot_id']}"
    cp.author = "污染场地土壤生态-生产功能重构监管系统"
    bio = io.BytesIO(); doc.save(bio)
    return bio.getvalue()
