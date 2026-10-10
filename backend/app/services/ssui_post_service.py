"""v1.1 课题三: 修复后 SSUI 独立导入 → 校验预览 → 事务确认 → 计算 → 持久化 → 导出。

C4 要求: 课题三使用修复后数据, 具有独立导入功能, 与课题一/二的修复前检测数据分离。

模板(SSUI-POST-v1.1) 四个工作表:
  说明            — 填写规则(只读)
  批次信息        — 场地编号/评价轨道/评价年份/修复后年数 t/管理强度或 M/数据来源
  指标得分        — D1–D25: 指标编码, 指标名称, 原始值(可选), 单位(可选), 得分 s_i(0–1, 必填), 备注
  修复后污染物检测 — 点位编号, pH, 污染物, 浓度, 单位(mg/kg) — 供法规安全门禁使用

预览不写任何业务数据(仅写一条 status=previewed 的批次记录, 便于确认与审计);
确认在单一事务中写入 ssui_records + 修复后测值(stage=post_remediation) + 评价结果, 任一步失败全部回滚。
"""
from __future__ import annotations

import hashlib
import io
import os
import sys
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from app.core.config import resource_root
from app.models import (POST_REMEDIATION, EvaluationResult, FactorDictionary, ImportBatch, Measurement,
                        SamplingPoint, Site, SSUIImportBatch, SSUIRecord)
from app.services import audit_service

ROOT = resource_root()
_ML = os.path.join(ROOT, "ml", "evaluation")
if _ML not in sys.path:
    sys.path.insert(0, _ML)
import ssui_v11 as SV  # noqa: E402
import utilization as U  # noqa: E402

SHEETS = ["说明", "批次信息", "指标得分", "修复后污染物检测"]
DATA_ORIGINS = {"client_real": "真实数据(甲方/课题组提供)", "field": "现场实测",
                "monte_carlo_demo": "模拟数据——仅供测试/演示", "test_fixture": "测试夹具"}
TRACKS = {"production": "生产利用", "ecology": "生态利用"}
INTENSITY = {"low": "低强度", "medium": "中等强度", "high": "高强度"}
_M_TABLE = {("production", "low"): 1.1, ("production", "medium"): 1.15, ("production", "high"): 1.2,
            ("ecology", "low"): 1.05, ("ecology", "medium"): 1.075, ("ecology", "high"): 1.1}
_UNIT_TO_MGKG = {"mg/kg": 1.0, "mgkg": 1.0, "mg·kg-1": 1.0, "mg kg-1": 1.0}
_CN_METAL = {"Cd": "镉", "Hg": "汞", "As": "砷", "Pb": "铅", "Cr": "铬", "Cu": "铜", "Ni": "镍",
             "Zn": "锌", "Cr(VI)": "六价铬"}
_INV_INTENSITY = {v: k for k, v in INTENSITY.items()}
_INV_TRACK = {v: k for k, v in TRACKS.items()}
_INV_ORIGIN = {v: k for k, v in DATA_ORIGINS.items()}


class SSUIImportError(ValueError):
    pass


# ───────────────────────── 模板 ─────────────────────────
def effective_status(ev) -> str | None:
    """v1.2.2(T01): 课题三结果的有效状态。v1.2.1 及更早版本入库的记录没有 status 字段, 按原值与等级定义域
    推导(定义域外 → out_of_domain), 只在读取时判定, 不改写库内原记录(原 grade 保留供审计)。"""
    if ev is None or not (ev.eval_type or "").startswith("ssui_post_"):
        return None
    dims = ev.dimensions or {}
    st = dims.get("status")
    if st in ("invalid", "insufficient"):
        return st
    if ev.score is None:
        return st or "insufficient"
    # v1.2.1 对所有可计算结果都写 status="ok"(含 SSUI>1), 因此不能信任旧 status: 一律按原值与定义域判定
    lo, hi = SV.validity_domain(SV.load_weights(ROOT))
    if not (lo <= float(ev.score) <= hi):
        return "out_of_domain"
    return st or "ok"


def is_legacy_record(ev) -> bool:
    """v1.2.2 起课题三结果均写入 classification_scope; 没有该字段即为旧版本(≤v1.2.1)入库记录。"""
    return bool(ev is not None and (ev.eval_type or "").startswith("ssui_post_")
                and "classification_scope" not in (ev.dimensions or {}))


def build_template(track: str = "production", site_code: str = "") -> bytes:
    if track not in TRACKS:
        raise SSUIImportError("track 必须为 production/ecology")
    W = SV.load_weights(ROOT)
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F5C99")
    ws = wb.active
    ws.title = "说明"
    lines = [
        f"SRS 课题三 修复后 SSUI 导入模板  {SV.TEMPLATE_VERSION}",
        "1. 本模板仅用于【修复后】数据; 修复前检测数据请使用'数据导入'页面(课题一/二)。",
        "2. '批次信息'表: 评价轨道填 生产利用 或 生态利用; t 为修复完成后年数(≥0); 管理强度填 低强度/中等强度/高强度, 或直接填 M。",
        f"   M 允许区间: 生产 {W['M_range']['production']}, 生态 {W['M_range']['ecology']}(方法 PPT 第 15 页)。",
        "3. '指标得分'表: 25 项指标的得分 s_i 必须填写, 取值 0–1(方法文件未给原始值→得分规则, 由课题组按分级标准赋分); 原始值/单位可选, 用于追溯。",
        "4. '修复后污染物检测'表: 用于法规安全门禁(生产 GB 15618 / 生态 GB 36600)。浓度单位必须为 mg/kg; 缺测项目不会被视为达标。",
        "5. 数据来源必须如实选择; '模拟数据——仅供测试/演示'的结果不得用于正式报告。",
        f"6. 权重来源: {W['_source']}",
        f"7. 方法状态: {W['_status']}",
    ]
    for i, t in enumerate(lines, 1):
        ws.cell(i, 1, t)
    ws.column_dimensions["A"].width = 130

    wb2 = wb.create_sheet("批次信息")
    meta = [("场地编号", site_code, "可留空(以页面所选场地为准); 填写时须与系统中场地编号一致"),
            ("评价轨道", TRACKS[track], "生产利用 / 生态利用"),
            ("评价年份", "", "如 2026"),
            ("修复后年数t", "", "≥0, 单位: 年"),
            ("管理强度", "中等强度", "低强度/中等强度/高强度; 若填写 M 则以 M 为准"),
            ("M", "", "可选, 须在方法区间内"),
            ("数据来源", DATA_ORIGINS["client_real"], " / ".join(DATA_ORIGINS.values())),
            ("备注", "", "")]
    for c, h in enumerate(["字段", "值", "说明"], 1):
        x = wb2.cell(1, c, h); x.font = head; x.fill = fill
    for r, (k, v, n) in enumerate(meta, 2):
        wb2.cell(r, 1, k); wb2.cell(r, 2, v); wb2.cell(r, 3, n)
    dv = DataValidation(type="list", formula1='"生产利用,生态利用"', allow_blank=False); wb2.add_data_validation(dv); dv.add("B3")
    dv2 = DataValidation(type="list", formula1='"低强度,中等强度,高强度"', allow_blank=True); wb2.add_data_validation(dv2); dv2.add("B6")
    dv3 = DataValidation(type="list", formula1='"' + ",".join(DATA_ORIGINS.values()) + '"', allow_blank=False)
    wb2.add_data_validation(dv3); dv3.add("B8")
    for col, w in zip("ABC", (16, 28, 60)):
        wb2.column_dimensions[col].width = w

    wb3 = wb.create_sheet("指标得分")
    hdr = ["指标编码", "指标名称", "准则层", f"权重({TRACKS[track]})", "原始值", "单位", "得分s_i(0-1)", "备注"]
    for c, h in enumerate(hdr, 1):
        x = wb3.cell(1, c, h); x.font = head; x.fill = fill
    for r, ind in enumerate(W["indicators"], 2):
        wb3.cell(r, 1, ind["code"]); wb3.cell(r, 2, ind["name"]); wb3.cell(r, 3, ind["criterion"])
        wb3.cell(r, 4, ind["w_" + track])
    dv4 = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
    wb3.add_data_validation(dv4); dv4.add("G2:G26")
    for col, w in zip("ABCDEFGH", (10, 46, 8, 14, 12, 10, 14, 30)):
        wb3.column_dimensions[col].width = w

    wb4 = wb.create_sheet("修复后污染物检测")
    for c, h in enumerate(["点位编号", "pH", "污染物", "浓度", "单位", "采样日期", "备注"], 1):
        x = wb4.cell(1, c, h); x.font = head; x.fill = fill
    wb4.cell(2, 7, "示例行请删除: 每个点位每个污染物一行; 单位 mg/kg")
    for col, w in zip("ABCDEFG", (12, 8, 22, 12, 10, 12, 40)):
        wb4.column_dimensions[col].width = w
    for s in wb.worksheets:
        for row in s.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=s.title == "说明")
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


# ───────────────────────── 解析与校验 ─────────────────────────
def _num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return "NaN"


def parse_and_validate(content: bytes, filename: str, site: Site, expected_track: str | None = None) -> dict:
    errors: list[dict] = []
    warnings: list[str] = []

    def err(sheet, row, col, msg):
        errors.append({"sheet": sheet, "row": row, "col": col, "message": msg})

    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001
        raise SSUIImportError(f"无法读取 Excel 文件: {e}")
    missing_sheets = [s for s in SHEETS[1:] if s not in wb.sheetnames]
    if missing_sheets:
        raise SSUIImportError(f"模板缺少工作表: {', '.join(missing_sheets)}(请下载 {SV.TEMPLATE_VERSION} 模板)")

    meta = {}
    for r, row in enumerate(wb["批次信息"].iter_rows(min_row=2, values_only=True), 2):
        if row and row[0]:
            meta[str(row[0]).strip()] = (row[1] if len(row) > 1 else None, r)
    code = str((meta.get("场地编号") or ("", 0))[0] or "").strip()
    if code and code not in (site.site_code, site.original_site_code or ""):
        err("批次信息", meta["场地编号"][1], "B", f"场地编号 {code} 与所选场地 {site.site_code} 不一致")
    track_cn = str((meta.get("评价轨道") or ("", 0))[0] or "").strip()
    track = _INV_TRACK.get(track_cn) or (track_cn if track_cn in TRACKS else None)
    if not track:
        err("批次信息", (meta.get("评价轨道") or (0, 3))[1], "B", "评价轨道须为 生产利用 / 生态利用")
    elif expected_track and track != expected_track:
        err("批次信息", meta["评价轨道"][1], "B", f"文件轨道 {track_cn} 与页面选择不一致")
    year = _num((meta.get("评价年份") or (None, 0))[0])
    if year in (None, "NaN") or not (2000 <= year <= 2100):
        err("批次信息", (meta.get("评价年份") or (0, 4))[1], "B", "评价年份缺失或无效")
    t = _num((meta.get("修复后年数t") or (None, 0))[0])
    if t in (None, "NaN") or t < 0:
        err("批次信息", (meta.get("修复后年数t") or (0, 5))[1], "B", "修复后年数 t 缺失或 < 0")
    M = _num((meta.get("M") or (None, 0))[0])
    inten_cn = str((meta.get("管理强度") or ("", 0))[0] or "").strip()
    intensity = _INV_INTENSITY.get(inten_cn)
    if M == "NaN":
        err("批次信息", meta["M"][1], "B", "M 不是数字"); M = None
    if M is None and track and intensity:
        M = _M_TABLE[(track, intensity)]
    if M is None:
        err("批次信息", (meta.get("管理强度") or (0, 6))[1], "B", "须填写管理强度或 M")
    elif track:
        lo, hi = SV.load_weights(ROOT)["M_range"][track]
        if not (lo - 1e-9 <= M <= hi + 1e-9):
            err("批次信息", (meta.get("M") or (0, 7))[1], "B", f"M={M} 超出 {TRACKS[track]} 方法区间 [{lo}, {hi}]")
    origin_cn = str((meta.get("数据来源") or ("", 0))[0] or "").strip()
    origin = _INV_ORIGIN.get(origin_cn) or (origin_cn if origin_cn in DATA_ORIGINS else None)
    if not origin:
        err("批次信息", (meta.get("数据来源") or (0, 8))[1], "B", "数据来源须从下拉选项中选择")

    catalog = {i["code"]: i for i in SV.load_weights(ROOT)["indicators"]}
    scores, records, seen = {}, [], set()
    for r, row in enumerate(wb["指标得分"].iter_rows(min_row=2, values_only=True), 2):
        if not row or all(v is None for v in row):
            continue
        code_i = str(row[0] or "").strip().upper()
        if code_i not in catalog:
            err("指标得分", r, "A", f"未知指标编码 {row[0]}"); continue
        if code_i in seen:
            err("指标得分", r, "A", f"指标 {code_i} 重复"); continue
        seen.add(code_i)
        s = _num(row[6] if len(row) > 6 else None)
        raw = _num(row[4] if len(row) > 4 else None)
        if s is None:
            err("指标得分", r, "G", f"{code_i} 得分 s_i 未填写"); s = None
        elif s == "NaN" or not (0 <= s <= 1):
            err("指标得分", r, "G", f"{code_i} 得分须为 0–1 的数字"); s = None
        if raw == "NaN":
            err("指标得分", r, "E", f"{code_i} 原始值不是数字"); raw = None
        scores[code_i] = s
        records.append({"indicator_code": code_i, "indicator_name": catalog[code_i]["name"], "raw_value": raw,
                        "unit": (str(row[5]).strip() if len(row) > 5 and row[5] else None), "score": s,
                        "source_sheet": "指标得分", "source_row": r, "source_col": "G",
                        "note": (str(row[7]) if len(row) > 7 and row[7] else None)})
    for code_i in catalog:
        if code_i not in seen:
            err("指标得分", None, "A", f"缺少指标行 {code_i}")

    pollutants, pts = [], {}
    for r, row in enumerate(wb["修复后污染物检测"].iter_rows(min_row=2, values_only=True), 2):
        if not row or all(v is None for v in row[:5]):
            continue
        pcode = str(row[0] or "").strip()
        factor = str(row[2] or "").strip()
        val = _num(row[3] if len(row) > 3 else None)
        unit = str(row[4] or "").strip().lower().replace(" ", "") if len(row) > 4 else ""
        ph = _num(row[1])
        if not pcode:
            err("修复后污染物检测", r, "A", "点位编号为空"); continue
        if not factor:
            err("修复后污染物检测", r, "C", "污染物名称为空"); continue
        if val in (None, "NaN"):
            err("修复后污染物检测", r, "D", f"{pcode}/{factor} 浓度缺失或非数字"); continue
        if val < 0:
            err("修复后污染物检测", r, "D", f"{pcode}/{factor} 浓度为负"); continue
        if unit not in _UNIT_TO_MGKG:
            err("修复后污染物检测", r, "E", f"单位 '{row[4]}' 不受支持(须为 mg/kg)"); continue
        if ph == "NaN" or (ph is not None and not (0 < ph < 14)):
            err("修复后污染物检测", r, "B", "pH 无效"); ph = None
        canon = U.canonical_factor(factor)
        pollutants.append({"point_code": pcode, "pH": ph, "factor": canon, "factor_raw": factor,
                           "value_mgkg": val * _UNIT_TO_MGKG[unit], "row": r})
        p = pts.setdefault(pcode, {"point": pcode, "pH": ph})
        if ph is not None:
            p["pH"] = ph
        p[canon] = max(p.get(canon, float("-inf")), val)
    if not pollutants:
        warnings.append("未提供修复后污染物检测数据: 法规安全门禁将判为'证据不足', 无法给出利用结论")

    calc = None
    if track and t not in (None, "NaN") and M is not None:
        calc = SV.compute(scores, track, t, M)
    sha = hashlib.sha256(content).hexdigest()
    return {"filename": filename, "sha256": sha, "template_version": SV.TEMPLATE_VERSION,
            "track": track, "evaluation_year": int(year) if year not in (None, "NaN") else None,
            "t": t if t not in (None, "NaN") else None, "M": M, "intensity": intensity,
            "data_origin": origin, "data_origin_label": DATA_ORIGINS.get(origin or "", None),
            "records": records, "pollutants": pollutants, "points": list(pts.values()),
            "errors": errors, "warnings": warnings, "n_errors": len(errors),
            "preview_calc": calc, "can_confirm": not errors}


def preview(db: Session, site: Site, content: bytes, filename: str, user_id: int | None,
            expected_track: str | None = None) -> dict:
    rep = parse_and_validate(content, filename, site, expected_track)
    dup = (db.query(SSUIImportBatch).filter_by(site_id=site.id, source_sha256=rep["sha256"], status="confirmed").first())
    if dup:
        rep["errors"].append({"sheet": None, "row": None, "col": None,
                              "message": f"同一文件已于批次 #{dup.id} 确认导入, 不可重复导入"})
        rep["can_confirm"] = False; rep["n_errors"] = len(rep["errors"])
    b = SSUIImportBatch(site_id=site.id, source_file=filename[:300], source_sha256=rep["sha256"],
                        template_version=SV.TEMPLATE_VERSION, stage=POST_REMEDIATION,
                        track=rep["track"] or "unknown", evaluation_year=rep["evaluation_year"],
                        years_since_remediation=rep["t"], multiplier_m=rep["M"],
                        data_origin=rep["data_origin"] or "unknown", status="previewed",
                        row_count=len(rep["records"]) + len(rep["pollutants"]),
                        valid_count=sum(1 for x in rep["records"] if x["score"] is not None) + len(rep["pollutants"]),
                        error_count=rep["n_errors"], validation_report=rep,
                        method_version=SV.METHOD_VERSION, imported_by=user_id)
    db.add(b); db.flush()
    audit_service.log(db, "ssui_post_preview", user_id, "ssui_import_batch", b.id,
                      result="success" if rep["can_confirm"] else "validation_failed",
                      detail={"sha256": rep["sha256"], "errors": rep["n_errors"]}, commit=False)
    db.commit()
    return {"batch_id": b.id, **{k: v for k, v in rep.items() if k not in ("pollutants",)},
            "n_pollutant_rows": len(rep["pollutants"]), "n_points": len(rep["points"])}


def _factor_id(db: Session, canon: str) -> int:
    name = _CN_METAL.get(canon, canon)
    f = db.query(FactorDictionary).filter_by(factor_code=name).first() or \
        db.query(FactorDictionary).filter_by(factor_name=name).first()
    if f is None:
        f = FactorDictionary(factor_code=name, factor_name=name, default_unit="mg/kg",
                             source="课题三修复后导入登记(v1.1)")
        db.add(f); db.flush()
    return f.id


def confirm(db: Session, batch_id: int, user_id: int | None) -> dict:
    b = db.get(SSUIImportBatch, batch_id)
    if b is None:
        raise SSUIImportError("批次不存在")
    if b.status != "previewed":
        raise SSUIImportError(f"批次状态为 {b.status}, 只能确认 previewed 批次")
    rep = b.validation_report or {}
    if not rep.get("can_confirm"):
        raise SSUIImportError(f"校验未通过({rep.get('n_errors')} 处错误), 不可确认")
    if db.query(SSUIImportBatch).filter_by(site_id=b.site_id, source_sha256=b.source_sha256, status="confirmed").first():
        raise SSUIImportError("同一文件已确认导入")
    try:
        for r in rep["records"]:
            db.add(SSUIRecord(batch_id=b.id, site_id=b.site_id, **r))
        mbatch = None
        if rep.get("pollutants"):
            mbatch = ImportBatch(site_id=b.site_id, source_file=b.source_file, source_sha256=b.source_sha256,
                                 row_count=len(rep["pollutants"]), valid_count=len(rep["pollutants"]),
                                 invalid_count=0, status="success", imported_by=user_id,
                                 stage=POST_REMEDIATION, subproject="S3", data_origin=b.data_origin,
                                 method_version=SV.METHOD_VERSION, script_version="ssui_post_v1.1",
                                 data_version=f"post-{b.id}")
            db.add(mbatch); db.flush()
            pid = {}
            for p in rep["pollutants"]:
                code = f"POST{b.id}-{p['point_code']}"[:50]
                if code not in pid:
                    sp = SamplingPoint(site_id=b.site_id, point_code=code)
                    db.add(sp); db.flush(); pid[code] = sp.id
                db.add(Measurement(site_id=b.site_id, sampling_point_id=pid[code], factor_id=_factor_id(db, p["factor"]),
                                   value=p["value_mgkg"], value_used_for_model=p["value_mgkg"], unit="mg/kg",
                                   source_file=b.source_file, import_batch_id=mbatch.id,
                                   original_value_text=str(p["value_mgkg"]), data_origin=b.data_origin,
                                   stage=POST_REMEDIATION, qa_status="raw"))
                if p.get("pH") is not None:
                    pass  # pH 随点位保存在 validation_report.points, 门禁直接读取
        calc = SV.compute({r["indicator_code"]: r["score"] for r in rep["records"]}, b.track,
                          b.years_since_remediation, b.multiplier_m)
        ev = EvaluationResult(site_id=b.site_id, eval_type=f"ssui_post_{b.track}", data_version=f"ssui-batch-{b.id}",
                              param_version=SV.METHOD_VERSION, score=calc.get("ssui"), grade=calc.get("grade"),
                              dimensions={"criterion_scores": calc.get("criterion_scores"), "f_t": calc.get("f_t"),
                                          "M": b.multiplier_m, "t": b.years_since_remediation,
                                          "exceeds_unit_range": calc.get("exceeds_unit_range"),
                                          "feasible": calc.get("feasible"), "status": calc.get("status"),
                                          "support_interpretation": calc.get("support_interpretation"),
                                          "classification_scope": calc.get("classification_scope"),
                                          "validity_domain": calc.get("validity_domain"),
                                          "domain_note": calc.get("domain_note"),
                                          "data_origin": b.data_origin},
                              weights=calc.get("criterion_weights"), limiting_factors={"contributions": calc.get("contributions")},
                              explanation="; ".join(calc.get("warnings", [])),
                              run_config={"batch_id": b.id, "track": b.track, "template": b.template_version},
                              input_fingerprint=b.source_sha256, stage=POST_REMEDIATION, subproject="S3",
                              method_version=SV.METHOD_VERSION, method_status="provisional", source_batch_id=b.id)
        db.add(ev)
        for old in db.query(SSUIImportBatch).filter(SSUIImportBatch.site_id == b.site_id,
                                                    SSUIImportBatch.track == b.track,
                                                    SSUIImportBatch.status == "confirmed").all():
            old.status = "superseded"
        b.status = "confirmed"; b.confirmed_at = datetime.utcnow()
        db.flush()
        audit_service.log(db, "ssui_post_confirm", user_id, "ssui_import_batch", b.id,
                          detail={"ssui": calc.get("ssui"), "grade": calc.get("grade"),
                                  "measurement_batch": mbatch.id if mbatch else None}, commit=False)
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"batch_id": b.id, "status": b.status, "evaluation_id": ev.id, "calc": calc,
            "measurement_batch_id": mbatch.id if mbatch else None}


def batch_points(db: Session, batch: SSUIImportBatch) -> list[dict]:
    return (batch.validation_report or {}).get("points", [])


def export_batch(db: Session, batch_id: int) -> bytes:
    b = db.get(SSUIImportBatch, batch_id)
    if b is None:
        raise SSUIImportError("批次不存在")
    rep = b.validation_report or {}
    ev = (db.query(EvaluationResult).filter_by(source_batch_id=b.id, subproject="S3")
          .order_by(EvaluationResult.id.desc()).first())
    calc = SV.compute({r["indicator_code"]: r["score"] for r in rep.get("records", [])}, b.track,
                      b.years_since_remediation, b.multiplier_m) if b.track in TRACKS else {}
    wb = Workbook(); ws = wb.active; ws.title = "结果"
    origin_label = DATA_ORIGINS.get(b.data_origin, b.data_origin)
    rows = [("场地ID", b.site_id), ("批次", b.id), ("状态", b.status), ("数据阶段", "修复后(课题三)"),
            ("数据来源", origin_label), ("源文件", b.source_file), ("SHA-256", b.source_sha256),
            ("模板版本", b.template_version), ("方法版本", SV.METHOD_VERSION), ("方法状态", "provisional(待课题组确认)"),
            ("评价轨道", TRACKS.get(b.track, b.track)), ("t(年)", b.years_since_remediation), ("M", b.multiplier_m),
            ("f(t)", calc.get("f_t")), ("Σ v_j·S_j", calc.get("weighted_sum")), ("SSUI(未截断)", calc.get("ssui")),
            ("计算状态", {"ok": "域内(暂定演示分级)", "out_of_domain": "超出等级定义域: 不分级、不作支持判断"}.get(calc.get("status"), calc.get("status"))),
            ("等级", calc.get("grade") if calc.get("grade") else "—(不分级)"), ("超出 0–1 区间", calc.get("exceeds_unit_range")),
            ("说明", calc.get("domain_note")),
            ("评价结果ID", ev.id if ev else None), ("导出时间(UTC)", datetime.utcnow().isoformat(timespec="seconds"))]
    if b.data_origin in ("monte_carlo_demo", "test_fixture"):
        rows.insert(0, ("警告", "模拟数据——仅供测试/演示, 不得用于正式报告"))
    for r, (k, v) in enumerate(rows, 1):
        ws.cell(r, 1, k); ws.cell(r, 2, v)
    r0 = len(rows) + 2
    ws.cell(r0, 1, "提示")
    for i, w in enumerate(calc.get("warnings", []), 1):
        ws.cell(r0 + i, 1, w)
    ws2 = wb.create_sheet("指标贡献")
    for c, h in enumerate(["编码", "名称", "准则层", "s_i", "w_ij", "贡献"], 1):
        ws2.cell(1, c, h)
    for r, x in enumerate(calc.get("contributions", []), 2):
        for c, k in enumerate(["code", "name", "criterion", "s", "w", "contribution"], 1):
            ws2.cell(r, c, x[k])
    ws3 = wb.create_sheet("准则层")
    for r, (c, s) in enumerate((calc.get("criterion_scores") or {}).items(), 1):
        ws3.cell(r, 1, c); ws3.cell(r, 2, s); ws3.cell(r, 3, (calc.get("criterion_weights") or {}).get(c))
        ws3.cell(r, 4, (calc.get("group_sums") or {}).get(c))
    ws4 = wb.create_sheet("修复后污染物")
    for c, h in enumerate(["点位", "pH", "污染物", "浓度mg/kg", "源行"], 1):
        ws4.cell(1, c, h)
    for r, p in enumerate(rep.get("pollutants", []), 2):
        for c, k in enumerate(["point_code", "pH", "factor_raw", "value_mgkg", "row"], 1):
            ws4.cell(r, c, p.get(k))
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
