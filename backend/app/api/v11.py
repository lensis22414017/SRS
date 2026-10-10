"""v1.1 API: 课题三修复后 SSUI 独立导入、利用决策、全流程追溯引导与真实进度、模板下载。"""
from __future__ import annotations

import io
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.deps import assert_site_access, get_current_user, require_permission
from app.db.session import get_db
from app.models import (POST_REMEDIATION, PRE_REMEDIATION, DiagnosisResult, EvaluationResult, ImportBatch,
                        ReportRecord, Site, SSUIImportBatch, User, UtilizationDecision, WorkflowRecord)
from app.services import ssui_post_service as SP
from app.services import utilization_service as US
from app.services import workflow_service as WS

router = APIRouter(prefix=get_settings().api_v1_prefix, tags=["v1.1"])
MAX_UPLOAD = 20 * 1024 * 1024
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _site(db: Session, user: User, site_id: int) -> Site:
    s = db.get(Site, site_id)
    if not s:
        raise HTTPException(404, "场地不存在")
    assert_site_access(db, user, s)
    return s


def _xlsx(data: bytes, name: str) -> StreamingResponse:
    return StreamingResponse(io.BytesIO(data), media_type=_XLSX,
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})


# ───────────── 模板下载(C5: 引导上传/下载) ─────────────
PRE_TEMPLATE_COLUMNS = ["采样点编号", "经度", "纬度", "深度_上限(cm)", "深度_下限(cm)", "土壤类型", "pH",
                        "镉_Cd(mg/kg)", "汞_Hg(mg/kg)", "砷_As(mg/kg)", "铅_Pb(mg/kg)", "铬_Cr(mg/kg)",
                        "铜_Cu(mg/kg)", "镍_Ni(mg/kg)", "锌_Zn(mg/kg)", "六价铬_Cr(VI)(mg/kg)",
                        "有机质(g/kg)", "全氮(g/kg)", "阳离子交换量(cmol/kg)", "备注"]


def build_pre_template() -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "检测数据"
    for c, h in enumerate(PRE_TEMPLATE_COLUMNS, 1):
        x = ws.cell(1, c, h); x.font = Font(bold=True, color="FFFFFF"); x.fill = PatternFill("solid", fgColor="2E7D32")
        ws.column_dimensions[x.column_letter].width = max(10, len(h) * 2)
    g = wb.create_sheet("说明")
    for i, t in enumerate([
        "SRS 修复前检测数据导入模板(课题一 障碍因子识别 / 课题二 功能重构可行性)  PRE-v1.1",
        "1. 每行一个采样点; 第一张工作表为数据表, 请勿改名或调整表头。",
        "2. 浓度单位以表头括号为准; 低于检出限填写 '<检出限' (如 <0.01), 系统按规则处理并保留原文。",
        "3. GB 15618 基本项目(镉汞砷铅铬铜镍锌)与 pH 为生产门禁必需; GB 36600 第一类用地门禁另需六价铬及 VOC/SVOC 基本项目, 可追加列。",
        "4. 修复后数据请勿使用本模板, 请在 '课题三 修复后 SSUI 导入' 页面下载专用模板。",
        "5. 模拟/演示数据请在数据来源中如实标注, 系统会在所有结果上显示'模拟数据——仅供测试/演示'。",
    ], 1):
        g.cell(i, 1, t)
    g.column_dimensions["A"].width = 120
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


@router.get("/templates/pre-remediation")
def download_pre_template(user: User = Depends(get_current_user)):
    return _xlsx(build_pre_template(), "SRS_修复前检测数据导入模板_PRE-v1.1.xlsx")


@router.get("/templates/ssui-post")
def download_ssui_template(track: str = Query("production"), site_code: str = Query(""),
                           user: User = Depends(get_current_user)):
    try:
        data = SP.build_template(track, site_code)
    except SP.SSUIImportError as e:
        raise HTTPException(422, str(e))
    return _xlsx(data, f"SRS_课题三修复后SSUI导入模板_{SP.TRACKS[track]}_{SP.SV.TEMPLATE_VERSION}.xlsx")


# ───────────── 课题三 修复后 SSUI(C4) ─────────────
@router.post("/sites/{site_id}/ssui-post/preview")
async def ssui_post_preview(site_id: int, file: UploadFile = File(...), track: str | None = Form(None),
                            user: User = Depends(require_permission("data:input")),
                            db: Session = Depends(get_db)):
    site = _site(db, user, site_id)
    content = await file.read()
    if len(content) > MAX_UPLOAD:
        raise HTTPException(413, "文件超过 20 MB")
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(422, "仅支持 .xlsx 模板文件")
    try:
        return SP.preview(db, site, content, file.filename or "upload.xlsx", user.id, track)
    except SP.SSUIImportError as e:
        raise HTTPException(422, str(e))


@router.post("/ssui-post/batches/{batch_id}/confirm")
def ssui_post_confirm(batch_id: int, user: User = Depends(require_permission("data:input")),
                      db: Session = Depends(get_db)):
    b = db.get(SSUIImportBatch, batch_id)
    if not b:
        raise HTTPException(404, "批次不存在")
    _site(db, user, b.site_id)
    try:
        return SP.confirm(db, batch_id, user.id)
    except SP.SSUIImportError as e:
        raise HTTPException(409, str(e))


@router.post("/ssui-post/batches/{batch_id}/reject")
def ssui_post_reject(batch_id: int, user: User = Depends(require_permission("data:input")),
                     db: Session = Depends(get_db)):
    b = db.get(SSUIImportBatch, batch_id)
    if not b:
        raise HTTPException(404, "批次不存在")
    _site(db, user, b.site_id)
    if b.status != "previewed":
        raise HTTPException(409, f"批次状态为 {b.status}, 不可放弃")
    b.status = "rejected"; db.commit()
    return {"batch_id": b.id, "status": b.status}


@router.get("/sites/{site_id}/ssui-post/batches")
def ssui_post_batches(site_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    _site(db, user, site_id)
    rows = db.query(SSUIImportBatch).filter_by(site_id=site_id).order_by(SSUIImportBatch.id.desc()).all()
    out = []
    for b in rows:
        ev = (db.query(EvaluationResult).filter_by(source_batch_id=b.id, subproject="S3")
              .order_by(EvaluationResult.id.desc()).first())
        out.append({"batch_id": b.id, "status": b.status, "track": b.track, "source_file": b.source_file,
                    "sha256": b.source_sha256, "data_origin": b.data_origin,
                    "data_origin_label": SP.DATA_ORIGINS.get(b.data_origin, b.data_origin),
                    "evaluation_year": b.evaluation_year, "t": b.years_since_remediation, "M": b.multiplier_m,
                    "error_count": b.error_count, "created_at": b.created_at.isoformat() if b.created_at else None,
                    "confirmed_at": b.confirmed_at.isoformat() if b.confirmed_at else None,
                    "ssui": ev.score if ev else None,
                    "grade": (None if SP.effective_status(ev) == "out_of_domain" else ev.grade) if ev else None,
                    "ssui_status": SP.effective_status(ev) if ev else None,
                    "exceeds_unit_range": (ev.dimensions or {}).get("exceeds_unit_range") if ev else None,
                    "warnings": (ev.explanation or "").split("; ") if ev and ev.explanation else []})
    return {"site_id": site_id, "batches": out}


@router.get("/ssui-post/batches/{batch_id}/report")
def ssui_post_report(batch_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    b = db.get(SSUIImportBatch, batch_id)
    if not b:
        raise HTTPException(404, "批次不存在")
    _site(db, user, b.site_id)
    rep = dict(b.validation_report or {})
    rep.pop("pollutants", None)
    return {"batch_id": b.id, "status": b.status, **rep}


@router.get("/ssui-post/batches/{batch_id}/export")
def ssui_post_export(batch_id: int, user: User = Depends(require_permission("file:download")),
                     db: Session = Depends(get_db)):
    b = db.get(SSUIImportBatch, batch_id)
    if not b:
        raise HTTPException(404, "批次不存在")
    _site(db, user, b.site_id)
    return _xlsx(SP.export_batch(db, batch_id), f"SRS_课题三SSUI结果_批次{batch_id}.xlsx")


# ───────────── 利用决策(C3) ─────────────
@router.post("/sites/{site_id}/utilization")
def utilization_run(site_id: int, stage: str = Query(PRE_REMEDIATION),
                    farmland_type: str | None = Query(None, description="水田/其他; 未选择 → 仅保守假设筛查, 无正式生产结论"),
                    eco_land_class: str | None = Query(None, description="第一类用地/第二类用地/非建设用地生态用途; 未选择 → 无正式生态结论"),
                    user: User = Depends(require_permission("data:input")), db: Session = Depends(get_db)):
    _site(db, user, site_id)
    if stage not in (PRE_REMEDIATION, POST_REMEDIATION):
        raise HTTPException(422, "stage 必须为 pre_remediation / post_remediation")
    if farmland_type not in (None, "水田", "其他"):
        raise HTTPException(422, "farmland_type 必须为 水田/其他")
    if eco_land_class not in (None, "第一类用地", "第二类用地", "非建设用地生态用途"):
        raise HTTPException(422, "eco_land_class 必须为 第一类用地/第二类用地/非建设用地生态用途")
    return US.run(db, site_id, stage, farmland_type=farmland_type, eco_land_class=eco_land_class, user_id=user.id)


@router.get("/sites/{site_id}/utilization")
def utilization_latest(site_id: int, stage: str | None = Query(None),
                       user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    _site(db, user, site_id)
    d = US.latest(db, site_id, stage)
    return {"site_id": site_id, "decision": US.to_dict(d) if d else None,
            "empty_reason": None if d else "尚未运行利用决策"}


# ───────────── 全流程追溯(C5) ─────────────
TRACE_GUIDE = {
    "title": "污染场地土壤生态-生产功能重构 全流程追溯",
    "principle": "查看本引导不会创建任何记录; 选择场地并点击'初始化追溯'后才建立五阶段记录。",
    "stages": [
        {"stage": "survey", "name": "调查评估", "data_stage": PRE_REMEDIATION, "subprojects": ["课题一", "课题二"],
         "system_steps": ["下载修复前检测数据模板并导入", "障碍因子识别(KOS)", "下载课题二重构指标模板(28 项)并导入",
                          "功能重构可行性评价(M-REC-2025)", "修复前情景利用判断"],
         "uploads": ["场地调查报告", "采样布点方案", "检测报告(CMA)"], "downloads": ["修复前检测数据模板", "课题二重构指标模板", "课题二结果导出"]},
        {"stage": "approval", "name": "方案审批", "data_stage": PRE_REMEDIATION, "subprojects": ["课题一", "课题二"],
         "system_steps": ["修复技术推荐", "方案比选"], "uploads": ["修复方案", "专家评审意见", "批复文件"], "downloads": []},
        {"stage": "construction", "name": "施工监理", "data_stage": None, "subprojects": [],
         "system_steps": ["记录施工进度与监理意见"], "uploads": ["施工日志", "监理报告", "二次污染监测记录"], "downloads": []},
        {"stage": "effect", "name": "效果评估", "data_stage": POST_REMEDIATION, "subprojects": ["课题三"],
         "system_steps": ["下载课题三修复后 SSUI 模板并导入", "SSUI 计算", "修复后利用决策(生产/生态)"],
         "uploads": ["效果评估报告", "修复后检测报告"], "downloads": ["课题三修复后 SSUI 模板", "SSUI 结果导出"]},
        {"stage": "maintenance", "name": "后期管护", "data_stage": POST_REMEDIATION, "subprojects": ["课题三"],
         "system_steps": ["按年份更新 t 并重新导入 SSUI", "生成全流程报告"],
         "uploads": ["管护记录", "长期监测数据"], "downloads": ["全流程追溯报告"]},
    ],
    "templates": [
        {"name": "修复前检测数据模板", "url": "/api/v1/templates/pre-remediation"},
        {"name": "课题二重构指标模板(28 项, 修复前)", "url": "/api/v1/templates/recon-pre"},
        {"name": "课题三修复后 SSUI 模板(生产)", "url": "/api/v1/templates/ssui-post?track=production"},
        {"name": "课题三修复后 SSUI 模板(生态)", "url": "/api/v1/templates/ssui-post?track=ecology"},
    ],
    "data_origin_rule": "模拟数据仅用于测试/演示, 所有页面与导出均显示'模拟数据——仅供测试/演示'。",
}


@router.get("/trace/guide")
def trace_guide(user: User = Depends(get_current_user)):
    return TRACE_GUIDE


@router.get("/sites/{site_id}/trace/progress")
def trace_progress(site_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """真实进度: 只读查询, 不创建任何记录。"""
    site = _site(db, user, site_id)
    wf = {w["stage"]: w for w in WS.get_stages(db, site_id)}
    milestones = WS.software_milestones(db, site_id)
    nxt = next((m for m in milestones if not m["done"]), None)
    stages = []
    for g in TRACE_GUIDE["stages"]:
        w = wf.get(g["stage"])
        stages.append({"stage": g["stage"], "name": g["name"],
                       "status": w["status"] if w else "not_initialized",
                       "is_completed": bool(w and w["is_completed"]), "n_attachments": w["n_attachments"] if w else 0})
    biz = WS.business_stage_status(db, site_id)
    return {"site_id": site_id, "site_code": site.site_code, "workflow_initialized": bool(wf),
            "stages": stages, "milestones": milestones,
            # v1.2.1(R03): completed/total 只统计七项软件操作里程碑, 不等于五阶段业务完成
            "milestone_kind": "software_operation",
            "milestone_label": "七项软件操作里程碑(不等于五阶段业务完成)",
            "completed": sum(m["done"] for m in milestones), "total": len(milestones),
            "business_stages": biz["stages"], "business_completed": biz["n_completed"],
            "business_with_documents": biz["n_with_documents"], "business_total": len(biz["stages"]),
            "next_step": nxt["name"] if nxt else None}


# ───────── v1.2.1(R03): 评价快照(界面摘要/Excel/报告/比较 JSON 同一来源) ─────────
@router.get("/sites/{site_id}/evaluation-snapshot")
def evaluation_snapshot(site_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    from app.services import evaluation_snapshot as ES
    _site(db, user, site_id)
    snap = ES.build(db, site_id)
    return {"headline": ES.headline(snap), "snapshot": snap}


@router.get("/sites/{site_id}/evaluation-snapshot.xlsx")
def evaluation_snapshot_xlsx(site_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    from app.services import evaluation_snapshot as ES
    site = _site(db, user, site_id)
    snap = ES.build(db, site_id)
    return _xlsx(ES.to_xlsx(snap), f"评价快照_{site.site_code}_{snap['snapshot_id']}.xlsx")


@router.get("/reports/{report_id}/snapshot")
def report_snapshot(report_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """报告生成时保存的快照(不随后续数据变化); verified=重算 SHA-256 与记录一致。"""
    from app.models import ReportRecord
    from app.services import evaluation_snapshot as ES
    rec = db.get(ReportRecord, report_id)
    if rec is None:
        raise HTTPException(404, "报告不存在")
    _site(db, user, rec.site_id)
    ds = rec.data_snapshot or {}
    snap = ds.get("evaluation_snapshot")
    return {"report_id": rec.id, "version": rec.version, "format": ds.get("format"),
            "snapshot_id": ds.get("snapshot_id"), "headline": ds.get("headline"),
            "verified": bool(snap) and ES.verify(snap), "snapshot": snap, "pdf_font": ds.get("pdf_font")}
