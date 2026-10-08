"""SRS v1.1 蒙特卡洛演示数据包 — 模拟数据——仅供测试/演示。

用法(均需显式指定独立演示库, 拒绝写入默认用户库):
  python scripts/mc_demo_v11.py generate --out demo/mc_v11
  python scripts/mc_demo_v11.py run      --out demo/mc_v11 --db demo/mc_v11/srs_demo.db
  python scripts/mc_demo_v11.py clean    --db demo/mc_v11/srs_demo.db

generate: 固定随机种子生成 修复前检测表(PRE-v1.1 格式) + 课题三修复后 SSUI 表(生产/生态, SSUI-POST-v1.1) + metadata.json
run:      在独立演示库中导入并依次运行 课题一 KOS → 课题二 重构可行性 → 修复前情景判断 → 课题三 SSUI(两轨) → 修复后利用结论,
          输出 demo_evidence.json
clean:    删除演示库中所有 site_code 以 MCDEMO 开头的场地及其关联记录
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import sys
from datetime import datetime, timezone

import numpy as np

LABEL = "模拟数据——仅供测试/演示"
FILE_LABEL = "模拟数据_仅供测试演示"  # 文件名安全版本(不含 /)
SEED = 20261008
GEN_VERSION = "mc_demo_v11.0"
SITE_CODE = "MCDEMO"
SITE_NAME = f"【{LABEL}】蒙特卡洛演示场地 · 复合重金属农用地"
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# 修复前: 对数正态(中位数, 几何标准差) — 设计为 Cd/Pb/As 超筛选值、低于管制值, 体现"安全利用类"情景
PRE_DIST = {
    "镉_Cd(mg/kg)": (0.9, 1.5), "汞_Hg(mg/kg)": (0.25, 1.4), "砷_As(mg/kg)": (32, 1.3), "铅_Pb(mg/kg)": (140, 1.4),
    "铬_Cr(mg/kg)": (85, 1.2), "铜_Cu(mg/kg)": (42, 1.3), "镍_Ni(mg/kg)": (38, 1.2), "锌_Zn(mg/kg)": (150, 1.3),
    "六价铬_Cr(VI)(mg/kg)": (0.8, 1.5), "有机质(g/kg)": (18, 1.3), "全氮(g/kg)": (1.1, 1.3), "阳离子交换量(cmol/kg)": (14, 1.2),
}
PH = (6.3, 0.35)                      # 正态
REMOVAL = (7.0, 3.0)                  # 修复去除率 Beta(a,b), 均值 0.70
SCORE_BETA = {"production": (7.0, 3.0), "ecology": (6.0, 4.0)}   # s_i ~ Beta, 均值 0.70 / 0.60
N_PRE, N_POST = 24, 12
T_YEARS, INTENSITY = 3, "中等强度"
CN = {"镉_Cd(mg/kg)": "镉", "汞_Hg(mg/kg)": "汞", "砷_As(mg/kg)": "砷", "铅_Pb(mg/kg)": "铅", "铬_Cr(mg/kg)": "铬",
      "铜_Cu(mg/kg)": "铜", "镍_Ni(mg/kg)": "镍", "锌_Zn(mg/kg)": "锌", "六价铬_Cr(VI)(mg/kg)": "六价铬"}


def _sha(p):
    return hashlib.sha256(open(p, "rb").read()).hexdigest()


def generate(out: str) -> dict:
    sys.path.insert(0, os.path.join(ROOT, "backend")); sys.path.insert(0, os.path.join(ROOT, "ml", "evaluation"))
    os.environ.setdefault("DATABASE_URL", "sqlite:///" + os.path.join(out, "_gen_unused.db"))
    from openpyxl import load_workbook
    from app.api.v11 import PRE_TEMPLATE_COLUMNS, build_pre_template
    from app.services.ssui_post_service import build_template
    os.makedirs(out, exist_ok=True)
    rng = np.random.default_rng(SEED)
    # ── 修复前 ──
    wb = load_workbook(io.BytesIO(build_pre_template()))
    ws = wb["检测数据"]
    pre_rows = []
    for i in range(N_PRE):
        row = {"采样点编号": f"MC{i + 1:02d}", "经度": round(120.10 + rng.uniform(0, 0.02), 6),
               "纬度": round(30.20 + rng.uniform(0, 0.02), 6), "深度_上限(cm)": 0, "深度_下限(cm)": 20,
               "土壤类型": "水稻土", "pH": round(float(np.clip(rng.normal(*PH), 4.5, 8.5)), 2), "备注": LABEL}
        for col, (med, gsd) in PRE_DIST.items():
            row[col] = round(float(med * np.exp(rng.normal(0, np.log(gsd)))), 4)
        pre_rows.append(row)
        for c, h in enumerate(PRE_TEMPLATE_COLUMNS, 1):
            ws.cell(i + 2, c, row.get(h))
    wb["说明"].cell(8, 1, f"本文件为{LABEL}; 生成器 {GEN_VERSION}, 种子 {SEED}")
    pre_path = os.path.join(out, f"MCDEMO_修复前检测数据_{FILE_LABEL}.xlsx")
    wb.save(pre_path)
    # ── 修复后(课题三) ──
    post_paths = {}
    removal = rng.beta(*REMOVAL, size=N_POST)
    post_pts = []
    for j in range(N_POST):
        src = pre_rows[j]
        pt = {"point": f"MC{j + 1:02d}", "pH": round(src["pH"] + float(rng.normal(0.2, 0.1)), 2)}
        for col in CN:
            pt[CN[col]] = round(src[col] * (1 - removal[j]), 4)
        post_pts.append(pt)
    for track in ("production", "ecology"):
        wb = load_workbook(io.BytesIO(build_template(track, SITE_CODE)))
        m = wb["批次信息"]
        m["B4"] = 2026; m["B5"] = T_YEARS; m["B6"] = INTENSITY; m["B8"] = LABEL
        m["B9"] = f"{GEN_VERSION} seed={SEED}"
        s = wb["指标得分"]
        a, b = SCORE_BETA[track]
        for r in range(2, 27):
            s.cell(r, 7, round(float(rng.beta(a, b)), 4)); s.cell(r, 8, LABEL)
        p = wb["修复后污染物检测"]; p.delete_rows(2, 1)
        r = 2
        for pt in post_pts:
            for name in CN.values():
                p.cell(r, 1, pt["point"]); p.cell(r, 2, pt["pH"]); p.cell(r, 3, name); p.cell(r, 4, pt[name])
                p.cell(r, 5, "mg/kg"); p.cell(r, 7, LABEL); r += 1
        path = os.path.join(out, f"MCDEMO_课题三修复后SSUI_{'生产' if track == 'production' else '生态'}_{FILE_LABEL}.xlsx")
        wb.save(path); post_paths[track] = path
    meta = {"label": LABEL, "generator": GEN_VERSION, "seed": SEED,
            "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "site_code": SITE_CODE, "site_name": SITE_NAME,
            "design": {"n_pre_points": N_PRE, "n_post_points": N_POST, "pre_lognormal_median_gsd": PRE_DIST,
                       "pH_normal_mean_sd": PH, "removal_beta_ab": REMOVAL, "ssui_score_beta_ab": SCORE_BETA,
                       "t_years": T_YEARS, "intensity": INTENSITY},
            "intended_scenario": "修复前 Cd/Pb/As 超 GB15618 筛选值但低于管制值(安全利用类); 修复后浓度按 Beta 去除率降低",
            "files": {os.path.basename(x): _sha(x) for x in [pre_path, *post_paths.values()]},
            "usage_rule": "仅用于测试/演示; 只能导入独立演示库; 所有结果显示模拟数据标签; 不得用于正式报告"}
    json.dump(meta, open(os.path.join(out, "metadata.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return meta


def _guard_db(db: str):
    p = os.path.abspath(db)
    bad = ("Application Support/SRS", "AppData", "/srs.db")
    if any(b in p for b in bad) or not os.path.basename(p).startswith("srs_demo"):
        raise SystemExit(f"拒绝写入: 演示库文件名须以 srs_demo 开头且不得为用户默认库 ({p})")
    return p


def _boot(db: str):
    os.environ["DATABASE_URL"] = "sqlite:///" + _guard_db(db)
    os.environ.setdefault("SRS_FIRST_ADMIN_PASSWORD", "Demo@2026x")
    for p in (os.path.join(ROOT, "backend"), os.path.join(ROOT, "ml", "evaluation"), os.path.join(ROOT, "ml", "models"),
              os.path.join(ROOT, "ml", "explain"), os.path.join(ROOT, "ml", "recommend")):
        if p not in sys.path:
            sys.path.insert(0, p)
    from app.db.init_db import create_all
    from app.db.seed_db import seed_if_empty
    create_all(); seed_if_empty()


def run(out: str, db: str) -> dict:
    _boot(db)
    from app.db.session import SessionLocal
    from app.models import ImportBatch, Measurement, Site, User
    from app.services import ssui_post_service as SP
    from app.services import utilization_service as US
    from app.services.evaluation_service import run_evaluation
    from app.services.import_service import smart_detect_and_map
    from app.services.pipeline import run_import_with_mapping
    meta = json.load(open(os.path.join(out, "metadata.json"), encoding="utf-8"))
    for name, sha in meta["files"].items():
        assert _sha(os.path.join(out, name)) == sha, f"演示文件被修改: {name}"
    db_ = SessionLocal()
    ev: dict = {"label": LABEL, "generator": meta["generator"], "seed": meta["seed"], "db": os.path.abspath(db)}
    try:
        admin = db_.query(User).filter_by(username="admin").first()
        pre = [n for n in meta["files"] if "修复前" in n][0]
        _, mapping, _ = smart_detect_and_map(os.path.join(out, pre))
        imp = run_import_with_mapping(db_, os.path.join(out, pre), mapping, imported_by=admin.id if admin else None,
                                      on_conflict="skip")
        site = db_.get(Site, imp["site_id"])
        site.site_code = SITE_CODE; site.name = SITE_NAME; site.land_use_type = "耕地"
        site.description = f"{LABEL}; {meta['generator']} seed={meta['seed']}"
        for b in db_.query(ImportBatch).filter_by(site_id=site.id).all():
            b.data_origin = "monte_carlo_demo"; b.subproject = "S1/S2"
        db_.query(Measurement).filter_by(site_id=site.id).update({"data_origin": "monte_carlo_demo"})
        db_.commit()
        ev["site"] = {"id": site.id, "code": site.site_code, "n_measurements":
                      db_.query(Measurement).filter_by(site_id=site.id).count()}
        # 课题一 KOS
        from app.services.kos_service import run_kos_diagnosis
        from app.services.recommend_service import run_recommendation
        rec = run_recommendation(db_, site.id, top_k=3)
        ev["S1_kos"] = {"upstream": rec.get("upstream_status"), "factors": [r.get("matched_factors") for r in rec.get("recommendations", [])][:3],
                        "recommendation_type": rec.get("recommendation_type")}
        # 课题二 重构可行性
        e = run_evaluation(db_, site.id)
        ev["S2_reconstruction"] = {k: {"score": e[k].get("score"), "grade": e[k].get("grade")}
                                   for k in ("reconstruction_prod", "reconstruction_eco")}
        d_pre = US.run(db_, site.id, "pre_remediation", farmland_type="水田", user_id=admin.id if admin else None)
        ev["decision_pre"] = {k: d_pre[k] for k in ("decision_state", "conclusion_text", "remediation_targets", "data_origin")}
        ev["decision_pre"]["gates"] = {t: d_pre[t]["gate"]["state"] for t in ("production", "ecology")}
        # 课题三
        ev["S3_ssui"] = {}
        for track in ("production", "ecology"):
            name = [n for n in meta["files"] if "课题三" in n and ("生产" if track == "production" else "生态") in n][0]
            content = open(os.path.join(out, name), "rb").read()
            pv = SP.preview(db_, site, content, name, admin.id if admin else None, track)
            assert pv["can_confirm"], pv["errors"]
            cf = SP.confirm(db_, pv["batch_id"], admin.id if admin else None)
            c = cf["calc"]
            ev["S3_ssui"][track] = {"batch_id": pv["batch_id"], "ssui": c["ssui"], "grade": c["grade"],
                                    "exceeds_unit_range": c["exceeds_unit_range"], "feasible": c["feasible"],
                                    "criterion_scores": c["criterion_scores"], "warnings": c["warnings"]}
            export = SP.export_batch(db_, pv["batch_id"])
            xp = os.path.join(out, f"export_SSUI_{track}_batch{pv['batch_id']}.xlsx")
            open(xp, "wb").write(export); ev["S3_ssui"][track]["export"] = os.path.basename(xp)
        d_post = US.run(db_, site.id, "post_remediation", farmland_type="水田", user_id=admin.id if admin else None)
        ev["decision_post"] = {k: d_post[k] for k in ("decision_state", "conclusion_text", "missing_evidence", "data_origin")}
        ev["decision_post"]["gates"] = {t: d_post[t]["gate"]["state"] for t in ("production", "ecology")}
        ev["decision_post"]["track_status"] = {t: d_post[t]["track_status"] for t in ("production", "ecology")}
    finally:
        db_.close()
    json.dump(ev, open(os.path.join(out, "demo_evidence.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)
    return ev


def clean(db: str) -> dict:
    _boot(db)
    from sqlalchemy import text
    from app.db.session import SessionLocal
    from app.models import Site
    db_ = SessionLocal()
    removed = []
    try:
        ids = [s.id for s in db_.query(Site).filter(Site.site_code.like(f"{SITE_CODE}%")).all()]
        child = ["utilization_decisions", "ssui_records", "ssui_import_batches", "evaluation_results", "recommendations",
                 "diagnosis_factor_details", "diagnosis_results", "measurements", "sampling_points", "import_batches",
                 "dataset_versions", "workflow_records", "report_records"]
        for sid in ids:
            for t in child:
                try:
                    if t == "diagnosis_factor_details":
                        db_.execute(text("DELETE FROM diagnosis_factor_details WHERE diagnosis_id IN "
                                         "(SELECT id FROM diagnosis_results WHERE site_id=:s)"), {"s": sid})
                    else:
                        db_.execute(text(f"DELETE FROM {t} WHERE site_id=:s"), {"s": sid})
                except Exception:  # noqa: BLE001 — 表/列不存在时跳过
                    db_.rollback()
            db_.execute(text("DELETE FROM sites WHERE id=:s"), {"s": sid}); removed.append(sid)
        db_.commit()
    finally:
        db_.close()
    return {"removed_site_ids": removed}


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["generate", "run", "clean"])
    ap.add_argument("--out", default=os.path.join(ROOT, "demo", "mc_v11"))
    ap.add_argument("--db")
    a = ap.parse_args()
    if a.cmd == "generate":
        print(json.dumps(generate(a.out)["files"], ensure_ascii=False, indent=1))
    elif a.cmd == "run":
        if not a.db:
            raise SystemExit("run 需要 --db 指定独立演示库")
        r = run(a.out, a.db)
        print(json.dumps({k: r[k] for k in ("decision_pre", "decision_post")}, ensure_ascii=False, indent=1, default=str))
    else:
        if not a.db:
            raise SystemExit("clean 需要 --db")
        print(clean(a.db))
