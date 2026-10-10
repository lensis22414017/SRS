"""v1.2.1(R03/R09): 跨渠道一致性断言 — 评价快照 API = 报告记录 = Excel = PDF 文本 = DOCX 文本。

不导入 app 包(Windows 验收时在 runner 的独立 Python 中运行, 只经 HTTP 取得文件)。
所有比较都基于 /api/v1/sites/{id}/evaluation-snapshot 的 headline; 文本渠道去掉空白后做包含判断。
"""
from __future__ import annotations

import io
import re
import zipfile

GATE_STATE_CN = {"pass": "未超筛选值", "conditional": "超筛选值、未超管制值(需风险评估/管控)",
                 "fail": "超管制值", "insufficient": "证据不足(缺测/条件不明)"}
DECISION_CN = {"both_supported": "生产与生态均支持", "production_supported": "支持生产利用",
               "ecology_supported": "支持生态利用", "neither_supported": "均不支持",
               "insufficient_evidence": "证据不足"}
OFFICIAL_STATUS_CN = {"available": "正式结果可用", "partial": "正式因子不足 3 个(部分结果)",
                      "insufficient_evidence": "证据不足, 无正式排名"}
FERTILITY = {"CEC_cmolkg", "TN_gkg", "OM_gkg", "OC_pct", "Total_P_gkg", "Total_K_gkg", "P_mgkg", "K_mgkg",
             "Hydrolyzable_N_mgkg"}


def fmt(v, nd=4) -> str:
    if v is None:
        return "—"
    if isinstance(v, float):
        s = f"{v:.{nd}f}".rstrip("0").rstrip(".")
        return s or "0"
    return str(v)


def _squash(t: str) -> str:
    return re.sub(r"\s+", "", t or "")


def pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)


def pdf_info(data: bytes) -> dict:
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    fonts, embedded = set(), True
    for p in r.pages:
        res = p.get("/Resources") or {}
        for _, f in ((res.get("/Font") or {}).items() if res.get("/Font") else []):
            f = f.get_object()
            fonts.add(str(f.get("/BaseFont")))
            desc = f.get("/FontDescriptor")
            if f.get("/Subtype") == "/Type0":
                d = f["/DescendantFonts"][0].get_object()
                desc = d.get("/FontDescriptor")
            if desc is not None:
                desc = desc.get_object()
                if not any(k in desc for k in ("/FontFile", "/FontFile2", "/FontFile3")):
                    embedded = False
            elif "Helvetica" not in str(f.get("/BaseFont")):
                embedded = False
    return {"pages": len(r.pages), "fonts": sorted(fonts), "all_cjk_fonts_embedded": embedded,
            "n_images": sum(len(getattr(p, "images", []) or []) for p in r.pages)}


def docx_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
        xml += z.read("word/footer1.xml").decode("utf-8", errors="ignore") if "word/footer1.xml" in z.namelist() else ""
    xml = re.sub(r"</w:p>", "\n", xml)
    return re.sub(r"<[^>]+>", "", xml)


def docx_media_count(data: bytes) -> int:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return sum(1 for n in z.namelist() if n.startswith("word/media/"))


def xlsx_headline(data: bytes) -> dict:
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(data), read_only=True)["首页关键数"]
    return {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]}


def expected_strings(h: dict) -> list[tuple[str, str]]:
    """报告首页摘要必须出现的字符串(与 report_document 摘要措辞对应)。"""
    out = [("snapshot_id", h["snapshot_id"]),
           ("修复前独立样品数", f"{h['pre_n_samples']} 个"),
           ("修复前检测记录数", f"{h['pre_n_measurements']} 条检测记录")]
    if h.get("post_n_batches"):
        out.append(("修复后独立样品数", f"{h['post_n_samples']} 个独立样品"))
    for k, lbl in (("pre_gate_production", "修复前生产门禁"), ("pre_gate_ecology", "修复前生态门禁"),
                   ("post_gate_production", "修复后生产门禁"), ("post_gate_ecology", "修复后生态门禁")):
        if h.get(k):
            out.append((lbl, GATE_STATE_CN[h[k]]))
    if h.get("kos_official_status"):
        out.append(("KOS 正式状态", OFFICIAL_STATUS_CN[h["kos_official_status"]]))
    for f in h.get("kos_official") or []:
        out.append((f"KOS 正式因子 {f}", f))
    for t, cn in (("production", "生产"), ("ecology", "生态")):
        v = h.get(f"ssui_post_{t}")
        if v is not None:
            out.append((f"修复后 SSUI {cn}", f"{cn} {fmt(v)}"))
    for k, lbl in (("decision_pre", "修复前情景判断"), ("decision_post", "修复后利用结论")):
        if h.get(k):
            out.append((lbl, DECISION_CN[h[k]]))
    out.append(("五阶段业务记录", f"{h['business_stages_completed']}/{h['business_stages_total']} 阶段已完成"))
    out.append(("软件操作里程碑", f"{h['software_milestones_completed']}/{h['software_milestones_total']}"))
    if h.get("is_synthetic"):
        out.append(("模拟数据标签", "模拟数据——仅供测试/演示"))
    return out


def forbidden_strings(h: dict) -> list[str]:
    bad = ["heavy_metal", "真实进度"]
    if h.get("is_synthetic"):
        bad += ["真实采样点", "实测因子"]
    return bad


def check_channels(api_headline: dict, report_snapshots: list[dict], xlsx: bytes | None,
                   pdf: bytes | None, docx: bytes | None) -> list[dict]:
    """返回 [{check, passed, detail}]。report_snapshots = GET /reports/{id}/snapshot 的返回列表。"""
    res = []

    def chk(name, cond, detail=None):
        res.append({"check": name, "passed": bool(cond), "detail": detail})
    h = api_headline
    for rs in report_snapshots:
        chk(f"报告 {rs.get('format')} 快照可校验(SHA-256)", rs.get("verified") is True, rs.get("snapshot_id"))
        chk(f"报告 {rs.get('format')} 快照 = 当前 API 快照", rs.get("headline") == h,
            None if rs.get("headline") == h else {k: (rs.get("headline") or {}).get(k) for k in h if (rs.get("headline") or {}).get(k) != h[k]})
    if len({rs.get("snapshot_id") for rs in report_snapshots}) > 1:
        chk("PDF 与 DOCX 使用同一快照", False, [rs.get("snapshot_id") for rs in report_snapshots])
    elif report_snapshots:
        chk("PDF 与 DOCX 使用同一快照", True, report_snapshots[0].get("snapshot_id"))
    if xlsx is not None:
        xh = xlsx_headline(xlsx)
        diff = {k: (xh.get(k), v) for k, v in h.items()
                if str(xh.get(k)) != ("、".join(map(str, v)) if isinstance(v, list) else str(v))
                and not (xh.get(k) is None and v in (None, [], ""))}
        chk("Excel 首页关键数 = API 快照", not diff, diff or None)
    chk("正式 Top-N 不含肥力/下限指标", not (set(h.get("kos_official") or []) & FERTILITY), h.get("kos_official"))
    for name, data, fn in (("PDF", pdf, pdf_text), ("DOCX", docx, docx_text)):
        if data is None:
            continue
        txt = _squash(fn(data))
        miss = [lbl for lbl, s in expected_strings(h) if _squash(s) not in txt]
        chk(f"{name} 文本含首页关键数", not miss, miss or None)
        bad = [s for s in forbidden_strings(h) if _squash(s) in txt]
        chk(f"{name} 无误导用语", not bad, bad or None)
    if pdf is not None:
        info = pdf_info(pdf)
        chk("PDF 中文字体已嵌入", info["all_cjk_fonts_embedded"], info["fonts"])
        chk("PDF 含图件", info["n_images"] >= 2, info["n_images"])
        chk("PDF 非纯文本降级(≥5 页)", info["pages"] >= 5, info["pages"])
    if docx is not None:
        chk("DOCX 含图件", docx_media_count(docx) >= 2, docx_media_count(docx))
    return res
