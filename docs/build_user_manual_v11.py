"""SRS v1.1.0 用户操作手册生成器: 同一内容源 → 可编辑 DOCX + PDF。

截图从 --shots 目录读取(capture_v110_screenshots.command 或 Windows CI 生成);
缺失的截图以醒目的"截图位"占位, 绝不复用旧版截图。
用法: python docs/build_user_manual_v11.py --shots <dir> --out <dir>
"""
from __future__ import annotations

import argparse
import html as H
import json
import os
from datetime import date

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
VERSION = open(os.path.join(ROOT, "VERSION"), encoding="utf-8").read().strip()
TITLE = "污染场地土壤生态-生产功能重构监管系统（SRS）"
SUB = f"用户操作手册 v{VERSION}"
FONT_DIR = "/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts"
LABEL = "模拟数据——仅供测试/演示"


def _weights():
    return json.load(open(os.path.join(ROOT, "data", "standards", "ssui_weights_pptx_v1.json"), encoding="utf-8"))


def content():
    W = _weights()
    C = []  # (kind, payload)
    h1 = lambda t: C.append(("h1", t)); h2 = lambda t: C.append(("h2", t))
    p = lambda t: C.append(("p", t)); ul = lambda *xs: C.append(("ul", list(xs)))
    ol = lambda *xs: C.append(("ol", list(xs))); tb = lambda hdr, rows: C.append(("table", (hdr, rows)))
    note = lambda t: C.append(("note", t)); shot = lambda f, cap: C.append(("shot", (f, cap)))

    h1("第一章 系统概述")
    h2("1.1 系统定位")
    p("SRS 面向污染场地修复与再利用监管，支持从修复前调查数据到修复后可持续利用评价的全过程记录、计算与追溯。"
      "系统由生态环境部土壤与农业农村生态环境监管技术中心委托，浙江大学环境与资源学院研制。")
    h2("1.2 三个课题与数据阶段")
    p("v1.1 起，系统按课题和数据阶段分开处理数据。修复前数据只进入课题一、课题二的计算；修复后数据只能通过课题三的独立导入功能进入系统。两类数据在数据库中物理分离，不会互相替代。")
    tb(["课题", "功能", "数据阶段", "入口", "主要输出"], [
        ["课题一", "障碍因子识别（KOS）", "修复前", "场地管理 → 数据导入；障碍因子分析（课题一）", "关键障碍因子 Top-N、补测建议"],
        ["课题二", "功能重构可行性", "修复前", "功能重构分析（课题二）", "生产/生态重构可行性；修复前情景判断"],
        ["课题三", "修复后土壤持续利用度（SSUI）", "修复后", "修复后SSUI（课题三）", "SSUI 值与等级；修复后利用方向结论"],
    ])
    note("课题划分依据 7 月 30 日李老师意见；课题二、课题三的方法与数据结构以陈亮、伟杰老师最终确认为准，系统中相关结果均标注“方法状态：暂定”。")
    h2("1.3 v1.1 主要变化（对应甲方意见 C1–C6）")
    tb(["编号", "甲方意见", "v1.1 实现"], [
        ["C1", "与陈亮/伟杰对齐方法粒度", "方法冲突逐项列入待确认清单；涉及结果标注“暂定”，不静默修正"],
        ["C2", "三个课题的数据都要演示", "提供带标签的蒙特卡洛演示数据包，覆盖课题一→二→三全流程"],
        ["C3", "修复后给出“宜生产/宜生态”的具体结论", "新增利用方向判定：法规安全门禁（GB 15618 / GB 36600）＋功能评分，共五类结论"],
        ["C4", "课题一/二用修复前数据，课题三用修复后数据并单独导入", "新增课题三修复后 SSUI 独立模板、校验预览、确认入库、计算与导出"],
        ["C5", "一进入就展示全流程追溯并引导上传/下载", "全流程追溯页面打开即显示五阶段引导与模板下载；选择场地后显示真实进度"],
        ["C6", "提供用户手册和 PPT", "本手册（DOCX/PDF）与操作演示 PPT"],
    ])
    h2("1.4 用户角色与权限")
    tb(["角色", "可进行的操作"], [
        ["系统管理员（admin）", "全部功能，含用户/角色管理、参数设置、备份恢复、技术库管理"],
        ["企业用户（enterprise）", "本企业场地的数据录入、查询、导出、报告生成、地图与追溯查看"],
        ["第三方机构（agency）", "获授权场地的数据录入、查询、追溯查看、文档下载"],
        ["监管人员（regulator）", "全部场地只读查询、地图、追溯、报告生成、日志审计；不能导入数据或运行利用判定"],
    ])

    h1("第二章 安装与首次启动")
    h2("2.1 运行环境")
    ul("操作系统：Windows 10/11 64 位", "内存 8 GB 及以上；可用磁盘 3 GB 及以上",
       "浏览器：系统启动后自动打开默认浏览器（推荐 Edge 或 Chrome）", "网络：核心功能可离线使用；在线底图与 AI 问答需联网")
    h2("2.2 安装")
    ol(f"双击 SRS-Setup-{VERSION}-Windows-x64.exe，按向导完成安装。安装路径可以包含中文和空格。",
       "无管理员权限时，可在命令行加 /CURRENTUSER 参数安装到当前用户目录。",
       "安装完成后，从开始菜单或桌面快捷方式启动“SRS”。")
    h2("2.3 首次启动")
    p("首次启动时数据库为空，系统进入“首启设置向导”。按提示设置管理员用户名和密码（至少 8 位，且包含大写字母、小写字母、数字、特殊字符中的至少 3 类），完成后用该账户登录。安装包不包含任何业务数据或测试账户。")
    shot("01_login.png", "登录页面")
    h2("2.4 数据存储位置")
    tb(["内容", "位置（Windows）"], [
        ["数据库", r"%APPDATA%\SRS\srs.db"], ["上传文件与报告", r"%APPDATA%\SRS\storage"],
        ["备份", r"%APPDATA%\SRS\backups"], ["启动日志", r"%APPDATA%\SRS\launcher.log"]])
    h2("2.5 从 v1.0.x 升级")
    p("直接安装 v1.1.0 即可。首次启动时，系统自动为旧数据库补充 v1.1 新增的字段和数据表（只新增，不删除、不修改已有数据）。旧数据一律视为“修复前”数据。建议升级前先在“系统管理 → 备份恢复”中创建一次备份。")
    h2("2.6 卸载")
    p("通过开始菜单“卸载 SRS”或系统“应用和功能”卸载。卸载不删除 %APPDATA%\\SRS 下的数据库和备份；需要彻底清除时请手动删除该目录。")

    h1("第三章 全流程追溯：进入即见的引导")
    p("点击左侧“全流程追溯”，页面顶部立即显示五阶段引导：每个阶段需要上传的资料、可下载的模板，以及对应的课题和数据阶段。查看引导不会创建任何记录。")
    shot("09_trace_guide.png", "全流程追溯引导（未选择场地时即显示）")
    tb(["阶段", "数据阶段 / 课题", "系统操作", "需上传", "可下载"], [
        ["调查评估", "修复前 / 课题一、二", "导入修复前数据 → KOS → 重构可行性 → 情景判断", "调查报告、布点方案、检测报告", "修复前检测数据模板"],
        ["方案审批", "修复前 / 课题一、二", "技术推荐、方案比选", "修复方案、专家意见、批复", "—"],
        ["施工监理", "—", "记录进度与监理意见", "施工日志、监理报告、二次污染监测", "—"],
        ["效果评估", "修复后 / 课题三", "导入修复后 SSUI → 计算 → 利用结论", "效果评估报告、修复后检测报告", "课题三模板、SSUI 结果"],
        ["后期管护", "修复后 / 课题三", "按年份更新 t 重新导入；生成全流程报告", "管护记录、长期监测数据", "全流程追溯报告"]])
    p("在列表中点击“进入追溯”选择场地后，页面顶部显示该场地的真实进度（7 个里程碑：修复前数据导入、障碍因子识别、重构可行性、修复前情景判断、修复后 SSUI 导入、修复后利用结论、全流程报告），并提示下一步。"
      "五阶段记录需点击“初始化五阶段”后才会建立；之后可按阶段上传附件、提交审核或退回。")
    shot("10_trace_detail.png", "场地追溯详情与真实进度")

    h1("第四章 修复前数据准备与导入（课题一、课题二）")
    h2("4.1 下载模板")
    p("在全流程追溯引导或“场地管理 → 数据导入”中下载《修复前检测数据导入模板（PRE-v1.1）》。第一张工作表为数据表，每行一个采样点。")
    h2("4.2 填写规则")
    ul("必填：采样点编号；建议填写经纬度、采样深度、土壤类型、pH。",
       "浓度单位以表头括号为准（如“铅_Pb(mg/kg)”）。低于检出限填写“<检出限”（如 <0.01），系统保留原文并按规则处理。",
       "生产门禁需要 GB 15618 基本项目（镉、汞、砷、铅、铬、铜、镍、锌）和 pH；生态门禁需要 GB 36600 表 1 的 45 项基本项目（含六价铬及 VOC/SVOC），可在模板后追加列。",
       f"演示/模拟数据须在备注列或说明页注明“{LABEL}”。系统识别到该标签后，会把该场地及其所有结果标为模拟数据。",
       "修复后数据不要使用本模板，请使用课题三模板（第七章）。")
    h2("4.3 导入与校验")
    ol("场地管理 → 导入数据 → 选择文件 → 系统自动识别字段（可在向导中调整映射）。",
       "查看校验报告：错误、警告、超标因子统计。同一文件重复导入时按“跳过 / 覆盖 / 新版本”策略处理（按内容指纹判重）。",
       "导入完成后，在场地详情中查看点位地图、检测数据和 EDA 分析。")
    shot("04_import.png", "数据导入页面")
    shot("03_sites.png", "场地管理列表")

    h1("第五章 课题一：障碍因子识别（KOS）")
    ol("左侧点击“障碍因子分析（课题一）”，选择场地和评价轨道（生产 / 生态）。", "点击“运行诊断”。系统使用修复前检测数据和 p3_alpha 模型计算。",
       "查看结果：法规明确障碍（B=1）、关键障碍因子 Top-N（KOS 综合评分）、五分量证据堆叠条和补测建议。")
    p("KOS = B ×（0.30R + 0.25W + 0.15M + 0.20S + 0.10E），只在规则判定 B = 1 时计算。R 为规则严重度，W 为用途权重，M 为模型贡献度，S 为稳定性，E 为证据强度。")
    shot("05_obstacle_S1.png", "障碍因子分析（课题一）")

    h1("第六章 课题二：功能重构可行性与修复前情景判断")
    ol("左侧点击“功能重构分析（课题二）”，选择场地并点击“运行”。",
       "查看生产、生态两轨结果：综合得分大于 50 为“可行”，不大于 50 为“不可行”；指标覆盖率低于门禁时显示“证据不足/无法评价”，并列出缺失指标。",
       "页面下方的“修复前情景利用判断”卡片中，可选择农用地类型和生态用地类别，然后点击“运行判定”。")
    note("修复前情景判断只说明“以当前（修复前）状态，哪些因子必须修复到什么水平”，结论前缀为【修复前情景判断，非修复后利用结论】，并给出修复目标清单。修复完成后，必须用修复后数据（课题三）重新判定。")
    shot("06_reconstruction_S2_pre_decision.png", "功能重构分析与修复前情景判断（课题二）")

    h1("第七章 课题三：修复后 SSUI 独立导入")
    h2("7.1 流程")
    ol("左侧点击“修复后SSUI（课题三）”，选择场地和评价轨道（生产利用 / 生态利用）。",
       "点击“下载模板”。模板已预填场地编号和 25 项指标的编码、名称与权重。",
       "按 7.2 填写后点击“上传已填写的模板”。此步骤只做校验预览，不写入业务数据。",
       "检查预览：错误逐条列出“工作表 / 行 / 列 / 问题”，有错误时不能确认，修正后重新上传。",
       "校验通过后点击“确认入库”。系统在一个事务中写入指标记录、修复后污染物检测值和 SSUI 结果，任一步失败则全部回滚。",
       "在“已导入批次”中可导出 SSUI 结果（xlsx：结果、指标贡献、准则层、修复后污染物）。同一场地同一轨道再次确认新批次时，旧批次标为“已被新批次替代”。")
    shot("07_ssui_post_S3.png", "修复后 SSUI 导入（课题三）")
    h2("7.2 模板填写（SSUI-POST-v1.1）")
    tb(["工作表", "内容", "要求"], [
        ["批次信息", "场地编号、评价轨道、评价年份、修复后年数 t、管理强度或 M、数据来源", f"M 须在方法区间内：生产 {W['M_range']['production'][0]}–{W['M_range']['production'][1]}，生态 {W['M_range']['ecology'][0]}–{W['M_range']['ecology'][1]}；数据来源须如实选择"],
        ["指标得分", "D1–D25：原始值（可选）、单位（可选）、得分 s{_i}（必填）", "s{_i} 取值 0–1，由课题组按分级标准赋分（方法文件未给出原始值到得分的换算规则）"],
        ["修复后污染物检测", "点位编号、pH、污染物、浓度、单位", "单位必须为 mg/kg；用于法规安全门禁，缺测项目不视为达标"]])
    h2("7.3 计算方法")
    p("SSUI = f(t) × Σ{_j}(v{_j} × S{_j}) × M，其中 S{_j} = Σ{_i} w{_ij} × s{_i}，f(t) = 1 + 0.03t。权重 v{_j}、w{_ij} 分轨道取自方法 PPT（生产第 14 页，生态第 13 页）。")
    tb(["SSUI", "等级"], [["0.8 – 1.0", "高度可持续"], ["0.6 – 0.8", "中度可持续"], ["0.4 – 0.6", "低度可持续"], ["< 0.4", "不可持续"]])
    note("由于 M ≥ 1.05、f(t) ≥ 1，SSUI 可能超过 1.0。系统不截断，显示原值，按“高度可持续”展示并标记“超出 0–1 区间，待课题组确认”。另外，方法文件中 C4 组内权重之和（生态 1.214，生产 1.047）不等于 1，系统按原值使用并提示。")
    p("旧版“SSUI参考评价（修复前）”页面仍保留，用于基于修复前检测数据的参考评价；正式 SSUI 以课题三页面为准。")
    shot("08_ssui_legacy_reference.png", "SSUI 参考评价（修复前数据，旧版口径）")

    h1("第八章 修复后利用方向结论（生产 / 生态）")
    h2("8.1 判定规则")
    ol("先过法规安全门禁：生产轨道按 GB 15618-2018（农用地筛选值 / 管制值，按 pH 和水田 / 其他分档），生态轨道按 GB 36600-2018 表 1 的 45 项基本项目（缺省按第一类用地）。",
       "门禁结果分四种：通过 / 条件通过（超筛选值但不超管制值）/ 未通过（超管制值）/ 证据不足。门禁不能被任何评分抵消。",
       "门禁通过（或生产轨道条件通过）后，再看功能评分：修复后用课题三 SSUI（暂定 ≥ 0.6 为支持），修复前用课题二重构可行性。",
       "必测项目缺测、阈值缺失、pH 或农用地类型未知导致无法判定时，一律为“证据不足”，并列出需补充的证据。")
    tb(["结论", "含义"], [
        ["支持生产利用", "生产门禁通过或条件通过，且功能评分支持；条件通过时附安全利用条件（农艺调控、替代种植、农产品协同监测）"],
        ["支持生态利用", "生态门禁通过且功能评分支持"],
        ["生产与生态均支持", "两轨都支持；两轨权重体系不同，分数差只作参考，最终方向由管理目标确定"],
        ["均不支持", "两轨都因门禁未通过或评分不支持而被否定"],
        ["证据不足", "至少一轨无法判定，且没有任何一轨得到支持"]])
    p("保守假设：生态用地类别未指定时按第一类用地；农用地类型未指定时，筛选值取水田 / 其他中较严者，管制值取各 pH 档中最宽者（只在确定超管制时判失败）。所有假设都写在结论的“假设与说明”中。")
    shot("13_post_decision.png", "修复后利用方向结论")
    h2("8.2 当前已知边界")
    note("GB 36600 表 1 中有 12 项挥发性有机物在系统中尚无经核实的阈值（氯甲烷、1,1-二氯乙烷、1,1-二氯乙烯、顺/反-1,2-二氯乙烯、1,2-二氯丙烷、1,1,1,2-/1,1,2,2-四氯乙烷、1,1,1-/1,1,2-三氯乙烷、1,2,3-三氯丙烷、间+对二甲苯）。因此生态门禁目前只能得到“未通过”或“证据不足”，不会得到“通过”。需由项目组按标准原文录入并复核后解除。")

    h1("第九章 方案推荐与报告")
    ul("方案推荐：基于 KOS 关键障碍因子匹配技术库，给出排序、匹配度、成本周期和禁用条件。",
       "报告：在场地追溯详情中生成 PDF、DOCX 或 HTML 报告。报告包含场地信息、数据来源与覆盖率、地图图件、检测摘要、障碍因子、重构可行性、SSUI、利用方向结论、推荐方案、五阶段追溯记录、附件清单、版本口径和人工复核意见区。",
       f"使用模拟数据的场地，报告封面显示红色“【{LABEL}，不得用于正式报告】”横幅。")
    shot("11_files.png", "文件管理")

    h1("第十章 备份与恢复")
    ol("系统管理 → 备份恢复 → “立即备份”：生成加密备份文件（%APPDATA%\\SRS\\backups）。系统每天 02:00 自动备份。",
       "恢复：选择备份 → 二次确认。恢复前，系统会自动保存当前数据库快照（文件名含 pre_restore）。",
       "恢复完成后重新登录，核对场地数量、课题三批次和利用结论。")
    shot("12_system.png", "系统管理")

    h1("第十一章 模拟数据与演示")
    ul(f"演示数据包（demo_montecarlo_v11）由固定随机种子生成，所有文件、场地名和结果均带“{LABEL}”标签，metadata.json 中记录分布参数和文件 SHA-256。",
       "演示请使用独立的演示数据库（文件名以 srs_demo 开头）；脚本会拒绝写入用户的正式数据库。",
       "演示脚本：python scripts/mc_demo_v11.py generate / run / clean。clean 只删除演示场地及其关联记录。",
       "模拟数据结果只说明软件流程，不代表任何真实场地，不得用于正式报告。")

    h1("第十二章 常见问题与故障排除")
    tb(["问题", "原因与处理"], [
        ["结论显示“证据不足”", "查看“需补充的证据”：通常是缺测必测项目、缺 pH 或农用地类型、功能评分未计算，或生态门禁阈值缺失（见 8.2）"],
        ["课题二显示“证据不足/无法评价”", "指标覆盖率低于门禁（生产 30%，生态 10%，暂定），请补充缺失指标"],
        ["课题三预览报错", "按“工作表/行/列”逐条修正；常见原因：得分超出 0–1、M 超出区间、单位不是 mg/kg、场地编号与所选场地不一致"],
        ["提示“同一文件已确认导入”", "相同内容的文件不能重复入库；如需更正，修改文件内容后重新上传"],
        ["SSUI 大于 1", "方法公式中 M ≥ 1.05 所致，系统按原值显示并标注待确认，不是错误"],
        ["启动后浏览器未打开或端口被占用", "手动访问 http://127.0.0.1:8000；若端口被占用，关闭占用程序或已在运行的 SRS 实例"],
        ["地图无底图", "离线时只显示点位；在线底图需联网，地图 Key 由安装包内置或通过环境变量 GAODE_KEY 配置（详见《地图离线方案与天地图配置》）"],
        ["监管账户无法导入", "监管角色只读，按设计不能导入数据或运行判定"]])

    h1("附录 A 课题三指标与权重（方法 PPT 第 13/14 页）")
    tb(["编码", "指标", "准则层", "生产权重", "生态权重"],
       [[i["code"], i["name"], i["criterion"], f"{i['w_production']:.4f}", f"{i['w_ecology']:.4f}"] for i in W["indicators"]])
    p(f"准则层权重 v{{_j}}：生产 {W['criterion_weights']['production']}；生态 {W['criterion_weights']['ecology']}。编号以第 13/14 页权重表为准，第 6 页层次图有 10 处编号不一致，已提交课题组确认。")
    h1("附录 B 法规阈值来源")
    ul("GB 15618-2018：表 1 风险筛选值（水田 / 其他，按 pH 分 4 档），表 3 风险管制值（镉、汞、砷、铅、铬）。文件：data/standards/gb15618_2018_v11.csv（含来源与置信度）。",
       "GB 36600-2018：表 1 重金属 7 项及已核实的有机物，第一类 / 第二类用地筛选值与管制值。文件：data/standards/gb36600_2018_v11.csv。")
    h1("附录 C 已知限制与待确认事项")
    ul("课题三尚无真实修复后数据，目前只用模拟数据演示。",
       "课题二子课题测试表中 7 项文本分级指标缺少“分级 → F 值”映射，暂不能完整计算。",
       "方法冲突（D 编号、C4 权重和、M 使 SSUI > 1、准则层组合方式、28 项权重和 1.0506 等）待陈亮、伟杰老师确认；相关结果均标注“暂定”。",
       "GB 36600 表 1 中 12 项 VOC 阈值待录入（见 8.2）。")
    return C


def _runs(par, text):
    import re
    for i, part in enumerate(re.split(r"\{_([a-z]+)\}", text)):
        if part:
            r = par.add_run(part)
            if i % 2 == 1:
                r.font.subscript = True
    return par


def _plain(text):
    import re
    return re.sub(r"\{_([a-z]+)\}", r"\1", text)


def build_docx(C, shots, out_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    doc = Document()
    sec = doc.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(2.2))
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for lvl, size in (("Heading 1", 16), ("Heading 2", 13)):
        s = doc.styles[lvl]; s.font.size = Pt(size); s.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        s.font.name = "Times New Roman"; s.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    for _ in range(6):
        doc.add_paragraph()
    for t, sz in ((TITLE, 22), (SUB, 18)):
        pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = pp.add_run(t); r.bold = True; r.font.size = Pt(sz)
    for t in ("委托单位：生态环境部土壤与农业农村生态环境监管技术中心", "研制单位：浙江大学环境与资源学院",
              f"版本：v{VERSION}　　日期：{date.today():%Y年%m月}"):
        pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; pp.add_run(t).font.size = Pt(12)
    doc.add_page_break()
    missing = []
    for kind, x in C:
        if kind == "h1":
            doc.add_heading(x, level=1)
        elif kind == "h2":
            doc.add_heading(x, level=2)
        elif kind == "p":
            _runs(doc.add_paragraph(), x)
        elif kind in ("ul", "ol"):
            for it in x:
                _runs(doc.add_paragraph(style="List Bullet" if kind == "ul" else "List Number"), it)
        elif kind == "note":
            pp = doc.add_paragraph(); r = pp.add_run("说明：" + _plain(x)); r.font.color.rgb = RGBColor(0x8A, 0x4B, 0x00)
        elif kind == "table":
            hdr, rows = x
            t = doc.add_table(rows=1, cols=len(hdr)); t.style = "Table Grid"
            for i, hh in enumerate(hdr):
                t.rows[0].cells[i].text = hh
                for r in t.rows[0].cells[i].paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = _plain(str(v))
            doc.add_paragraph()
        elif kind == "shot":
            f, cap = x
            path = os.path.join(shots, f) if shots else None
            if path and os.path.isfile(path):
                doc.add_picture(path, width=Cm(16.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                missing.append(f)
                pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = pp.add_run(f"【截图位：{f} —— 待从 v{VERSION} 实机采集，不使用旧版截图】"); r.bold = True
                r.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)
            cp = doc.add_paragraph(f"图：{cap}"); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(out_path)
    return missing


def _sub_html(t):
    import re
    return re.sub(r"\{_([a-z]+)\}", r"<sub>\1</sub>", H.escape(t))


def build_pdf(C, shots, out_path):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)
    pdfmetrics.registerFont(TTFont("Deng", os.path.join(FONT_DIR, "Deng.ttf")))
    pdfmetrics.registerFont(TTFont("DengB", os.path.join(FONT_DIR, "Dengb.ttf")))
    pdfmetrics.registerFont(TTFont("SimHei", os.path.join(FONT_DIR, "SimHei.ttf")))
    from reportlab.pdfbase.pdfmetrics import registerFontFamily
    registerFontFamily("Deng", normal="Deng", bold="DengB", italic="Deng", boldItalic="DengB")
    base = dict(fontName="Deng", fontSize=10.5, leading=16, wordWrap="CJK")
    sty = {"p": ParagraphStyle("p", **base, spaceAfter=4),
           "h1": ParagraphStyle("h1", fontName="SimHei", fontSize=16, leading=22, textColor=colors.HexColor("#1F3A5F"),
                                spaceBefore=10, spaceAfter=8, wordWrap="CJK"),
           "h2": ParagraphStyle("h2", fontName="SimHei", fontSize=12.5, leading=18, textColor=colors.HexColor("#1F3A5F"),
                                spaceBefore=8, spaceAfter=4, wordWrap="CJK"),
           "cell": ParagraphStyle("cell", fontName="Deng", fontSize=9, leading=12.5, wordWrap="CJK"),
           "th": ParagraphStyle("th", fontName="DengB", fontSize=9, leading=12.5, wordWrap="CJK"),
           "note": ParagraphStyle("note", **{**base, "fontSize": 10}, textColor=colors.HexColor("#8A4B00"),
                                  backColor=colors.HexColor("#FFF7E6"), borderPadding=5, spaceBefore=4, spaceAfter=8),
           "ph": ParagraphStyle("ph", **base, alignment=TA_CENTER, textColor=colors.HexColor("#B91C1C"),
                                borderColor=colors.HexColor("#B91C1C"), borderWidth=1, borderPadding=10, spaceBefore=18, spaceAfter=16),
           "cap": ParagraphStyle("cap", fontName="Deng", fontSize=9, leading=12, alignment=TA_CENTER,
                                 textColor=colors.HexColor("#555555"), spaceAfter=8),
           "cv1": ParagraphStyle("cv1", fontName="SimHei", fontSize=22, leading=32, alignment=TA_CENTER, wordWrap="CJK"),
           "cv2": ParagraphStyle("cv2", fontName="SimHei", fontSize=17, leading=26, alignment=TA_CENTER),
           "cv3": ParagraphStyle("cv3", fontName="Deng", fontSize=12, leading=20, alignment=TA_CENTER)}
    width = A4[0] - 4 * cm
    story = [Spacer(1, 6 * cm), Paragraph(H.escape(TITLE), sty["cv1"]), Spacer(1, 0.4 * cm),
             Paragraph(H.escape(SUB), sty["cv2"]), Spacer(1, 3 * cm)]
    for t in ("委托单位：生态环境部土壤与农业农村生态环境监管技术中心", "研制单位：浙江大学环境与资源学院",
              f"版本：v{VERSION}　　日期：{date.today():%Y年%m月}"):
        story.append(Paragraph(t, sty["cv3"]))
    story.append(PageBreak())
    for kind, x in C:
        if kind in ("h1", "h2"):
            story.append(Paragraph(_sub_html(x), sty[kind]))
        elif kind == "p":
            story.append(Paragraph(_sub_html(x), sty["p"]))
        elif kind == "note":
            story.append(Paragraph("说明：" + _sub_html(x), sty["note"]))
        elif kind in ("ul", "ol"):
            items = [ListItem(Paragraph(_sub_html(i), sty["p"]), leftIndent=14) for i in x]
            story.append(ListFlowable(items, bulletType="bullet" if kind == "ul" else "1", start=None if kind == "ul" else 1,
                                      bulletFontName="Deng", bulletFontSize=9, leftIndent=14))
        elif kind == "table":
            hdr, rows = x
            n = len(hdr)
            lens = [max([len(str(hdr[i]))] + [len(str(r[i])) for r in rows]) for i in range(n)]
            tot = sum(min(l, 40) + 4 for l in lens)
            cw = [width * (min(l, 40) + 4) / tot for l in lens]
            data = [[Paragraph(_sub_html(h), sty["th"]) for h in hdr]] + \
                   [[Paragraph(_sub_html(str(v)), sty["cell"]) for v in r] for r in rows]
            t = Table(data, colWidths=cw, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            story += [t, Spacer(1, 8)]
        elif kind == "shot":
            f, cap = x
            path = os.path.join(shots, f) if shots else None
            if path and os.path.isfile(path):
                from PIL import Image as PI
                w0, h0 = PI.open(path).size
                w1 = width; h1 = w1 * h0 / w0
                if h1 > 18 * cm:
                    h1 = 18 * cm; w1 = h1 * w0 / h0
                block = [Image(path, width=w1, height=h1)]
            else:
                block = [Paragraph(f"【截图位：{H.escape(f)} —— 待从 v{VERSION} 实机采集，不使用旧版截图】", sty["ph"])]
            story.append(KeepTogether(block + [Paragraph("图：" + H.escape(cap), sty["cap"])]))

    def footer(canvas, doc):
        canvas.saveState(); canvas.setFont("Deng", 8); canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawCentredString(A4[0] / 2, 1.0 * cm, f"SRS 用户操作手册 v{VERSION} · 第 {doc.page} 页")
        canvas.restoreState()
    doc = SimpleDocTemplate(out_path, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                            bottomMargin=2 * cm, title=f"{TITLE} {SUB}", author="浙江大学环境与资源学院")
    doc.build(story, onFirstPage=lambda c, d: None, onLaterPages=footer)
    return True


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--shots", default=None)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    C = content()
    base = os.path.join(a.out, f"SRS_用户操作手册_v{VERSION}")
    miss = build_docx(C, a.shots, base + ".docx")
    ok = build_pdf(C, a.shots, base + ".pdf")
    print(json.dumps({"docx": base + ".docx", "pdf": base + ".pdf", "pdf_ok": ok, "missing_screenshots": miss},
                     ensure_ascii=False, indent=1))
